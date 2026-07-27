#!/usr/bin/env python3
"""将 scores.jsonl 导出为带指标说明子页的 Excel。"""

from __future__ import annotations

import argparse
import json
import re
import sys
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

REPO_ROOT = Path(__file__).resolve().parents[1]
if str(REPO_ROOT) not in sys.path:
    sys.path.insert(0, str(REPO_ROOT))

from plugins.text_description_evaluator.evaluation_export import (  # noqa: E402
    backfill_gate_thresholds_in_jsonl,
    load_gate_thresholds_from_fashion_config,
)

SPEC_PATH = REPO_ROOT / "plugins/text_description_evaluator/fashion_prompt_optimizer_spec.json"

COL_MAP = [
    ("chapter_idx", "章节"),
    ("look_number", "Look 序号"),
    ("source_txt", "源 txt 文件"),
    ("chapter_dir", "章节目录"),
    ("md_index", "序号"),
    ("source_name", "源文件名"),
    ("evaluator_spec", "评估规范"),
    ("spec_version", "规范版本"),
    ("judge_model", "裁判模型"),
    ("total_score_S_fp", "综合总分 S_fp"),
    ("score_band", "评分等级"),
    ("gates_both_passed", "双门限通过"),
    ("score_gate_passed", "得分门限通过"),
    ("score_gate_value", "得分门限判定值"),
    ("score_gate_threshold", "得分门限阈值"),
    ("penalty_gate_passed", "惩罚门限通过"),
    ("penalty_gate_total_penalty", "惩罚门限判定值"),
    ("penalty_gate_threshold", "惩罚门限阈值"),
    ("coverage_axis_score", "覆盖度轴得分"),
    ("quality_axis_score_raw", "质量轴得分(原始)"),
    ("quality_base_score", "质量基础分"),
    ("penalties_total_penalty_mean", "惩罚项均值"),
    ("total_defined_metrics", "定义指标总数"),
    ("total_applicable_metrics", "适用指标数"),
    ("total_hit_metrics", "命中指标数"),
    ("total_skipped_metrics", "跳过指标数"),
    ("R_content_scalar", "R_content 标量"),
    ("r_content_enabled", "R_content 启用"),
    ("prompt_char_len", "文本字符长度"),
    ("source_md", "源文件路径"),
]

FIELD_DEFS = [
    ("score_gate_value", "得分门限判定值", "coverage×0.4 + quality_base×0.6（不含 penalty 扣减），用于与得分门限阈值比较。"),
    ("score_gate_threshold", "得分门限阈值", "规范/配置中的 score_gate_min，默认 0.7；判定值 ≥ 阈值则 score_gate_passed=true。"),
    ("penalty_gate_total_penalty", "惩罚门限判定值", "四项惩罚的算术平均，用于与惩罚门限阈值比较。"),
    ("penalty_gate_threshold", "惩罚门限阈值", "规范/配置中的 penalty_gate_max，默认 0.5；判定值 ≤ 阈值则 penalty_gate_passed=true。"),
    ("penalties_total_penalty_mean", "惩罚项均值", "与惩罚门限判定值相同，来自 quality.penalties.total_penalty。"),
    ("total_score_S_fp", "综合总分 S_fp", "min(覆盖度×0.4 + 质量有效分×0.6, total_cap)，范围 0~1；惩罚不直接扣减此项。"),
]


def _style_workbook(path: Path) -> None:
    wb = load_workbook(path)
    header_fill = PatternFill("solid", fgColor="4472C4")
    header_font = Font(color="FFFFFF", bold=True)
    wrap = Alignment(wrap_text=True, vertical="top")
    for ws in wb.worksheets:
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.freeze_panes = "A2"
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                if cell.value is not None:
                    max_len = max(max_len, min(len(str(cell.value)), 80))
                cell.alignment = wrap
            ws.column_dimensions[col_letter].width = min(max(max_len + 2, 10), 60)
    wb.save(path)


def _normalize_score_row(obj: dict) -> dict:
    if obj.get("chapter_idx") is None and obj.get("chapter_dir"):
        m = re.search(r"chapter_(\d+)", str(obj["chapter_dir"]).replace("\\", "/"))
        if m:
            obj["chapter_idx"] = int(m.group(1))
    mt = obj.pop("metric_totals", None) or {}
    obj.update(
        {
            "total_defined_metrics": mt.get("total_defined_metrics"),
            "total_applicable_metrics": mt.get("total_applicable_metrics"),
            "total_hit_metrics": mt.get("total_hit_metrics"),
            "total_skipped_metrics": mt.get("total_skipped_metrics"),
        }
    )
    return obj


def _read_jsonl_rows(jsonl_paths: list[Path]) -> list[dict]:
    rows: list[dict] = []
    for jsonl_path in jsonl_paths:
        with Path(jsonl_path).open(encoding="utf-8") as f:
            for line in f:
                line = line.strip()
                if not line:
                    continue
                rows.append(_normalize_score_row(json.loads(line)))
    rows.sort(key=lambda r: (r.get("chapter_idx") or 0, r.get("look_number") or 0, r.get("md_index") or 0))
    return rows


def export_scores_jsonl_to_excel(
    jsonl_paths: Path | list[Path],
    out_path: Path,
    *,
    spec_path: Path = SPEC_PATH,
) -> int:
    if isinstance(jsonl_paths, Path):
        paths = [jsonl_paths]
    else:
        paths = list(jsonl_paths)
    rows = _read_jsonl_rows(paths)

    df = pd.DataFrame(rows)
    ordered_cols = [k for k, _ in COL_MAP if k in df.columns]
    df = df[ordered_cols].rename(columns={k: v for k, v in COL_MAP if k in df.columns})

    numeric_cols = [
        "综合总分 S_fp", "覆盖度轴得分", "质量轴得分(原始)", "质量基础分",
        "惩罚项均值", "得分门限判定值", "得分门限阈值",
        "惩罚门限判定值", "惩罚门限阈值", "R_content 标量", "文本字符长度",
        "适用指标数", "命中指标数",
    ]
    summary_rows = []
    for c in numeric_cols:
        if c not in df.columns:
            continue
        s = pd.to_numeric(df[c], errors="coerce")
        summary_rows.append(
            {
                "指标": c,
                "样本数": int(s.notna().sum()),
                "均值": round(float(s.mean()), 4) if s.notna().any() else None,
                "最小值": round(float(s.min()), 4) if s.notna().any() else None,
                "最大值": round(float(s.max()), 4) if s.notna().any() else None,
                "标准差": round(float(s.std()), 4) if s.notna().sum() > 1 else None,
            }
        )

    band_counts = df["评分等级"].value_counts().reset_index() if "评分等级" in df.columns else pd.DataFrame()
    if not band_counts.empty:
        band_counts.columns = ["评分等级", "数量"]

    chapter_summary = pd.DataFrame()
    if "章节" in df.columns and "综合总分 S_fp" in df.columns:
        chapter_summary = (
            df.groupby("章节", dropna=False)
            .agg(
                样本数=("综合总分 S_fp", "count"),
                平均S_fp=("综合总分 S_fp", "mean"),
                最低S_fp=("综合总分 S_fp", "min"),
                最高S_fp=("综合总分 S_fp", "max"),
                平均R_content=("R_content 标量", "mean"),
                双门限通过数=("双门限通过", lambda s: int(pd.Series(s).astype(bool).sum())),
            )
            .reset_index()
        )
        for col in ("平均S_fp", "最低S_fp", "最高S_fp", "平均R_content"):
            if col in chapter_summary.columns:
                chapter_summary[col] = chapter_summary[col].round(4)

    spec = json.loads(Path(spec_path).read_text(encoding="utf-8"))
    axis_zh = {"coverage_score": "覆盖度", "quality_score": "质量", "bonus_score": "加分项"}

    all_field_defs = [
        ("chapter_idx", "章节", "workflow 章节编号（chapter_01 → 1）。"),
        ("look_number", "Look 序号", "章节内 look 编号。"),
        ("source_txt", "源 txt 文件", "被评分的 look_*.txt 文件名。"),
        ("chapter_dir", "章节目录", "章节所在目录路径。"),
        ("md_index", "序号", "源 markdown 文件在批次中的编号（1-based）。"),
        ("source_name", "源文件名", "被评估的 text_description 文件名。"),
        ("source_md", "源文件路径", "源文件在项目中的相对路径。"),
        ("evaluator_spec", "评估规范", "评估器规范名称。"),
        ("spec_version", "规范版本", "评估规范版本号。"),
        ("judge_model", "裁判模型", "LLM-as-a-judge 模型名称。"),
        ("total_score_S_fp", "综合总分 S_fp", "min(覆盖度×0.4 + 质量有效分×0.6, total_cap)，范围 0~1；惩罚经 R_content 与惩罚门限单独体现。"),
        ("score_band", "评分等级", "≥0.90 Excellent；≥0.75 Strong；≥0.55 Usable；≥0.35 Weak；<0.35 Poor。"),
        ("gates_both_passed", "双门限通过", "得分门限与惩罚门限均通过。"),
        ("score_gate_passed", "得分门限通过", "得分门限判定值 ≥ 得分门限阈值。"),
        ("score_gate_value", "得分门限判定值", "coverage×0.4 + quality_base×0.6（不含 penalty）。"),
        ("score_gate_threshold", "得分门限阈值", "score_gate_min，默认 0.7。"),
        ("penalty_gate_passed", "惩罚门限通过", "惩罚门限判定值 ≤ 惩罚门限阈值。"),
        ("penalty_gate_total_penalty", "惩罚门限判定值", "四项惩罚算术平均，用于惩罚门限比较。"),
        ("penalty_gate_threshold", "惩罚门限阈值", "penalty_gate_max，默认 0.5。"),
        ("coverage_axis_score", "覆盖度轴得分", "适用 coverage 指标命中率均值，权重 0.4。"),
        ("quality_axis_score_raw", "质量轴得分(原始)", "适用 quality 指标 score_value 均值。"),
        ("quality_base_score", "质量基础分", "质量轴基础分，参与 S_fp 加权与 score_gate 计算。"),
        ("penalties_total_penalty_mean", "惩罚项均值", "四项惩罚算术平均；用于惩罚门限与 R_content，不直接扣减质量基础分。"),
        ("total_defined_metrics", "定义指标总数", "规范中定义的指标总数。"),
        ("total_applicable_metrics", "适用指标数", "本条 applicable 的指标数。"),
        ("total_hit_metrics", "命中指标数", "适用且 hit=1 的指标数。"),
        ("total_skipped_metrics", "跳过指标数", "不适用指标数。"),
        ("R_content_scalar", "R_content 标量", "RL/GRPO 内容主导奖励标量。"),
        ("r_content_enabled", "R_content 启用", "是否启用 R_content。"),
        ("prompt_char_len", "文本字符长度", "text_description 字符数。"),
    ]
    df_fields = pd.DataFrame(all_field_defs, columns=["字段名", "中文名", "含义说明"])

    registry = spec.get("metric_registry", {})
    module_map = {}
    for mod_list_key in ("coverage_modules", "quality_modules", "bonus_modules"):
        for mod in spec.get(mod_list_key, []):
            for m in mod.get("metrics", []):
                module_map[m] = mod.get("name", "")

    metric_rows = []
    for name, cfg in sorted(registry.items()):
        axis = cfg.get("axis", "")
        qdim = ""
        if axis == "quality_score":
            for qm in spec.get("quality_modules", []):
                if name in qm.get("metrics", []):
                    qreg = spec.get("quality_dimension_registry", {}).get(qm.get("name"), {})
                    qdim = qreg.get("zh_name", qm.get("name", ""))
                    break
        metric_rows.append(
            {
                "指标英文名": name,
                "所属模块": module_map.get(name, ""),
                "评分轴": axis_zh.get(axis, axis),
                "适用条件": cfg.get("applicability", ""),
                "判定规则": cfg.get("rule", ""),
                "质量维度": qdim,
            }
        )
    df_metrics = pd.DataFrame(metric_rows)

    penalty_rows = []
    for key, cfg in spec.get("penalty_registry", {}).items():
        penalty_rows.append(
            {
                "惩罚项英文名": key,
                "中文名": cfg.get("zh_name", ""),
                "说明": cfg.get("description", ""),
                "允许分值": str(cfg.get("allowed_scores", [])),
            }
        )
    df_penalties = pd.DataFrame(penalty_rows)

    gates = spec.get("optimization_gates", {})
    df_gates = pd.DataFrame(
        [
            {"项目": "得分门限阈值", "值": gates.get("score_gate_min", 0.7), "说明": "score_gate_value ≥ 此值 → 通过"},
            {"项目": "惩罚门限阈值", "值": gates.get("penalty_gate_max", 0.5), "说明": "penalty_gate_total_penalty ≤ 此值 → 通过"},
        ]
    )

    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(out_path, engine="openpyxl") as writer:
        df.to_excel(writer, sheet_name="评分明细", index=False)
        pd.DataFrame(summary_rows).to_excel(writer, sheet_name="汇总统计", index=False)
        if not band_counts.empty:
            band_counts.to_excel(writer, sheet_name="等级分布", index=False)
        if not chapter_summary.empty:
            chapter_summary.to_excel(writer, sheet_name="章节汇总", index=False)
        df_fields.to_excel(writer, sheet_name="字段说明", index=False)
        df_gates.to_excel(writer, sheet_name="门限说明", index=False)
        df_metrics[df_metrics["评分轴"] == "覆盖度"].to_excel(writer, sheet_name="覆盖度指标", index=False)
        df_metrics[df_metrics["评分轴"] == "质量"].to_excel(writer, sheet_name="质量指标", index=False)
        df_penalties.to_excel(writer, sheet_name="惩罚项说明", index=False)

    _style_workbook(out_path)
    return len(df)


def parse_args() -> argparse.Namespace:
    p = argparse.ArgumentParser(description="导出 scores.jsonl 为 Excel")
    p.add_argument("--jsonl", type=Path, nargs="+", required=True, help="一个或多个 scores.jsonl 路径")
    p.add_argument("--out", type=Path, required=True, help="输出 xlsx 路径")
    p.add_argument(
        "--backfill-thresholds",
        action="store_true",
        help="导出前为 jsonl 补写门限字段（缺省时用 fashion_config.yaml）",
    )
    p.add_argument("--config", type=Path, default=REPO_ROOT / "fashion_config.yaml")
    return p.parse_args()


def main() -> None:
    args = parse_args()
    if args.backfill_thresholds:
        th = load_gate_thresholds_from_fashion_config(args.config)
        for jsonl_path in args.jsonl:
            n = backfill_gate_thresholds_in_jsonl(
                jsonl_path,
                score_gate_threshold=th["score_gate_threshold"],
                penalty_gate_threshold=th["penalty_gate_threshold"],
            )
            print(f"backfilled {n} rows in {jsonl_path}")
    n = export_scores_jsonl_to_excel(args.jsonl, args.out)
    print(f"exported {n} rows -> {args.out}")


if __name__ == "__main__":
    main()
