"""Generate a 16:9 spec briefing PPT matching reflection_prompt.pptx style."""

from __future__ import annotations

import math
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

NSMAP_A = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}

SLIDE_W = Emu(12192000)
SLIDE_H = Emu(6858000)
LEFT = Emu(634702)
TITLE_TOP = Emu(517738)
TITLE_W = Emu(10800000)
TITLE_H = Emu(620000)
LABEL_TOP = Emu(1180000)
LABEL_H = Emu(320000)
BODY_TOP = Emu(1580000)
BODY_W = Emu(10880000)
BODY_H = Emu(4900000)
FOOT_TOP = Emu(6480000)

BLACK = RGBColor(0x1A, 0x1A, 0x1A)
GRAY = RGBColor(0x5A, 0x5A, 0x5A)
MUTED = RGBColor(0x8A, 0x8A, 0x8A)
LINE = RGBColor(0xD0, 0xD0, 0xD0)
HEADER_BG = RGBColor(0xF4, 0xF4, 0xF4)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
FONT = "微软雅黑"
FONT_EN = "Calibri"
FONT_MONO = "Consolas"

A = 0.002948
X0 = 7.3478
L_REF = round(math.exp(X0) - 1)
W_MODULES = {
    "DesignMerit": 10.0,
    "ConcisenessAndDensity": 3.0,
    "GenerationReadiness": 3.0,
    "BindingAccuracy": 1.0,
    "LanguageClarity": 1.0,
    "StructuralClarity": 1.0,
}
W_TOTAL = sum(W_MODULES.values())

# Distinct module colors: saturated chips + light row tints.
# Keep these in one place so the overview table and later identifier bars stay in sync.
MODULE_PALETTE = [
    {
        "key": "DesignMerit",
        "zh": "设计价值",
        "color": RGBColor(0x2F, 0x4B, 0x8A),  # indigo
    },
    {
        "key": "ConcisenessAndDensity",
        "zh": "可见性",
        "color": RGBColor(0x1A, 0x7A, 0x72),  # teal
    },
    {
        "key": "GenerationReadiness",
        "zh": "生成适配",
        "color": RGBColor(0xC2, 0x78, 0x00),  # amber
    },
    {
        "key": "BindingAccuracy",
        "zh": "属性绑定",
        "color": RGBColor(0xA3, 0x3B, 0x5C),  # rose
    },
    {
        "key": "LanguageClarity",
        "zh": "语言清晰",
        "color": RGBColor(0x4A, 0x7C, 0x3F),  # green
    },
    {
        "key": "StructuralClarity",
        "zh": "结构清晰",
        "color": RGBColor(0x6B, 0x4C, 0x9A),  # violet
    },
]
MODULE_BY_KEY = {m["key"]: m for m in MODULE_PALETTE}
CHIP_PREFIX = "ModuleChip_"
BAR_PREFIX = "ModuleBar_"
CHIP_TOP = Emu(90000)
CHIP_H = Emu(250000)
CHIP_GAP = Emu(36000)


def mix_white(rgb: RGBColor, p: float) -> RGBColor:
    return RGBColor(
        int(rgb[0] + (255 - rgb[0]) * p),
        int(rgb[1] + (255 - rgb[1]) * p),
        int(rgb[2] + (255 - rgb[2]) * p),
    )


def hex_of(rgb: RGBColor) -> str:
    return f"{rgb[0]:02X}{rgb[1]:02X}{rgb[2]:02X}"


def set_run_font(run, *, name=FONT, size=16, bold=False, color=BLACK, italic=False):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    run.font.name = name
    rPr = run._r.get_or_add_rPr()
    rFonts = rPr.find(qn("a:latin"))
    if rFonts is None:
        rFonts = etree.SubElement(rPr, qn("a:latin"))
    rFonts.set("typeface", name)
    ea = rPr.find(qn("a:ea"))
    if ea is None:
        ea = etree.SubElement(rPr, qn("a:ea"))
    ea.set("typeface", FONT)
    cs = rPr.find(qn("a:cs"))
    if cs is None:
        cs = etree.SubElement(rPr, qn("a:cs"))
    cs.set("typeface", name)


def add_textbox(slide, left, top, width, height, *, wrap=True):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = wrap
    tf.auto_size = None
    return tf


def add_title(slide, text):
    tf = add_textbox(slide, LEFT, TITLE_TOP, TITLE_W, TITLE_H, wrap=False)
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = text
    set_run_font(run, name=FONT_EN, size=32, bold=True)
    return tf


def add_label(slide, text):
    tf = add_textbox(slide, LEFT, LABEL_TOP, TITLE_W, LABEL_H, wrap=False)
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = text
    set_run_font(run, name=FONT, size=16, bold=True, color=GRAY)
    return tf


def add_footer(slide, text="fashion_prompt_optimizer_spec.json  v0.4.5"):
    tf = add_textbox(slide, LEFT, FOOT_TOP, BODY_W, Emu(280000), wrap=False)
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = text
    set_run_font(run, name=FONT_EN, size=10, color=MUTED)


def clear_first_paragraph(tf):
    p = tf.paragraphs[0]
    p.clear()
    return p


def add_blank(tf, size=8):
    p = tf.add_paragraph()
    p.space_before = Pt(0)
    p.space_after = Pt(0)
    run = p.add_run()
    run.text = ""
    set_run_font(run, size=size)
    return p


def add_rich_paragraph(tf, parts, *, first=False, size=16, space_after=6):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.space_before = Pt(0)
    p.space_after = Pt(space_after)
    p.line_spacing = 1.15
    for text, kwargs in parts:
        run = p.add_run()
        run.text = text
        set_run_font(run, size=size, **kwargs)
    return p


def section_parts(label, text, *, label_en=False):
    label_font = FONT_EN if label_en else FONT
    return [
        (f"[ {label} ] ", {"name": label_font, "bold": True, "color": BLACK}),
        (text, {"name": FONT, "bold": False, "color": BLACK}),
    ]


def add_sections(tf, sections, *, size=16, gap=True):
    first = True
    for item in sections:
        if not first and gap:
            add_blank(tf, size=6)
        if isinstance(item, str):
            add_rich_paragraph(tf, [(item, {"name": FONT})], first=first, size=size)
        else:
            add_rich_paragraph(tf, item, first=first, size=size)
        first = False


def set_cell_font(cell, text, *, size=12, bold=False, color=BLACK, name=FONT, align=None):
    cell.text = ""
    tf = cell.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    if align is not None:
        p.alignment = align
    p.space_before = Pt(2)
    p.space_after = Pt(2)
    run = p.add_run()
    run.text = text
    set_run_font(run, name=name, size=size, bold=bold, color=color)
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE


def shade_cell(cell, rgb: RGBColor):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    for child in list(tcPr):
        if child.tag == qn("a:solidFill"):
            tcPr.remove(child)
    solid = etree.SubElement(tcPr, qn("a:solidFill"))
    srgb = etree.SubElement(solid, qn("a:srgbClr"))
    srgb.set("val", f"{rgb[0]:02X}{rgb[1]:02X}{rgb[2]:02X}")


def iter_table_cells(table):
    for row in table.rows:
        for cell in row.cells:
            yield cell


def set_table_borders(table, color="D0D0D0", width="6350"):
    for cell in iter_table_cells(table):
        tc = cell._tc
        tcPr = tc.get_or_add_tcPr()
        for edge in ("lnL", "lnR", "lnT", "lnB"):
            for old in tcPr.findall(qn(f"a:{edge}")):
                tcPr.remove(old)
            ln = etree.SubElement(tcPr, qn(f"a:{edge}"))
            ln.set("w", width)
            ln.set("cap", "flat")
            ln.set("cmpd", "sng")
            ln.set("algn", "ctr")
            sf = etree.SubElement(ln, qn("a:solidFill"))
            srgb = etree.SubElement(sf, qn("a:srgbClr"))
            srgb.set("val", color)
            etree.SubElement(ln, qn("a:prstDash")).set("val", "solid")


def add_table(slide, left, top, width, rows, col_widths, *, header=True, font_size=12, row_h=360000):
    n_rows = len(rows)
    n_cols = len(rows[0])
    table_shape = slide.shapes.add_table(n_rows, n_cols, left, top, width, Emu(int(row_h * n_rows)))
    table = table_shape.table
    total_w = sum(col_widths)
    for i, w in enumerate(col_widths):
        table.columns[i].width = int(width * w / total_w)
    for r_idx, row in enumerate(rows):
        for c_idx, val in enumerate(row):
            cell = table.cell(r_idx, c_idx)
            is_header = header and r_idx == 0
            set_cell_font(
                cell,
                val,
                size=font_size,
                bold=is_header,
                name=FONT,
            )
            if is_header:
                shade_cell(cell, HEADER_BG)
    set_table_borders(table)
    return table_shape


def add_accent_line(slide, color=BLACK):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        LEFT,
        Emu(1480000),
        Emu(900000),
        Emu(14000),
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()


def remove_prefixed_shapes(slide, prefix):
    for shape in list(slide.shapes):
        if shape.name.startswith(prefix):
            shape._element.getparent().remove(shape._element)


def _set_shape_text_anchor(shape, anchor="ctr"):
    tf = shape.text_frame
    tf.word_wrap = False
    tf.auto_size = None
    body_pr = tf._txBody.find(qn("a:bodyPr"))
    if body_pr is not None:
        body_pr.set("anchor", anchor)
        body_pr.set("lIns", "40000")
        body_pr.set("rIns", "40000")
        body_pr.set("tIns", "20000")
        body_pr.set("bIns", "20000")


def add_module_chips(slide, active="all"):
    """Color blocks at the top. Overview lights all six; later slides light the current module(s)."""
    remove_prefixed_shapes(slide, CHIP_PREFIX)
    if active is None:
        return
    active_keys = {m["key"] for m in MODULE_PALETTE} if active == "all" else set(active)
    n = len(MODULE_PALETTE)
    chip_w = int((int(BODY_W) - int(CHIP_GAP) * (n - 1)) / n)
    x = int(LEFT)
    for mod in MODULE_PALETTE:
        is_on = mod["key"] in active_keys
        fill = mod["color"] if is_on else mix_white(mod["color"], 0.82)
        text_color = WHITE if is_on else mix_white(mod["color"], 0.25)
        shape = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            Emu(x),
            CHIP_TOP,
            Emu(chip_w),
            CHIP_H,
        )
        shape.name = f"{CHIP_PREFIX}{mod['key']}"
        shape.fill.solid()
        shape.fill.fore_color.rgb = fill
        shape.line.fill.background()
        _set_shape_text_anchor(shape)
        tf = shape.text_frame
        p = tf.paragraphs[0]
        p.clear()
        p.alignment = PP_ALIGN.CENTER
        p.space_before = Pt(0)
        p.space_after = Pt(0)
        run = p.add_run()
        run.text = mod["zh"]
        set_run_font(run, name=FONT, size=11, bold=True, color=text_color)
        x += chip_w + int(CHIP_GAP)


def color_module_table(table_shape):
    table = table_shape.table
    if len(table.rows) < 2:
        return
    for i, mod in enumerate(MODULE_PALETTE):
        r_idx = i + 1
        if r_idx >= len(table.rows):
            break
        tint = mix_white(mod["color"], 0.88)
        strong = mix_white(mod["color"], 0.72)
        for c_idx in range(len(table.columns)):
            cell = table.cell(r_idx, c_idx)
            shade_cell(cell, strong if c_idx == 0 else tint)
            color = mod["color"] if c_idx == 0 else BLACK
            bold = c_idx == 0
            for p in cell.text_frame.paragraphs:
                for run in p.runs:
                    if run.text:
                        run.font.color.rgb = color
                        run.font.bold = bold
        first = table.cell(r_idx, 0)
        tc_pr = first._tc.get_or_add_tcPr()
        for old in tc_pr.findall(qn("a:lnL")):
            tc_pr.remove(old)
        ln = etree.SubElement(tc_pr, qn("a:lnL"))
        ln.set("w", "50800")
        ln.set("cap", "flat")
        ln.set("cmpd", "sng")
        ln.set("algn", "ctr")
        sf = etree.SubElement(ln, qn("a:solidFill"))
        srgb = etree.SubElement(sf, qn("a:srgbClr"))
        srgb.set("val", hex_of(mod["color"]))
        etree.SubElement(ln, qn("a:prstDash")).set("val", "solid")


def add_table_module_bars(slide, table_shape):
    remove_prefixed_shapes(slide, BAR_PREFIX)
    n_rows = len(table_shape.table.rows)
    if n_rows < 2:
        return
    row_h = int(table_shape.height / n_rows)
    bar_w = Emu(70000)
    for i, mod in enumerate(MODULE_PALETTE):
        top = int(table_shape.top) + row_h * (i + 1)
        bar = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            int(table_shape.left) - int(bar_w),
            top,
            bar_w,
            row_h,
        )
        bar.name = f"{BAR_PREFIX}{mod['key']}"
        bar.fill.solid()
        bar.fill.fore_color.rgb = mod["color"]
        bar.line.fill.background()


def slide_heading(slide):
    for shape in slide.shapes:
        if not shape.has_text_frame:
            continue
        if abs(int(shape.top) - int(TITLE_TOP)) > 90000:
            continue
        text = shape.text_frame.text.strip()
        if text and not text.startswith("fashion_prompt"):
            return text
    return ""


def active_modules_for_title(title):
    t = title.lower()
    if "quality modules" in t:
        return "all"
    if "design merit" in t:
        return ["DesignMerit"]
    if "visibility" in t and "generation" in t:
        return ["ConcisenessAndDensity", "GenerationReadiness"]
    if "visibility" in t:
        return ["ConcisenessAndDensity"]
    if "generation" in t:
        return ["GenerationReadiness"]
    if "binding" in t and "language" in t:
        return ["BindingAccuracy", "LanguageClarity", "StructuralClarity"]
    if "binding" in t:
        return ["BindingAccuracy"]
    if "language" in t:
        return ["LanguageClarity"]
    if "structur" in t:
        return ["StructuralClarity"]
    return None


def recolor_accent_line(slide, color):
    for shape in slide.shapes:
        if abs(int(shape.top) - 1480000) > 30000:
            continue
        if int(shape.height) > 40000:
            continue
        try:
            shape.fill.solid()
            shape.fill.fore_color.rgb = color
            shape.line.fill.background()
        except Exception:
            continue


def colorize_existing_presentation(prs):
    """Paint module colors onto an already-built deck without rebuilding custom slides."""
    for slide in prs.slides:
        title = slide_heading(slide)
        active = active_modules_for_title(title)
        if active is None:
            continue
        add_module_chips(slide, active=active)
        if active == "all":
            for shape in slide.shapes:
                if shape.has_table and shape.table.cell(0, 0).text.strip() == "模块":
                    color_module_table(shape)
                    add_table_module_bars(slide, shape)
                    break
        else:
            recolor_accent_line(slide, MODULE_BY_KEY[active[0]]["color"])


def save_presentation(prs, path: Path):
    try:
        path.parent.mkdir(parents=True, exist_ok=True)
        prs.save(path)
        print(f"Saved: {path}")
        return path
    except PermissionError:
        extra = "_updated" if "模块配色" in path.stem else "_模块配色"
        fallback = path.with_name(path.stem + extra + path.suffix)
        prs.save(fallback)
        print(f"File locked, saved: {fallback}")
        return fallback


def blank_layout(prs):
    for layout in prs.slide_layouts:
        if layout.name in ("空白", "Blank"):
            return layout
    return prs.slide_layouts[min(6, len(prs.slide_layouts) - 1)]


def new_slide(prs):
    slide = prs.slides.add_slide(blank_layout(prs))
    add_footer(slide)
    return slide


def move_slide(prs, old_index, new_index):
    sld_id_lst = prs.slides._sldIdLst
    els = list(sld_id_lst)
    el = els[old_index]
    sld_id_lst.remove(el)
    sld_id_lst.insert(new_index, el)


def add_slide_at(prs, index):
    prs.slides.add_slide(blank_layout(prs))
    move_slide(prs, len(prs.slides) - 1, index)
    return prs.slides[index]


def tint_metric_rows(table_shape, row_keys):
    table = table_shape.table
    for i, key in enumerate(row_keys):
        r_idx = i + 1
        if r_idx >= len(table.rows):
            break
        mod = MODULE_BY_KEY[key]
        for c_idx in range(len(table.columns)):
            cell = table.cell(r_idx, c_idx)
            shade_cell(cell, mix_white(mod["color"], 0.72 if c_idx == 0 else 0.90))
            if c_idx == 0:
                for p in cell.text_frame.paragraphs:
                    for run in p.runs:
                        if run.text:
                            run.font.color.rgb = mod["color"]
                            run.font.bold = True


def populate_visibility_metrics(slide):
    key = "ConcisenessAndDensity"
    add_title(slide, "Visibility Metrics")
    add_label(slide, "打分量规")
    add_accent_line(slide, MODULE_BY_KEY[key]["color"])
    add_module_chips(slide, active=[key])
    tf = add_textbox(slide, LEFT, BODY_TOP, BODY_W, Emu(700000))
    add_sections(tf, [
        section_parts(
            "怎么打分",
            "单指标五档 0 / 0.25 / 0.5 / 0.75 / 1.0。模块分就是这一档。标准 T2I 开场白不计罚。",
        ),
    ], size=15, gap=False)
    tbl = add_table(
        slide,
        LEFT,
        Emu(2380000),
        BODY_W,
        [
            ["分数", "标准"],
            ["1.0", "可见主体（廓形 / 品类 / 层次 / 主色 / 关键工艺 / 可见配件）始终居前且占主导；隐藏信息只作轻量补充"],
            ["0.75", "可见主体明确优先；最多 1 处轻量氛围句，不分散生图焦点"],
            ["0.5", "氛围 / 姿态 / 过渡铺垫占明显篇幅，或低可见细节与主体争抢；或可见事实重复堆砌"],
            ["0.25", "姿态、抽象场域 / 身份评论喧宾夺主；或 ≥3 处沙龙 / 漫步 / 仿佛 / 暗示类句稀释可见主体"],
            ["0.0", "几乎不以可见服装为主体；非成像叙述淹没成像信息"],
        ],
        [1.1, 8.4],
        font_size=12,
        row_h=400000,
    )
    tint_metric_rows(tbl, [key] * 5)
    tf2 = add_textbox(slide, LEFT, Emu(5000000), BODY_W, Emu(900000))
    add_sections(tf2, [
        section_parts(
            "硬封顶",
            "姿态 / 走位占显著篇幅 → 最高 0.5；非成像评论明显多于可见事实 → 最高 0.25。",
        ),
    ], size=14, gap=False)


def populate_score_band_slide(
    slide,
    *,
    key,
    title,
    intro,
    bands,
    footnote=None,
    footnote_label="备注",
):
    add_title(slide, title)
    add_label(slide, "打分量规")
    add_accent_line(slide, MODULE_BY_KEY[key]["color"])
    add_module_chips(slide, active=[key])
    tf = add_textbox(slide, LEFT, BODY_TOP, BODY_W, Emu(700000))
    add_sections(tf, [section_parts("怎么打分", intro)], size=15, gap=False)
    rows = [["分数", "标准"]] + [[score, text] for score, text in bands]
    tbl = add_table(
        slide,
        LEFT,
        Emu(2380000),
        BODY_W,
        rows,
        [1.1, 8.4],
        font_size=12,
        row_h=400000,
    )
    tint_metric_rows(tbl, [key] * len(bands))
    if footnote:
        tf2 = add_textbox(slide, LEFT, Emu(5000000), BODY_W, Emu(900000))
        add_sections(tf2, [section_parts(footnote_label, footnote)], size=14, gap=False)


def populate_generation_metrics(slide):
    key = "GenerationReadiness"
    add_title(slide, "Generation Metrics")
    add_label(slide, "打分量规")
    add_accent_line(slide, MODULE_BY_KEY[key]["color"])
    add_module_chips(slide, active=[key])
    tf = add_textbox(slide, LEFT, BODY_TOP, BODY_W, Emu(620000))
    add_sections(tf, [
        section_parts(
            "怎么打分",
            "五档 1.0 / 0.75 / 0.5 / 0.25 / 0.0。适用指标取平均。生图适配始终计；左右、空间只在用得上时计。",
        ),
    ], size=14, gap=False)
    tbl = add_table(
        slide,
        LEFT,
        Emu(2280000),
        BODY_W,
        [
            ["分数", "生图适配（始终）", "左右一致性（有左右差）", "空间关系（有空间关系）"],
            ["1.0", "几乎可直接作为生图 prompt，可见事实充分可成像", "主干左右差有设计逻辑，且已收束", "内外、上下、附着点可还原"],
            ["0.75", "轻微整理即可用；可见事实为主，最多 1 句氛围收尾", "主干左右差已对齐，仅轻度含糊", "主层次清楚，个别附着点略含糊"],
            ["0.5", "可见事实与评论 / 姿态各占一半", "仅配饰轻度左右差，主干衣裤鞋已统一", "层次大体清楚，但固定点含糊"],
            ["0.25", "抽象评论或姿态指令为主，可见事实偏少", "多处主干左右对撞，未合并也未删除", "层次或束系含糊，难以成像"],
            ["0.0", "无法从文本稳定生成可见造型", "左右互斥无法收束，无法稳定生成", "空间关系混乱，无法还原穿戴"],
        ],
        [0.9, 3.1, 3.0, 3.0],
        font_size=11,
        row_h=400000,
    )
    tint_metric_rows(tbl, [key] * 5)
    tf2 = add_textbox(slide, LEFT, Emu(4900000), BODY_W, Emu(900000))
    add_sections(tf2, [
        section_parts(
            "硬封顶",
            "姿态 / 走位占显著篇幅 → 生图适配最高 0.5；可见事实稀疏且抽象评论为主 → 最高 0.25。",
        ),
    ], size=14, gap=False)


def populate_binding_metrics(slide):
    populate_score_band_slide(
        slide,
        key="BindingAccuracy",
        title="Binding Metrics",
        intro="五档 1.0 / 0.75 / 0.5 / 0.25 / 0.0。两指标共用。实体绑定始终计；多单品绑定只在多件衣服时计，适用的取平均。",
        bands=[
            ("1.0", "所有属性归属清晰准确，实体、部位、层次和左右关系稳定无歧义"),
            ("0.75", "基本准确，仅轻微模糊或少量跨句依赖，不误导生成"),
            ("0.5", "存在不确定归属或局部串线风险，但不影响主体理解"),
            ("0.25", "多处属性可能错绑，需反复解析才能判断归属"),
            ("0.0", "核心属性严重张冠李戴，实体归属混乱到无法稳定生成"),
        ],
        footnote="属性绑定模块分 < 0.5 时，整条质量有效分 Q 最高只能到 0.6。绑错了，不允许靠设计价值把质量轴拉很高。",
        footnote_label="上限",
    )


def populate_language_metrics(slide):
    populate_score_band_slide(
        slide,
        key="LanguageClarity",
        title="Language Metrics",
        intro="五档 1.0 / 0.75 / 0.5 / 0.25 / 0.0。两指标共用。指代清晰度始终计；数量准确性有数量词时才计，适用的取平均。",
        bands=[
            ("1.0", "数量、侧别、指代均清晰，无需回头解析主语或对象"),
            ("0.75", "整体清楚，仅少量轻微模糊或局部承接略紧，不影响理解"),
            ("0.5", "大意可读，但数量 / 侧别 / 指代仍有明显阅读负担"),
            ("0.25", "数量或指代需反复确认，可读性明显受损"),
            ("0.0", "数量、侧别、指代严重混乱，难以稳定解析"),
        ],
    )


def populate_structure_metrics(slide):
    key = "StructuralClarity"
    add_title(slide, "Structure Metrics")
    add_label(slide, "打分量规")
    add_accent_line(slide, MODULE_BY_KEY[key]["color"])
    add_module_chips(slide, active=[key])
    tf = add_textbox(slide, LEFT, BODY_TOP, BODY_W, Emu(620000))
    add_sections(tf, [
        section_parts(
            "怎么打分",
            "五档 1.0 / 0.75 / 0.5 / 0.25 / 0.0。信息顺序始终计；层级单品聚合只在多件衣服时计，适用的取平均。",
        ),
    ], size=14, gap=False)
    tbl = add_table(
        slide,
        LEFT,
        Emu(2280000),
        BODY_W,
        [
            ["分数", "信息顺序（始终）", "层级单品聚合（多单品时）"],
            ["1.0", "主体 → 廓形 → 材质颜色 → 配件，主干先立", "按件或按层分述，无跨件穿插"],
            ["0.75", "主线清楚，最多 1 处局部逆序或压缩", "整体按件可分，仅少量属性交叉"],
            ["0.5", "材质 / 配件与主体穿插，尚可整理", "2–3 处跨件穿插，需来回对照"],
            ["0.25", "配件 / 评论先于主干，或层次滞后", "频繁交错，层次不清或滞后"],
            ["0.0", "无清晰主干递进，无法还原结构", "无法判断属性属于哪一件 / 哪一层"],
        ],
        [1.1, 4.4, 4.5],
        font_size=12,
        row_h=400000,
    )
    tint_metric_rows(tbl, [key] * 5)
    tf2 = add_textbox(slide, LEFT, Emu(5000000), BODY_W, Emu(900000))
    add_sections(tf2, [
        section_parts(
            "不评什么",
            "不评段落数、分节标题或列表排版。只评信息层级与归属。",
        ),
    ], size=14, gap=False)


def find_slide_index(prs, pred):
    for i, slide in enumerate(prs.slides):
        if pred(slide_heading(slide).lower()):
            return i
    return None


def clear_slide(slide):
    for shape in list(slide.shapes):
        shape._element.getparent().remove(shape._element)


def rebuild_slide(slide, populate_fn):
    clear_slide(slide)
    add_footer(slide)
    populate_fn(slide)


def insert_remaining_rubrics(prs):
    """Insert 打分量规 slides for the five modules that only had concept pages."""
    if find_slide_index(prs, lambda t: t == "visibility metrics") is not None:
        return False

    vis_idx = find_slide_index(
        prs,
        lambda t: "visibility" in t and "generation" in t and "metrics" not in t,
    )
    if vis_idx is None:
        vis_idx = find_slide_index(prs, lambda t: "visibility" in t and "metrics" not in t)
    if vis_idx is not None:
        s = add_slide_at(prs, vis_idx + 1)
        add_footer(s)
        populate_visibility_metrics(s)
        s = add_slide_at(prs, vis_idx + 2)
        add_footer(s)
        populate_generation_metrics(s)

    bind_idx = find_slide_index(
        prs,
        lambda t: "binding" in t and "language" in t and "metrics" not in t,
    )
    if bind_idx is not None:
        s = add_slide_at(prs, bind_idx + 1)
        add_footer(s)
        populate_binding_metrics(s)
        s = add_slide_at(prs, bind_idx + 2)
        add_footer(s)
        populate_language_metrics(s)
        s = add_slide_at(prs, bind_idx + 3)
        add_footer(s)
        populate_structure_metrics(s)
    return True


def upgrade_existing_rubrics(prs):
    """Split the combined BLS metrics page and expand all rubrics to five bands."""
    gen_idx = find_slide_index(prs, lambda t: t == "generation metrics")
    if gen_idx is not None:
        rebuild_slide(prs.slides[gen_idx], populate_generation_metrics)

    combined = find_slide_index(
        prs,
        lambda t: "binding" in t and "language" in t and "metrics" in t,
    )
    if combined is not None:
        rebuild_slide(prs.slides[combined], populate_binding_metrics)
        s = add_slide_at(prs, combined + 1)
        add_footer(s)
        populate_language_metrics(s)
        s = add_slide_at(prs, combined + 2)
        add_footer(s)
        populate_structure_metrics(s)
        return True

    bind_idx = find_slide_index(prs, lambda t: t == "binding metrics")
    if bind_idx is None:
        return False
    rebuild_slide(prs.slides[bind_idx], populate_binding_metrics)
    lang_idx = find_slide_index(prs, lambda t: t == "language metrics")
    struct_idx = find_slide_index(prs, lambda t: t == "structure metrics")
    if lang_idx is not None:
        rebuild_slide(prs.slides[lang_idx], populate_language_metrics)
    if struct_idx is not None:
        rebuild_slide(prs.slides[struct_idx], populate_structure_metrics)
    return True


def build():
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H

    # 1. Cover
    s = new_slide(prs)
    add_title(s, "Quality Scoring")
    add_label(s, "规范")
    add_accent_line(s)
    tf = add_textbox(s, LEFT, BODY_TOP, BODY_W, BODY_H)
    add_sections(tf, [
        section_parts("版本", "fashion_text_prompt_optimizer  v0.4.5"),
        section_parts("讲什么", "质量轴怎么设计，分数怎么从指标汇总到最终 S_fp。"),
        section_parts("评什么", "质量轴问「好不好」：写得具体、准确、清晰，有识别性设计想法，并能直接拿去生图。"),
        section_parts("不评什么", "不评规格有没有写全，不评系列共用成衣语法，惩罚分也不并入总分。"),
    ], size=17)

    # 2. Score axes
    s = new_slide(prs)
    add_title(s, "Score Axes")
    add_label(s, "总览")
    add_accent_line(s)
    tf = add_textbox(s, LEFT, BODY_TOP, BODY_W, BODY_H)
    add_sections(tf, [
        section_parts("覆盖轴  20%", "问有没有。该写的点按 0 或 1 命中，再取平均，得到 C。"),
        section_parts("质量轴  80%", "问好不好。六个子模块先各自均分，再按权重合成 Q。这是主轴。"),
        section_parts("加分项  不计分", "风格、品牌、叙事增强。会记下来，但不进入最终分 S_fp。"),
        section_parts("惩罚项  不扣分", "冗余、左右互斥、气质打架、不合理、套公式。不从总分里扣，只用来卡门。"),
    ], size=16)

    # 3. Quality modules + weights
    s = new_slide(prs)
    add_title(s, "Quality Modules")
    add_label(s, "模块设计")
    add_accent_line(s)
    tf = add_textbox(s, LEFT, BODY_TOP, BODY_W, Emu(900000))
    add_sections(tf, [
        section_parts(
            "怎么加权",
            "质量轴先按模块打分，再按权重合成。权重大，说明这件事更决定「好不好」。",
        ),
    ], size=16, gap=False)

    share = lambda w: f"{w / W_TOTAL * 100:.1f}%"
    tbl = add_table(
        s,
        LEFT,
        Emu(2550000),
        BODY_W,
        [
            ["模块", "权重", "占质量轴", "评什么"],
            ["设计价值  DesignMerit", "×10.0", share(10), "有没有不可替换的识别性设计想法"],
            ["可见性优先级  ConcisenessAndDensity", "×3.0", share(3), "可见主体是否压过姿态、评论、隐藏细节"],
            ["生成适配  GenerationReadiness", "×3.0", share(3), "文本能否直接支撑 text-to-image"],
            ["属性绑定  BindingAccuracy", "×1.0", share(1), "颜色、材质、配件有没有绑到正确对象"],
            ["语言清晰  LanguageClarity", "×1.0", share(1), "数量、代词、省略能不能一次读懂"],
            ["结构清晰  StructuralClarity", "×1.0", share(1), "信息是否按主体到细节递进、按件聚合"],
        ],
        [3.4, 0.9, 1.1, 4.6],
        font_size=12,
        row_h=400000,
    )
    color_module_table(tbl)
    add_table_module_bars(s, tbl)
    add_module_chips(s, active="all")

    tf2 = add_textbox(s, LEFT, Emu(5680000), BODY_W, Emu(700000))
    add_sections(tf2, [
        section_parts(
            "怎么读",
            "设计价值单独约占质量轴 53%。可见性和生成适配合起来约占 32%。绑定、语言、结构各约 5%，用来兜底可读性。",
        ),
    ], size=14, gap=False)

    # 4. DesignMerit philosophy
    s = new_slide(prs)
    add_title(s, "Design Merit")
    add_label(s, "×10.0")
    add_accent_line(s, MODULE_BY_KEY["DesignMerit"]["color"])
    add_module_chips(s, active=["DesignMerit"])
    tf = add_textbox(s, LEFT, BODY_TOP, BODY_W, BODY_H)
    add_sections(tf, [
        section_parts("评什么", "有没有不可替换的可见设计想法。写细、写全、露腰、系腰带，都不等于独特。"),
        section_parts("怎么检验", "把颜色、材质、品牌词换掉，造型身份是否还在。换完仍可读成同一套，就不是高分。"),
        section_parts(
            "什么算高分",
            "整面表面场（肌理 / 印花 / 密铺装饰）；画出轮廓的沿边路径；开合后内层仍是另一件可读的衣服；两件主干在表面或量感上对撞。",
        ),
        section_parts(
            "什么算低分",
            "成衣收口、常规叠穿、系列共用的短夹克 + 衬衫 + 短裤/裤 + 腰带语法、面料情绪词、主题对仗。",
        ),
        section_parts(
            "打完分后",
            "裁判打完后，若文本检测不到识别性想法，相关指标会被压到 0.25–0.5，避免把「写细」当成「写好」。",
        ),
    ], size=15)

    # 5. DesignMerit metrics
    s = new_slide(prs)
    add_title(s, "Design Merit Metrics")
    add_label(s, "打分量规")
    add_accent_line(s, MODULE_BY_KEY["DesignMerit"]["color"])
    add_module_chips(s, active=["DesignMerit"])
    tf = add_textbox(s, LEFT, BODY_TOP, BODY_W, Emu(700000))
    add_sections(tf, [
        section_parts(
            "怎么打分",
            "五档 0 / 0.25 / 0.5 / 0.75 / 1.0。模块分 = 适用指标取平均。工艺装饰只在工艺明显时才计分。",
        ),
    ], size=15, gap=False)
    add_table(
        s,
        LEFT,
        Emu(2380000),
        BODY_W,
        [
            ["指标", "1.0  有识别性想法", "0.25  没有识别性想法"],
            ["设计独特性", "整面表面场、沿边路径、开合后可读内层、或主干对撞", "成衣收口、常规叠穿、系列共用廓形、面料情绪或主题对仗"],
            ["视觉观察锚定", "想法落在部位、沿边路径或开合层次上", "部位写得很满，锚定的只是收口、露腰、五金或主题句"],
            ["工艺装饰显著度", "装饰性沿边、整面肌理/印花场、图形裁片画出轮廓", "只有明线、暗襟、滚边、品牌五金等成衣收口"],
            ["组合原创性", "开合后内层有自己的品类或表面身份；或主干对撞", "常规叠穿，换一件主干读感不变"],
            ["设计信号纯度", "几乎全是在写那个识别性想法", "可成像，但信号是规格清单、主题对仗或姿态指令"],
        ],
        [1.7, 3.7, 3.7],
        font_size=11,
        row_h=420000,
    )

    # 6. Visibility & Generation
    s = new_slide(prs)
    add_title(s, "Visibility & Generation")
    add_label(s, "各 ×3.0")
    add_accent_line(s, MODULE_BY_KEY["ConcisenessAndDensity"]["color"])
    add_module_chips(s, active=["ConcisenessAndDensity", "GenerationReadiness"])
    tf = add_textbox(s, LEFT, BODY_TOP, BODY_W, BODY_H)
    add_sections(tf, [
        section_parts(
            "可见性优先级",
            "看可见且决定成像的主体（廓形、品类、层次、主色、工艺路径、可见配件）是否压过隐藏细节、模特姿态和抽象评论。",
        ),
        [
            ("1.0  ", {"bold": True}),
            ("可见主体始终居前且占主导。  ", {"name": FONT}),
            ("0.5  ", {"bold": True}),
            ("氛围 / 姿态与主体争抢。  ", {"name": FONT}),
            ("0.25  ", {"bold": True}),
            ("非成像内容喧宾夺主。标准 T2I 开场白不计罚。", {"name": FONT}),
        ],
        section_parts(
            "生成适配",
            "看文本能否直接支撑生图：可见廓形、品类、材质、颜色、工艺位置、层次与配件是否足够具体、可成像。",
        ),
        [
            ("1.0  ", {"bold": True}),
            ("几乎可直接作为生图 prompt。  ", {"name": FONT}),
            ("0.5  ", {"bold": True}),
            ("可见事实与评论/姿态各占一半。  ", {"name": FONT}),
            ("0.25  ", {"bold": True}),
            ("抽象评论或姿态指令为主。", {"name": FONT}),
        ],
        section_parts(
            "硬封顶",
            "姿态/走位占显著篇幅 → 可见性与生成适配最高 0.5；抽象评论明显多于可见事实 → 最高 0.25。",
        ),
        section_parts(
            "生成适配还看",
            "左右一致性：主干衣裤鞋的左右互斥要收束。空间关系：内外、上下、附着点要能还原。",
        ),
    ], size=15)

    s = new_slide(prs)
    populate_visibility_metrics(s)

    s = new_slide(prs)
    populate_generation_metrics(s)

    # 7. Binding / language / structure
    s = new_slide(prs)
    add_title(s, "Binding, Language, Structure")
    add_label(s, "各 ×1.0")
    add_accent_line(s, MODULE_BY_KEY["BindingAccuracy"]["color"])
    add_module_chips(s, active=["BindingAccuracy", "LanguageClarity", "StructuralClarity"])
    tf = add_textbox(s, LEFT, BODY_TOP, BODY_W, BODY_H)
    add_sections(tf, [
        section_parts(
            "属性绑定",
            "颜色、材质、配件、结构、左右信息有没有绑到正确对象。多件衣服时不能张冠李戴。",
        ),
        section_parts(
            "语言清晰",
            "数量词准不准、有没有打架；代词和省略能不能一次读懂，不用回头找主语。",
        ),
        section_parts(
            "结构清晰",
            "按「主体 → 廓形 → 材质颜色 → 配件」往下写。多件衣服时同一件的属性写在一起，层次先交代再分述。不评段落数或列表排版。",
        ),
        section_parts(
            "上限",
            "属性绑定模块分 < 0.5 时，整条质量有效分 Q 最高只能到 0.6。绑错了，不允许靠设计价值把质量轴拉很高。",
        ),
    ], size=16)

    s = new_slide(prs)
    populate_binding_metrics(s)

    s = new_slide(prs)
    populate_language_metrics(s)

    s = new_slide(prs)
    populate_structure_metrics(s)

    # 8. Pipeline
    s = new_slide(prs)
    add_title(s, "Score Pipeline")
    add_label(s, "计算步骤")
    add_accent_line(s)
    tf = add_textbox(s, LEFT, BODY_TOP, BODY_W, BODY_H)
    add_sections(tf, [
        section_parts("1  预处理", "去掉生图固定句、章节标题等，得到正文。字数和裁判都按这段正文来。"),
        section_parts("2  裁判打分", "覆盖按 0 或 1，质量五档 0~1，惩罚五档 0~1。用不上的项跳过。"),
        section_parts("3  模块均分", "每个质量子模块 = 适用指标取平均。"),
        section_parts("4  按权重合成", "六个模块按 ×10 / ×3 / ×3 / ×1 / ×1 / ×1 加权，得到 Q_w。"),
        section_parts("5  质量有效分", "Q = min(Q_w, cap_q)。默认上限 1.0；绑定 < 0.5 时上限改成 0.6。惩罚不参与。"),
        section_parts("6  内容主分", "s_fp_base = 0.2·C + 0.8·Q。"),
        section_parts("7  长度校正", "相对典型长度微调一下，得到最终 S_fp。"),
        section_parts("8  双门限", "S_fp ≥ 0.7 且惩罚平均 ≤ 0.5。惩罚只挡门，不改分数。"),
    ], size=14)

    # 9. Formulas
    s = new_slide(prs)
    add_title(s, "Score Formulas")
    add_label(s, "计算公式")
    add_accent_line(s)
    tf = add_textbox(s, LEFT, BODY_TOP, BODY_W, BODY_H)
    add_sections(tf, [
        section_parts("覆盖分", "C = 适用覆盖指标取平均"),
        [
            ("[ 质量分 ]  ", {"name": FONT, "bold": True}),
            ("Q_w = Σ(w_i · m_i) / Σ(w_i)", {"name": FONT_MONO, "bold": False}),
        ],
        [
            ("              ", {"name": FONT_MONO}),
            ("Q = min(Q_w, cap_q)", {"name": FONT_MONO}),
        ],
        [
            ("[ 内容主分 ]  ", {"name": FONT, "bold": True}),
            ("s_fp_base = 0.2 · C + 0.8 · Q", {"name": FONT_MONO}),
        ],
        [
            ("[ 长度校正 ]  ", {"name": FONT, "bold": True}),
            ("x = ln(1 + L)     Δ_len = a · (x − x_0)", {"name": FONT_MONO}),
        ],
        [
            ("[ 最终得分 ]  ", {"name": FONT, "bold": True}),
            ("S_fp = clip(s_fp_base − Δ_len, 0, 1)", {"name": FONT_MONO}),
        ],
        section_parts(
            "当前参数",
            f"a = {A}，x_0 = {X0}，典型长度约 {L_REF} 字。写到这个长度时不校正，S_fp 就等于内容主分。",
        ),
        section_parts(
            "一行总式",
            "S_fp = clip( 0.2·C + 0.8·Q  −  a·(ln(1+L) − x_0) , 0, 1 )",
        ),
    ], size=16)

    # 10. Length & gates
    s = new_slide(prs)
    add_title(s, "Length & Gates")
    add_label(s, "门限与校正")
    add_accent_line(s)
    tf = add_textbox(s, LEFT, BODY_TOP, BODY_W, BODY_H)
    add_sections(tf, [
        section_parts(
            "长度怎么调",
            "以常见长度 x_0 为锚，只校正相对偏长 / 偏短。写得比常见更长就扣一点，更短就略加一点，避免靠写长刷分。",
        ),
        section_parts(
            "得分门限",
            "S_fp ≥ 0.7。过了这道门，才算质量达标。",
        ),
        section_parts(
            "惩罚门限",
            "五项惩罚取平均，平均分 ≤ 0.5。当前惩罚不乘进总分，也不进奖励，只用来卡门。",
        ),
        section_parts(
            "奖励分",
            "单条评分时，奖励分就等于 S_fp。同一组里多条一起比时，才按字数微调一点，避免字数差主导奖励。",
        ),
        section_parts(
            "记住",
            "覆盖问有没有，质量问好不好。最终分数 80% 来自质量轴，质量轴一半以上来自设计价值。",
        ),
    ], size=16)

    desktop = Path(r"c:\Users\lsh\Desktop\quality_score_spec.pptx")
    local = Path(r"e:\fashion_agents_project\汇报\质量轴与计分.pptx")
    save_presentation(prs, desktop)
    save_presentation(prs, local)


def patch_named_desktop_pptx():
    desktop = Path(r"C:\Users\lsh\Desktop")
    candidates = [
        f for f in desktop.glob("*.pptx")
        if "(2)" in f.name and not f.name.startswith("~")
    ]
    if not candidates:
        print("No 改进方案及训练(2).pptx found on Desktop")
        return
    # Prefer the already-colored copy; fall back to the original (2) file.
    colored = [f for f in candidates if "模块配色" in f.name]
    source = colored[0] if colored else candidates[0]
    prs = Presentation(str(source))
    colorize_existing_presentation(prs)
    inserted = insert_remaining_rubrics(prs)
    upgraded = upgrade_existing_rubrics(prs)
    print(f"Source: {source}  rubrics_inserted={inserted}  rubrics_upgraded={upgraded}")
    save_presentation(prs, source)
    original = next((f for f in candidates if "模块配色" not in f.name), None)
    if original and original != source:
        save_presentation(prs, original)


if __name__ == "__main__":
    import sys
    if "patch-only" in sys.argv:
        patch_named_desktop_pptx()
    else:
        build()
        patch_named_desktop_pptx()
