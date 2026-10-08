"""三页横向 PPT：全部 / 逆解析 / 生成图，组内最高改写相对原文。"""

from __future__ import annotations

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

OUT = r"c:\Users\lsh\Desktop\记录10.1\改写器效果对比.pptx"

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)
BLACK = RGBColor(0x1A, 0x1A, 0x1A)
GRAY = RGBColor(0x5C, 0x5C, 0x5C)
MUTED = RGBColor(0x8A, 0x8A, 0x8A)
RED = RGBColor(0xC0, 0x22, 0x22)
ORIG = RGBColor(0x9A, 0x9A, 0x9A)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
FONT = "微软雅黑"


def _run(paragraph, text, *, size=14, bold=False, color=BLACK):
    run = paragraph.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = FONT
    return run


def _box(slide, l, t, w, h, lines, *, align=PP_ALIGN.LEFT):
    shape = slide.shapes.add_textbox(l, t, w, h)
    tf = shape.text_frame
    tf.word_wrap = True
    for i, parts in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(2)
        for text, kwargs in parts:
            _run(p, text, **kwargs)
    return shape


def _solid(series, color: RGBColor) -> None:
    fill = series.format.fill
    fill.solid()
    fill.fore_color.rgb = color
    series.format.line.fill.background()


def _chart(slide, l, t, w, h, title, categories, orig, rewrite, *, percent=False):
    data = CategoryChartData()
    data.categories = categories
    data.add_series("原文", orig)
    data.add_series("组内最高改写", rewrite)
    graphic = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, l, t, w, h, data)
    chart = graphic.chart
    chart.has_title = True
    chart.has_legend = True
    chart.legend.position = XL_LEGEND_POSITION.BOTTOM
    chart.legend.include_in_layout = False
    chart.legend.font.size = Pt(11)
    chart.legend.font.name = FONT
    title_p = chart.chart_title.text_frame.paragraphs[0]
    title_p.clear()
    run = title_p.add_run()
    run.text = title
    run.font.size = Pt(14)
    run.font.bold = True
    run.font.color.rgb = BLACK
    run.font.name = FONT
    plot = chart.plots[0]
    plot.gap_width = 70
    plot.has_data_labels = True
    _solid(plot.series[0], ORIG)
    _solid(plot.series[1], RED)
    for series in plot.series:
        labels = series.data_labels
        labels.font.size = Pt(10)
        labels.font.bold = True
        labels.font.name = FONT
        labels.font.color.rgb = BLACK
        labels.number_format = "0%" if percent else "0.000"
        # percent series are already 0-100 numbers, not fractions
        if percent:
            labels.number_format = "0.0"
    axis = chart.value_axis
    axis.minimum_scale = 0
    axis.maximum_scale = 100 if percent else 1
    axis.has_major_gridlines = True
    axis.major_gridlines.format.line.color.rgb = RGBColor(0xE6, 0xE6, 0xE6)
    axis.tick_labels.font.size = Pt(10)
    axis.tick_labels.font.name = FONT
    axis.tick_labels.font.color.rgb = GRAY
    if percent:
        axis.tick_labels.number_format = "0"
    else:
        axis.tick_labels.number_format = "0.0"
    chart.category_axis.tick_labels.font.size = Pt(11)
    chart.category_axis.tick_labels.font.name = FONT
    chart.category_axis.tick_labels.font.color.rgb = BLACK
    # hide the extra "0%" misread: values are percents already
    return chart


def add_slide(prs, page):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.background.fill
    bg.solid()
    bg.fore_color.rgb = WHITE

    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), SLIDE_W, Inches(0.08))
    bar.fill.solid()
    bar.fill.fore_color.rgb = RED
    bar.line.fill.background()

    _box(
        slide,
        Inches(0.45),
        Inches(0.22),
        Inches(10.2),
        Inches(0.42),
        [[(page["kicker"], {"size": 12, "color": GRAY})]],
    )
    _box(
        slide,
        Inches(0.45),
        Inches(0.52),
        Inches(12.4),
        Inches(0.5),
        [[(page["title"], {"size": 26, "bold": True})]],
    )
    _box(
        slide,
        Inches(0.45),
        Inches(1.12),
        Inches(12.4),
        Inches(0.85),
        page["red_lines"],
    )

    _chart(
        slide,
        Inches(0.35),
        Inches(2.05),
        Inches(6.35),
        Inches(4.55),
        "分数：原文 vs 组内最高改写",
        ["总分", "质量轴", "覆盖"],
        page["score_orig"],
        page["score_best"],
        percent=False,
    )
    _chart(
        slide,
        Inches(6.75),
        Inches(2.05),
        Inches(6.2),
        Inches(4.55),
        "门通过率（%）",
        ["得分门", "惩罚门", "双门"],
        page["gate_orig"],
        page["gate_best"],
        percent=True,
    )
    _box(
        slide,
        Inches(0.45),
        Inches(6.85),
        Inches(12.4),
        Inches(0.4),
        [[(page["foot"], {"size": 11, "color": MUTED})]],
    )


PAGES = [
    {
        "kicker": "01  /  全部 400 组",
        "title": "改写器把组内最高分整体抬过原文",
        "red_lines": [
            [
                ("总分 0.714 → 0.898，提升 0.184。", {"size": 16, "bold": True, "color": RED}),
                ("  更高 297 组，更低 33 组，持平 70 组。", {"size": 16, "color": BLACK}),
            ],
            [
                ("得分门 60.8% → 88.3%（+27.5 个百分点），惩罚门 72.8% → 99.8%（+27.0 个百分点）。", {"size": 16, "bold": True, "color": RED}),
            ],
        ],
        "score_orig": (0.714, 0.761, 0.914),
        "score_best": (0.898, 0.885, 0.973),
        "gate_orig": (60.8, 72.8, 60.8),
        "gate_best": (88.3, 99.8, 88.3),
        "foot": "每组取改写总分最高的一条，对比同组原文。双门与得分门重合：最高分几乎都过了惩罚门。红色柱为改写器。",
    },
    {
        "kicker": "02  /  逆解析 200 组",
        "title": "原文已经较高，改写器仍小幅抬分并收紧门限",
        "red_lines": [
            [
                ("总分 0.884 → 0.928，提升 0.044。", {"size": 16, "bold": True, "color": RED}),
                ("  更高 121 组，更低 27 组，持平 52 组。", {"size": 16, "color": BLACK}),
            ],
            [
                ("得分门 90.0% → 98.5%（+8.5 个百分点），惩罚门 96.0% → 100%。", {"size": 16, "bold": True, "color": RED}),
            ],
        ],
        "score_orig": (0.884, 0.885, 0.928),
        "score_best": (0.928, 0.917, 0.979),
        "gate_orig": (90.0, 96.0, 90.0),
        "gate_best": (98.5, 100.0, 98.5),
        "foot": "逆解析原文中位数已是 0.918，组内最高改写的中位差为 +0.012。属性绑定和语言清晰两边都接近满分。",
    },
    {
        "kicker": "03  /  生成图 200 组",
        "title": "改写器主要拉高的是生成图",
        "red_lines": [
            [
                ("总分 0.543 → 0.868，提升 0.325。", {"size": 16, "bold": True, "color": RED}),
                ("  更高 176 组，更低 6 组，持平 18 组。", {"size": 16, "color": BLACK}),
            ],
            [
                ("得分门 31.5% → 78.0%（+46.5 个百分点），惩罚门 49.5% → 99.5%（+50.0 个百分点）。", {"size": 16, "bold": True, "color": RED}),
            ],
        ],
        "score_orig": (0.543, 0.637, 0.900),
        "score_best": (0.868, 0.852, 0.967),
        "gate_orig": (31.5, 49.5, 31.5),
        "gate_best": (78.0, 99.5, 78.0),
        "foot": "生成适配 +0.393，信息密度 +0.359。原文总惩罚 0.207，组内最高改写降到 0.018。",
    },
]


def main() -> None:
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    prs.core_properties.title = "改写器效果对比"
    for page in PAGES:
        add_slide(prs, page)
    prs.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
