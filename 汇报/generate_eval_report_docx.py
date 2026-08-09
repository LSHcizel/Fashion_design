"""Generate consolidated evaluation spec Word document from three axes."""

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml.ns import qn
from docx.shared import Pt, Cm
from pathlib import Path


def set_cell_text(cell, text, bold=False, size=9):
    cell.text = ""
    p = cell.paragraphs[0]
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")


def add_table(doc, headers, rows, col_widths=None):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    for i, h in enumerate(headers):
        set_cell_text(table.rows[0].cells[i], h, bold=True)
    for r_idx, row in enumerate(rows):
        for c_idx, val in enumerate(row):
            set_cell_text(table.rows[r_idx + 1].cells[c_idx], val)
    if col_widths:
        for row in table.rows:
            for i, w in enumerate(col_widths):
                row.cells[i].width = Cm(w)
    doc.add_paragraph()
    return table


def add_page_break(doc):
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def add_heading(doc, text, level=1, page_break_before=False):
    if page_break_before:
        add_page_break(doc)
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        run.font.name = "Microsoft YaHei"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    return h


def add_para(doc, text, bold=False):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = bold
    run.font.name = "Microsoft YaHei"
    run.font.size = Pt(10.5)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(item, style="List Bullet")
        for run in p.runs:
            run.font.name = "Microsoft YaHei"
            run.font.size = Pt(10.5)
            run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")


def add_formula(doc, text):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.name = "Consolas"
    run.font.size = Pt(10.5)
    p.paragraph_format.left_indent = Cm(0.5)
    return p


S_FP_FORMULA_SECTIONS = [
    (
        "第 1 行：把字符数换成对数长度",
        ["x = ln(1 + L)"],
        "对字符数取 ln(1+L)，使篇幅越长、每多写几百字对分数的边际影响越小，避免用 raw 字数线性比较。",
    ),
    (
        "第 2 行：确定基准长度（典型长度）",
        ["x_0 = log_len_center", "L_ref = exp(x_0) - 1"],
        "x_0 是留出集上的典型对数长度，当 L = L_ref 时长度校正量为 0，S_fp 等于 s_fp_base。",
    ),
    (
        "第 3 行：计算长度校正量",
        ["Δ_len = a · (x - x_0)"],
        "a为斜率：比基准更长则 Δ_len > 0 待扣分，更短则 Δ_len < 0 略加分，接近基准则 Δ_len ≈ 0 不校正，整体只做微调。",
    ),
    (
        "第 4 行：从内容主分中减去校正量",
        ["S_fp = clip(s_fp_base - Δ_len, 0, 1)"],
        "从 s_fp_base 减去 Δ_len 得到最终 S_fp，并截断到 0~1。",
    ),
]


def add_s_fp_section(doc):
    add_heading(doc, "四、S_fp 长度去相关", 1, page_break_before=True)
    for title, formulas, explanation in S_FP_FORMULA_SECTIONS:
        add_para(doc, title, bold=True)
        for formula in formulas:
            add_formula(doc, formula)
        add_para(doc, explanation)


def build_doc():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)

    title = doc.add_heading("服装设计文本评估体系", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in title.runs:
        run.font.name = "Microsoft YaHei"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")

    add_para(doc, "本文档由惩罚分、覆盖轴、质量轴、S_fp 长度去相关四部分组成。")

    # ── 一、惩罚分 ──
    add_heading(doc, "一、惩罚分", 1, page_break_before=True)
    add_para(doc, "作用：暂不直接从总分扣分，仅作惩罚门限判定。")
    add_bullets(doc, [
        "每项 0~1 分五档（0 最好，1 最差），仅取 0 / 0.25 / 0.5 / 0.75 / 1.0",
        "综合惩罚分 P̄ = 五项算术平均",
        "通过条件：P̄ ≤ 0.5",
    ])

    add_heading(doc, "惩罚项概览", 2)
    add_table(doc, ["惩罚项", "核心判定"], [
        ["非生成导向内容", "冗余重复、氛围评论、分析解说、风格符号堆砌"],
        ["一致性", "主干（衣裤鞋/里衬）左右或同件互斥"],
        ["协调性", "主干间造型气质/场域冲突"],
        ["合理性", "违反物理/穿着常识（材质、腿带绑具等）"],
        ["公式模板", "巡航度假公式、品牌符号堆叠、可互换造型"],
    ], [3.5, 12.5])

    penalty_items = [
        ("非生成导向内容", None, [
            ["0", "无显著冗余或评论句；生成导向清晰"],
            ["0.25", "少量重复或 1–2 处轻量概念/评论表述"],
            ["0.5", "多处重复或 runway 评论；或 ≥3 处氛围/姿态/沙龙/漫步/仿佛/暗示类句"],
            ["0.75", "叠加可互换度假公式、品牌符号堆叠，或冗长重复铺垫"],
            ["1.0", "大量互异风格符号堆叠，未先交代主干，难以还原可见造型"],
        ]),
        ("一致性", "主干：全部衣裤鞋及明确写出的里衬/内衬。仅配饰轻度左右差异且主干统一 → 从宽不加罚。", [
            ["0", "所有主干结构/身份统一，无互斥"],
            ["0.25", "已收束为单一主干语法，或仅一笔轻描不对称"],
            ["0.5", "任一件主干左右割裂；或 ≥2 类主干层级同时左右互斥"],
            ["0.75", "多区主干对撞；单件「半件正装半件另一语汇」；壳与里衬互斥"],
            ["1.0", "多条互斥主体身份并列，仅以「解构」一句带过"],
        ]),
        ("协调性", "违和仅来自小配饰且轻微、主干已统一 → 从宽不加罚。", [
            ["0", "外套、内搭、裤、鞋及里衬场域统一，层次衔接可读"],
            ["0.25", "整体协调，仅配饰或局部轻微气质差异"],
            ["0.5", "同件上装左右互斥场域；内外层语气打架；鞋与套装严重冲突；腿带绑具违和"],
            ["0.75", "多区主干场域对撞；同件主干上哑光与高光金属护臂等对撞且无收束"],
            ["1.0", "整体造型场域完全脱节，无法收束为单一 look"],
        ]),
        ("合理性", None, [
            ["0", "材质、结构、穿着关系符合现实常识"],
            ["0.25", "极轻微不合理，文本中有可解释空间"],
            ["0.5", "部分设定勉强成立；腿带/吊带与裤管关系交代不清"],
            ["0.75", "明显违反现实材质或物理条件；腿带+harness 与裤型难以同时成立"],
            ["1.0", "严重违背穿着常识，装束关系无法成立"],
        ]),
        ("公式模板", "总领性判断：看整体是否像可互换公式稿，不看段落格式。典型示例：短夹克+条纹衬衫+修身短裤/长裤+腰带+平底鞋。", [
            ["0~0.25", "紧凑、基于可见观察，无公式特征"],
            ["0.25~0.5", "明显命中 1 项：① 巡航度假公式 ② 品牌符号无工艺点 ③ 氛围句稀释设计 ④ 换材质词可套用"],
            ["0.5~0.75", "明显命中 2 项"],
            ["0.75~1.0", "命中 3 项及以上，或全面公式化"],
        ]),
    ]

    for name, note, rows in penalty_items:
        add_heading(doc, name, 2)
        if note:
            add_para(doc, note)
        add_table(doc, ["分数", "标准"], rows, [2, 14])

    # ── 二、覆盖轴 ──
    add_heading(doc, "二、覆盖轴", 1, page_break_before=True)
    add_bullets(doc, [
        "轴权重：0.2；问题：有没有",
        "每项命中 1 分、未命中 0 分，只对「适用」指标取平均",
        "不适用指标跳过，不计入模块平均",
    ])

    coverage_modules = [
        ("服装主体（5 项）", [
            ["服装品类", "始终", "明确给出主体品类"],
            ["廓形", "廓形显著", "给出整体轮廓或结构趋势"],
            ["长度/下摆", "长度显著", "给出长度或下摆信息"],
            ["肩部结构", "肩线显著", "点明肩部结构语言，不能仅写「上装」或笼统廓形"],
            ["裸露/包裹", "裸露显著", "描述显著裸露或包裹区域"],
        ]),
        ("材质颜色（6 项）", [
            ["材质类别", "材质可推断", "给出主体材质类别"],
            ["表面性质", "表面特征显著", "给出光泽、哑光、垂坠或硬挺等"],
            ["主色", "始终", "给出主色"],
            ["副色", "副色显著", "描述明显副色；内里/开衩内衬等高对比色块须说明可见条件"],
            ["配色逻辑", "多色/强对比", "交代配色逻辑；仅罗列颜色词不算命中"],
            ["图案", "有图案", "给出图案类型"],
        ]),
        ("结构细节（5 项）", [
            ["开合方式", "开合显著", "提到显著开合方式"],
            ["功能细节", "功能细节显著", "提到显著口袋、挂带或功能细节"],
            ["结构工艺", "工艺显著", "同时给出工艺类型与大致落位；仅泛称「有细节」不算"],
            ["解构", "有解构", "提到解构设计"],
            ["五金装饰", "装饰显著", "提到显著五金或装饰"],
        ]),
        ("造型配件（5 项）", [
            ["包袋", "全套造型且有包", "视觉突出时须覆盖至少 2 项：品类/持拿、轮廓/体量、颜色与材质"],
            ["鞋履", "全套造型且有鞋", "视觉突出时须覆盖至少 2 项：鞋型族、跟高/靴筒/鞋头、颜色与材质"],
            ["首饰", "首饰显著", "提到显著首饰或身体装饰"],
            ["腰带/腰胯附件", "腰带/束系显著", "须说明与外衣/裤装附着关系；剪裁收腰、省道等不算腰带"],
            ["叠搭层次", "有多层", "给出层次关系，遮盖与固定顺序可还原"],
        ]),
        ("造型关系（3 项）", [
            ["上下比例", "比例显著", "提到明显上下比例；长外套/阔腿/厚底鞋叠加时须交代剪影主次"],
            ["不对称", "有不对称", "提到不对称设计"],
            ["跨单品区分", "多单品", "多单品时能区分属性属于哪件单品"],
        ]),
    ]

    for mod_name, rows in coverage_modules:
        add_heading(doc, mod_name, 2)
        add_table(doc, ["指标", "触发条件", "命中标准"], rows, [3, 3.5, 9.5])

    add_para(doc, "模块得分 = 该模块适用指标的命中率（命中数 / 适用数）。")

    # ── 三、质量轴 ──
    add_heading(doc, "三、质量轴", 1, page_break_before=True)
    add_bullets(doc, [
        "轴权重：0.8；问题：好不好",
        "各指标五档计分（0 / 0.25 / 0.5 / 0.75 / 1.0），模块分 = 适用指标算术平均",
    ])

    add_heading(doc, "模块权重", 2)
    add_table(doc, ["模块", "权重"], [
        ["设计价值", "10.0"],
        ["可见性优先级", "5.0"],
        ["生成适配", "5.0"],
        ["属性绑定", "1.0"],
        ["语言清晰", "1.0"],
        ["结构清晰", "1.0"],
    ], [6, 3])

    design_merit_soft_rules = {
        "设计独特性": [
            "优先识别沿具体部位路径展开的工艺/装饰锚点（neckline、front edges、cuffs、hem、pocket edges）",
            "异材质碰撞、非常规比例、强识别图案/print、手工感 texture 写清为记忆点",
            "结构亮点（boxy/cropped/dropped waist/deep V/asymmetric panels）可与工艺并列计为锚点",
            "金属扣/吊坠/链饰写清位置与视觉作用可作辅助识别点",
            "参考逆解析 corpus（wgsn_batch 20260524T044337Z，n≈79）",
        ],
        "视觉观察锚定": [
            "正文按主干单品 → 内外层次 → 下装/鞋配 → styling 配件递进",
            "部位锚定：neckline/collar、shoulders、sleeves、waist、hem、center front、cuffs、pockets",
            "层次用 worn open over / beneath / underneath / partly obscured 等可见语言",
            "不确定处用 appears / could be / not confirmed，优于编造 unseen 细节",
            "段末 1 句 overall mood/palette 可接受；中段大量 mood/essay 仍低分",
        ],
        "工艺装饰显著度": [
            "须写清类型 + 路径/位置 + 视觉作用三要素",
            "偏好 appliqué、braid trim、fringe、piping、embellishment、woven pattern 等具体 craft 词",
            "craft 为 look 主 hook 时在前 2/3 展开",
            "texture 应连位置与作用一起写",
            "仅 refined finish / hardware 泛称低于参考水平",
        ],
        "组合原创性": [
            "常以外穿 + 内搭 + 下装 + 包鞋配四层组合描述完整 look",
            "高分组合用 structured vs sporty、voluminous outer vs flat graphic inner 等 contrast 语言",
            "layering 逻辑具体时，即使品类偏常规仍可中高分",
            "鞋包首饰若写入组合逻辑视为组合一部分",
            "按换一件主干单品是否改读感区分强弱",
        ],
        "设计信号纯度": [
            "正文以可成像设计事实为主",
            "Styling includes / carries 引导的配件鞋履计为设计事实",
            "段末 1 句 overall mood/palette 可接受",
            "不写品牌名、不编造 unseen back / hidden closure",
            "具体色名 + 作用优于抽象奢华/度假评论",
        ],
    }

    quality_modules = [
        ("设计价值（×10.0）", ["设计独特性", "视觉观察锚定", "工艺装饰显著度", "组合原创性", "设计信号纯度"], {
            "设计独特性": [
                ["1.0", "≥2 个清晰非套路锚点，至少 1 个来自工艺/结构/非常规组合"],
                ["0.75", "1 个明确非公式化记忆点"],
                ["0.5", "偏常规奢华/度假语言，识别点弱"],
                ["0.25", "几乎全是可互换巡航度假公式"],
                ["0.0", "纯品类堆叠或仅品牌符号"],
            ],
            "视觉观察锚定": [
                ["1.0", "多处具体部位锚定，可见/不可见层次清楚，几乎无氛围/评论句"],
                ["0.75", "主要结构有可见锚定，概念句极少（≤1 处）"],
                ["0.5", "观察与概念/氛围句混合"],
                ["0.25", "以主题/氛围/沙龙/漫步等修辞为主"],
                ["0.0", "纯概念叙事，无法对应可见造型"],
            ],
            "工艺装饰显著度": [
                ["1.0", "工艺/装饰的类型、位置、视觉作用均清晰，且为造型主识别点之一"],
                ["0.75", "主要工艺点清楚，少量泛化"],
                ["0.5", "提到工艺但偏泛（明线、五金、精致饰面等）"],
                ["0.25", "仅「装饰/五金」级泛称"],
                ["0.0", "应写工艺处完全缺失或空泛"],
            ],
            "组合原创性": [
                ["1.0", "组合高度具体且难预测，换一件主干单品会改变整体读感"],
                ["0.75", "组合有明确取向，但仍有部分常规元素"],
                ["0.5", "安全奢华组合可预测；或仍属度假公式"],
                ["0.25", "典型 formula 组合，可预测；或主干单品可互换而不改变整体读感"],
                ["0.0", "最常见巡航/工装模板拼接，多条造型间可互换主干"],
            ],
            "设计信号纯度": [
                ["1.0", "几乎全是可成像的设计事实，无氛围/姿态/身份评论"],
                ["0.75", "设计事实为主，氛围/姿态句 ≤1 句"],
                ["0.5", "设计与叙事各占一半；或重复氛围铺垫"],
                ["0.25", "≥3 处沙龙/漫步/航海/仿佛/暗示等修辞稀释设计信号"],
                ["0.0", "主要是评论文，设计事实被淹没"],
            ],
        }),
        ("可见性优先级（×5.0）", ["可见性优先级"], {
            "可见性优先级": [
                ["1.0", "廓形、品类、层次、主色、关键工艺与可见配件始终居前且占主导"],
                ["0.75", "可见主体明确优先；可有 1 处轻量氛围句"],
                ["0.5", "氛围/姿态/过渡铺垫占明显篇幅，或低可见细节与主体争抢"],
                ["0.25", "氛围/身份/品牌评论喧宾夺主；或 ≥3 处氛围类句稀释可见主体"],
                ["0.0", "几乎不以可见服装为主体"],
            ],
        }),
        ("生成适配（×5.0）", ["生图提示词适配", "左右一致性", "空间关系"], {
            "（三指标共用维度分）": [
                ["1.0", "几乎可直接作为生图提示词"],
                ["0.75", "轻微整理后即可直接使用"],
                ["0.5", "偏说明文或设计解说，不够「可直接生图」"],
                ["0.25", "更像分析文本或风格评论"],
                ["0.0", "无法直接用于生成"],
            ],
        }),
        ("属性绑定（×1.0）", ["属性实体绑定", "多单品绑定"], {
            "（两指标共用维度分）": [
                ["1.0", "所有属性归属清晰准确，实体、部位、层次和左右关系稳定"],
                ["0.75", "基本准确，仅轻微模糊"],
                ["0.5", "存在不确定归属或局部串线风险"],
                ["0.25", "多处属性可能错绑，需反复解析"],
                ["0.0", "核心属性严重张冠李戴"],
            ],
        }),
        ("语言清晰（×1.0）", ["数量准确性", "指代清晰度"], {
            "（两指标共用维度分）": [
                ["1.0", "数量、侧别、指代均清晰"],
                ["0.75", "整体清楚，仅少量轻微模糊"],
                ["0.5", "数量/侧别/指代仍有明显阅读负担"],
                ["0.25", "数量或指代需反复确认"],
                ["0.0", "数量、侧别、指代严重混乱"],
            ],
        }),
        ("结构清晰（×1.0）", ["信息顺序", "层级单品聚合"], {
            "信息顺序": [
                ["1.0", "主干→廓形结构→材质颜色→配件，四段递进自然"],
                ["0.75", "主线清楚，≤1 处局部逆序"],
                ["0.5", "材质/配件与主体穿插，但尚可整理"],
                ["0.25", "配件/氛围明显先于主干"],
                ["0.0", "无清晰主干递进"],
            ],
            "层级单品聚合（多单品时）": [
                ["1.0", "按单品或层次分块/连贯分述，无跨件穿插"],
                ["0.75", "整体按件可分，仅少量属性轻微交叉"],
                ["0.5", "2–3 处跨件穿插，读者需来回对照"],
                ["0.25", "多单品属性频繁交错"],
                ["0.0", "无法判断属性归属哪一件/哪一层"],
            ],
        }),
    ]

    for mod_title, indicators, score_tables in quality_modules:
        add_heading(doc, mod_title, 2)
        add_para(doc, "指标：" + "、".join(indicators))
        if mod_title.startswith("可见性优先级"):
            add_bullets(doc, [
                "硬封顶：氛围/姿态/身份评论占显著篇幅 → 最高 0.5",
                "不可见信息或重复过渡铺垫喧宾夺主 → 最高 0.25",
            ])
        for sub_name, rows in score_tables.items():
            if sub_name.startswith("（"):
                add_para(doc, sub_name)
            else:
                add_para(doc, sub_name, bold=True)
            add_table(doc, ["分数", "标准"], rows, [2, 14])
            soft = design_merit_soft_rules.get(sub_name.rstrip("（多单品时）").split("（")[0])
            if soft and mod_title.startswith("设计价值"):
                add_para(doc, "软性规则（参考逆解析 corpus）", bold=True)
                add_bullets(doc, soft)

    add_heading(doc, "上限规则与综合计分", 2)
    add_bullets(doc, [
        "属性绑定模块分 < 0.5 时，质量轴得分最高只能到 0.6",
        "Q_w = Σ(模块权重 × 模块分) / Σ(模块权重)",
        "Q = min(Q_w, cap_q)；默认 cap_q = 1.0，属性绑定 < 0.5 时 cap_q = 0.6",
        "内容主分：s_fp_base = 0.2 × C + 0.8 × Q（C 为覆盖轴得分，Q 为质量有效分）",
    ])

    add_s_fp_section(doc)

    return doc


def main():
    out_dir = Path(__file__).resolve().parent
    out = out_dir / "服装设计文本评估体系.docx"
    doc = build_doc()
    try:
        doc.save(out)
    except PermissionError:
        out = out_dir / "服装设计文本评估体系_新版.docx"
        doc.save(out)
    print(f"Saved: {out}")


if __name__ == "__main__":
    main()
