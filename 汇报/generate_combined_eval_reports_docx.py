"""Merge evaluation markdown reports into one Word document."""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
REPORT_DIR = Path(__file__).resolve().parent
OUTPUT = REPORT_DIR / "完整评分报告合集.docx"

REPORT_ASSETS = [
    {
        "report_md": REPORT_DIR / "010_ps27_031_完整评分报告.md",
        "source_text": ROOT
        / "fashion_research_dir/wgsn_batch_image_inverse/20260524T044337Z/"
        "10_01_图片库＆设计资源_中文_女士_Chanel_Pre-Summer_2027_服装__010_media_cha_biarritz_ps27_031_text_description.md",
        "source_image": ROOT
        / "downloads/wgsn_latest_batch/01_图片库＆设计资源_中文_女士_Chanel_Pre-Summer_2027_服装/"
        "010_media_cha_biarritz_ps27_031.jpg",
    },
    {
        "report_md": REPORT_DIR / "chapter01_look01_完整评分报告.md",
        "source_text": ROOT / "fashion_research_dir/workflow_0/2026-06-06/chapter_01/look_01.txt",
        "source_image": ROOT
        / "fashion_research_dir/workflow_0/2026-06-06/chapter_01/image2_generation/look_01_gpt-image-2.png",
    },
]


def set_run_font(run, *, bold=False, size=10.5, name="Microsoft YaHei"):
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)


def set_cell_text(cell, text, *, bold=False, size=9):
    cell.text = ""
    p = cell.paragraphs[0]
    run = p.add_run(strip_md(text))
    set_run_font(run, bold=bold, size=size)


def strip_md(text: str) -> str:
    text = re.sub(r"`([^`]+)`", r"\1", text)
    text = re.sub(r"\*\*([^*]+)\*\*", r"\1", text)
    return text.strip()


def add_page_break(doc):
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


HEADING_SIZES = {0: 18, 1: 16, 2: 14, 3: 12, 4: 11}


def add_heading(doc, text, level=1):
    h = doc.add_heading(strip_md(text), level=min(level, 4))
    for run in h.runs:
        set_run_font(run, bold=True, size=HEADING_SIZES.get(level, 11))
    pf = h.paragraph_format
    if level == 2:
        pf.space_before = Pt(14)
        pf.space_after = Pt(6)
    elif level == 3:
        pf.space_before = Pt(10)
        pf.space_after = Pt(4)
    elif level == 4:
        pf.space_before = Pt(8)
        pf.space_after = Pt(2)
        pf.left_indent = Cm(0.4)
    return h


def add_para(doc, text, *, bold=False):
    p = doc.add_paragraph()
    run = p.add_run(strip_md(text))
    set_run_font(run, bold=bold)
    return p


def is_field_legend(text: str) -> bool:
    return "｜" in text and " — " in text


def add_legend_para(doc, text):
    p = doc.add_paragraph()
    run = p.add_run(strip_md(text))
    set_run_font(run, size=9.5)
    run.italic = True
    run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
    p.paragraph_format.space_after = Pt(4)
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(strip_md(text), style="List Bullet")
    for run in p.runs:
        set_run_font(run)
    return p


def load_source_text(path: Path) -> str:
    raw = path.read_text(encoding="utf-8").strip()
    if path.suffix.lower() == ".md":
        match = re.search(r"^##\s*text_description\s*\n+([\s\S]*?)(?:\n##\s|\Z)", raw, re.MULTILINE)
        if match:
            return match.group(1).strip()
    return raw


def add_source_block(doc, source_text_path: Path, source_image_path: Path):
    add_para(doc, "被评文本", bold=True)
    content = load_source_text(source_text_path)
    for paragraph in content.split("\n\n"):
        paragraph = paragraph.strip()
        if not paragraph:
            continue
        for line in paragraph.splitlines():
            line = line.strip()
            if line:
                add_para(doc, line)

    if source_image_path.exists():
        doc.add_paragraph()
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run()
        run.add_picture(str(source_image_path), width=Cm(12))
    else:
        add_para(doc, f"[图片未找到: {source_image_path.name}]")

    doc.add_paragraph()


def parse_table_lines(lines: list[str]) -> tuple[list[str], list[list[str]]] | None:
    if len(lines) < 2:
        return None
    if not all("|" in line for line in lines[:2]):
        return None
    if not re.match(r"^\|\s*[-: |]+\|\s*$", lines[1]):
        return None

    def split_row(row: str) -> list[str]:
        row = row.strip().strip("|")
        return [strip_md(cell) for cell in row.split("|")]

    headers = split_row(lines[0])
    rows = [split_row(line) for line in lines[2:] if line.strip()]
    return headers, rows


def add_table_as_bullets(doc, headers: list[str], rows: list[list[str]]):
    for row in rows:
        parts = []
        for header, val in zip(headers, row):
            val = strip_md(val)
            if not val:
                continue
            if header in {"项目", "得分"} and len(headers) == 2:
                parts = [f"{strip_md(row[0])}：{strip_md(row[1])}"]
                break
            parts.append(f"{strip_md(header)} {val}".strip())
        if parts:
            add_bullet(doc, "，".join(parts) if len(parts) > 1 and headers[0] not in {"项目"} else parts[0])
    doc.add_paragraph()


def render_report(
    doc,
    md_path: Path,
    *,
    page_break_before: bool = False,
    source_text_path: Path | None = None,
    source_image_path: Path | None = None,
):
    text = md_path.read_text(encoding="utf-8")
    lines = text.splitlines()
    i = 0

    if page_break_before:
        add_page_break(doc)

    while i < len(lines):
        line = lines[i].rstrip()

        if not line:
            i += 1
            continue

        if line.strip() == "---":
            i += 1
            continue

        if line.startswith("# "):
            add_heading(doc, line[2:], level=1)
            i += 1
            continue

        if line.startswith("## "):
            add_heading(doc, line[3:], level=2)
            i += 1
            continue

        if line.startswith("#### "):
            add_heading(doc, line[5:], level=4)
            i += 1
            continue

        if line.startswith("### "):
            add_heading(doc, line[4:], level=3)
            i += 1
            continue

        if line.startswith("- ") and "Spec" in line:
            i += 1
            continue

        if line.startswith("- ") and "评测时间" in line:
            i += 1
            continue

        if line.startswith("- ") and "文本来源" in line:
            if source_text_path and source_image_path:
                add_source_block(doc, source_text_path, source_image_path)
            else:
                add_bullet(doc, line[2:])
            i += 1
            continue

        if line.startswith("- "):
            add_bullet(doc, line[2:])
            i += 1
            continue

        if line.startswith("|"):
            table_lines = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                table_lines.append(lines[i].strip())
                i += 1
            parsed = parse_table_lines(table_lines)
            if parsed:
                headers, rows = parsed
                add_table_as_bullets(doc, headers, rows)
            continue

        if is_field_legend(line):
            add_legend_para(doc, line)
        else:
            add_para(doc, line)
        i += 1


def build_doc() -> Document:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)

    title = doc.add_heading("服装设计文本完整评分报告", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in title.runs:
        set_run_font(run, bold=True, size=18)

    add_para(doc, "本文档汇总两份 Look 文本评估结果，各报告独立成节。")
    doc.add_paragraph()

    for idx, assets in enumerate(REPORT_ASSETS):
        render_report(
            doc,
            assets["report_md"],
            page_break_before=idx > 0,
            source_text_path=assets["source_text"],
            source_image_path=assets["source_image"],
        )

    return doc


def main():
    missing = [
        assets["report_md"]
        for assets in REPORT_ASSETS
        if not assets["report_md"].exists()
    ]
    if missing:
        raise FileNotFoundError(f"Missing report files: {missing}")

    doc = build_doc()
    try:
        doc.save(OUTPUT)
    except PermissionError:
        alt = REPORT_DIR / "完整评分报告合集_新版.docx"
        doc.save(alt)
        print(f"Saved: {alt}")
        return
    print(f"Saved: {OUTPUT}")


if __name__ == "__main__":
    main()
