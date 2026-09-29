"""三类一致性矛盾各一例：改写前后的图、全文，以及红字标出的改写点。"""

from __future__ import annotations

import json
from io import BytesIO
from pathlib import Path

from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

REPO = Path(__file__).resolve().parents[1]
CASES = REPO / "fashion_research_dir" / "consistency_defect_cases"
BEFORE_DIR = CASES / "generated_gpt-image-2"
AFTER_DIR = CASES / "generated_rewrite_v6"
OUT = Path(__file__).resolve().parent / "一致性改写三类对照.pptx"

SLIDE_W = Emu(12192000)
SLIDE_H = Emu(6858000)
FONT = "微软雅黑"

BLACK = RGBColor(0x1A, 0x1A, 0x1A)
GRAY = RGBColor(0x5C, 0x5C, 0x5C)
MUTED = RGBColor(0x8A, 0x8A, 0x8A)
LINE = RGBColor(0xE2, 0xE2, 0xE2)
PANEL = RGBColor(0xF6, 0xF6, 0xF6)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
RED = RGBColor(0xC0, 0x39, 0x2B)
RED_SOFT = RGBColor(0xF8, 0xEE, 0xEC)


def emu(inches: float) -> Emu:
    return Emu(int(round(inches * 914400)))


def set_run(run, text, *, size=12, bold=False, color=BLACK, font=FONT):
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font
    rPr = run._r.get_or_add_rPr()
    for tag in ("latin", "ea", "cs"):
        node = rPr.find(qn(f"a:{tag}"))
        if node is None:
            node = etree.SubElement(rPr, qn(f"a:{tag}"))
        node.set("typeface", font)


def add_shape(slide, l, t, w, h, fill):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, l, t, w, h)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.fill.background()
    return shape


def add_runs(slide, l, t, w, h, segments, *, size=12, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(l, t, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    tf.margin_left = emu(0.06)
    tf.margin_right = emu(0.06)
    tf.margin_top = emu(0.02)
    tf.margin_bottom = emu(0.02)
    tf._txBody.bodyPr.set("anchor", {MSO_ANCHOR.TOP: "t", MSO_ANCHOR.MIDDLE: "ctr"}[anchor])
    p = tf.paragraphs[0]
    p.alignment = align
    p.space_before = Pt(0)
    p.space_after = Pt(0)
    p.line_spacing = 1.0
    started = False
    for seg in segments:
        text, red, bold = seg[0], seg[1], seg[2]
        color = seg[3] if len(seg) > 3 else (RED if red else BLACK)
        for i, part in enumerate(text.split("\n")):
            if i > 0:
                p = tf.add_paragraph()
                p.alignment = align
                p.space_before = Pt(0)
                p.space_after = Pt(0)
                p.line_spacing = 1.0
            if part == "":
                continue
            run = p.add_run()
            set_run(run, part, size=size, bold=bold or red, color=color)
            started = True
        if not started:
            continue
    return box


def paint(text: str, reds: list[str]) -> list[tuple[str, bool, bool]]:
    marks = []
    for span in reds:
        start = text.find(span)
        if start < 0:
            raise SystemExit(f"红字片段不在原文里：{span}")
        marks.append((start, start + len(span)))
    marks.sort()
    segments: list[tuple[str, bool, bool]] = []
    cursor = 0
    for start, end in marks:
        if start < cursor:
            raise SystemExit(f"红字片段重叠：{text[start:end]}")
        if start > cursor:
            segments.append((text[cursor:start], False, False))
        segments.append((text[start:end], True, True))
        cursor = end
    if cursor < len(text):
        segments.append((text[cursor:], False, False))
    if "".join(part for part, _, _ in segments) != text:
        raise SystemExit("红字切分后与原文不一致")
    return segments


def fit_picture(slide, path: Path, box_l, box_t, box_w, box_h):
    with Image.open(path) as im:
        iw, ih = im.size
        buf = BytesIO()
        im.convert("RGB").save(buf, format="JPEG", quality=90)
        buf.seek(0)
    scale = min(box_w / iw, box_h / ih)
    pw, ph = int(iw * scale), int(ih * scale)
    left = int(box_l + (box_w - pw) / 2)
    top = int(box_t + (box_h - ph) / 2)
    slide.shapes.add_picture(buf, left, top, pw, ph)


SAMPLES = Path(r"c:\Users\lsh\Desktop\consistency_defect_k18_v6\samples.jsonl")


def load_selected() -> dict[str, dict]:
    path = AFTER_DIR / "selected_rewrites.jsonl"
    rows = {}
    for line in path.read_text(encoding="utf-8").splitlines():
        if line.strip():
            rec = json.loads(line)
            rows[rec["group_id"]] = rec
    if SAMPLES.is_file():
        with SAMPLES.open(encoding="utf-8") as handle:
            for line in handle:
                if not line.strip():
                    continue
                sample = json.loads(line)
                rec = rows.get(sample.get("group_id"))
                if not rec or sample.get("candidate_index") != rec.get("candidate_index"):
                    continue
                compact = sample.get("scores_compact") or {}
                rec["scores_coverage"] = float(compact.get("coverage_axis_score") or 0.0)
                rec["scores_quality"] = float(compact.get("quality_base_score") or 0.0)
    for rec in rows.values():
        rec.setdefault("scores_coverage", 0.0)
        rec.setdefault("scores_quality", 0.0)
    return rows


CASES_SPEC = [
    {
        "id": "01",
        "group": "consistency_01_asymmetry",
        "kind": "左右不对称",
        "theme": "黑色手工对比",
        "before_file": "01_asymmetry.txt",
        "before_png": "01_asymmetry_gpt-image-2.png",
        "after_png": "consistency_01_asymmetry_gpt-image-2.png",
        "before_red": [
            "The right half of the same jacket is high-gloss silk and sequin: no lapel, a cold shoulder, and a bishop sleeve to the knee.",
            "the right leg is a red silk culotte.",
            "the right foot is an open red silk evening sandal.",
        ],
        "after_red": [
            "matte black wool jacket with a notched lapel, set-in tailored sleeve, and notched cuff,",
            "black wool trouser with a crease,",
            "closed black satin court pump.",
        ],
        "point": [
            ("改写点：右半整侧删除，", False, True),
            ("亮片丝绸袖、冷肩主教袖、红丝绸阔腿、露趾凉鞋", True, True),
            ("不再出现。只留左半", False, False),
            ("哑光羊毛翻领夹克、压线西裤、闭口缎面浅口鞋。", True, True),
            ("\n钩针贴花和红钩针包没有写进这条入选稿。", False, False),
        ],
        "deduction": [
            {
                "tone": "ok",
                "title": "惩罚没有扣分",
                "body": [
                    ("五项都是 0：生成内容、一致性、协调性、合理性、公式模板。左右两套已经收成一套，惩罚不再扣。", False, False),
                ],
            },
            {
                "tone": "bad",
                "title": "扣在设计价值",
                "body": [
                    ("钩针贴花没留下。正文只写了翻领、袖子、西裤和鞋。", False, False),
                    ("换掉颜色和面料之后，看不出这是哪一套，", True, True),
                    ("所以设计价值三项最高只能到 0.5，合起来是 0.625，质量是 0.80。覆盖是 1.00，总分停在 0.84。", False, False),
                ],
            },
            {
                "tone": "bad",
                "title": "图上的体现",
                "body": [
                    ("入选图是一套对称的哑光黑西装：翻领、装袖、压线西裤、浅口鞋都在。", False, False),
                    ("门襟和胸前没有钩针花，红钩针包也不在。", True, True),
                ],
            },
        ],
    },
    {
        "id": "08",
        "group": "consistency_08_style_clash",
        "kind": "风格不一致",
        "theme": "夜花园刺绣",
        "before_file": "08_style_clash.txt",
        "before_png": "08_style_clash_gpt-image-2.png",
        "after_png": "consistency_08_style_clash_gpt-image-2.png",
        "before_red": [
            "a white bridal cathedral gown with orange-blossom clusters, a floor-length tulle veil, and white satin court shoes.",
            "The handmade day jacket and the bridal ceremony do not belong to one theme.",
        ],
        "after_red": [
            "colored crochet florals on a black ground,",
            "white piping.",
            "Black bike shorts",
            "red crochet handbag.",
        ],
        "point": [
            ("改写点：另一套主题删除。", False, True),
            ("白纱新娘礼服、橙花、及地头纱、白缎浅口鞋", True, True),
            ("不再出现。日装侧留下：", False, False),
            ("黑底彩色钩针、白滚边、黑骑行短裤、红钩针包。", True, True),
        ],
        "deduction": [
            {
                "tone": "ok",
                "title": "没有剩余扣分",
                "body": [
                    ("惩罚 0.00，覆盖 1.00，质量 1.00，总分 1.00。五项惩罚都是 0。", False, False),
                    ("领口和门襟上的白滚边，换掉颜色和面料之后仍然认得出，", True, True),
                    ("所以设计价值没有被压低。", False, False),
                ],
            },
            {
                "tone": "ok",
                "title": "图上对得上留下来的那一侧",
                "body": [
                    ("黑底彩色钩针、白滚边、金扣、黑骑行短裤、红钩针包都在。", True, True),
                    ("头纱、橙花和教堂纱裙不在。", False, False),
                ],
            },
        ],
    },
    {
        "id": "17",
        "group": "consistency_17_same_element_contradiction",
        "kind": "同要素矛盾",
        "theme": "悬垂丝带",
        "before_file": "17_same_element_contradiction.txt",
        "before_png": "17_same_element_contradiction_gpt-image-2.png",
        "after_png": "consistency_17_same_element_contradiction_gpt-image-2.png",
        "before_red": [
            "the same neckline is a high stand collar with no straps.",
            "the same dress is entirely black with no red.",
            "the same hem is a cropped top ending at the waist.",
        ],
        "after_red": [
            "Square-necked, red silk column dress",
            "black vertical panels",
            "loose midi to mid-calf.",
            "red, black, white, and pale beige ribbon strips.",
        ],
        "point": [
            ("改写点：同一条裙子上的第二套绑定删除。", False, True),
            ("立领、全黑、短上衣下摆", True, True),
            ("不再出现。只留", False, False),
            ("方领、红底黑竖条、中长下摆", True, True),
            ("，肩部金饰上的红、黑、白、米白丝带还在。", False, False),
        ],
        "deduction": [
            {
                "tone": "ok",
                "title": "惩罚没有扣分",
                "body": [
                    ("五项都是 0。立领、全黑、短上衣下摆已经不在，同一条裙子不再写成两套。", False, False),
                ],
            },
            {
                "tone": "bad",
                "title": "扣在覆盖和质量",
                "body": [
                    ("覆盖 0.67，质量 0.71，总分 0.70，低于 0.80 的分数门。", True, True),
                    ("它仍入选，是因为高于原文，并且惩罚更低。规则只认三种一眼能认出来的写法：", False, False),
                    ("整片面料上的花纹或肌理、沿着领口门襟下摆走的装饰、外衣敞开后单独也能看懂的内层。", True, True),
                    ("肩上的丝带没写成这三种里的任何一种，所以这三项最高只能到 0.5，质量停在 0.71。采样没有留下覆盖轴具体是哪一项没命中。", False, False),
                ],
            },
            {
                "tone": "bad",
                "title": "图上的体现",
                "body": [
                    ("一条方领红底中长裙，金饰上垂着红、黑、白、米白丝带。", False, False),
                    ("文本里的黑竖条没有画成裙身上另一块面板，而是并进了丝带。", True, True),
                    ("全文没写鞋，图上的凉鞋是生成时补上的。", True, True),
                ],
            },
        ],
    },
]


def add_title(prs: Presentation, selected: dict[str, dict]) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_shape(slide, emu(0), emu(0), SLIDE_W, emu(0.08), RED)
    add_runs(
        slide,
        emu(0.48),
        emu(0.38),
        emu(12.3),
        emu(0.55),
        [("三类一致性矛盾的改写前后对照", False, True)],
        size=28,
    )
    add_runs(
        slide,
        emu(0.48),
        emu(1.05),
        emu(12.3),
        emu(0.7),
        [(
            "各取一条入选改写。入选条件是总分高于原文，并且惩罚低于原文。\n"
            "红字标出被删掉的矛盾，以及改写后留下来的那一侧。",
            False,
            False,
        )],
        size=15,
    )
    cards = [
        ("01", "左右不对称", "同一件衣服左右两套语法"),
        ("08", "风格不一致", "手工日装上叠了另一套婚礼"),
        ("17", "同要素矛盾", "同一条裙子同时写成两套"),
    ]
    for i, (cid, kind, line) in enumerate(cards):
        spec = CASES_SPEC[i]
        rec = selected[spec["group"]]
        left = 0.48 + i * 4.2
        add_shape(slide, emu(left), emu(2.15), emu(3.95), emu(4.55), PANEL)
        add_shape(slide, emu(left), emu(2.15), emu(0.08), emu(4.55), RED)
        add_runs(
            slide,
            emu(left + 0.28),
            emu(2.4),
            emu(3.45),
            emu(0.4),
            [(cid, True, True)],
            size=22,
        )
        add_runs(
            slide,
            emu(left + 0.28),
            emu(2.95),
            emu(3.45),
            emu(0.4),
            [(kind, False, True)],
            size=18,
        )
        add_runs(
            slide,
            emu(left + 0.28),
            emu(3.5),
            emu(3.45),
            emu(0.7),
            [(line, False, False)],
            size=14,
        )
        add_runs(
            slide,
            emu(left + 0.28),
            emu(4.45),
            emu(3.45),
            emu(0.9),
            [
                (f"原文  {rec['baseline_S_fp']:.2f}  /  惩罚 {rec['baseline_penalty']:.2f}\n", False, False),
                (f"入选  {rec['S_fp']:.2f}  /  惩罚 {rec['total_penalty']:.2f}", True, True),
            ],
            size=14,
        )
        add_runs(
            slide,
            emu(left + 0.28),
            emu(5.7),
            emu(3.45),
            emu(0.6),
            [(spec["theme"], False, False)],
            size=13,
        )


def add_case(prs: Presentation, spec: dict, selected: dict[str, dict]) -> None:
    rec = selected[spec["group"]]
    before = (CASES / spec["before_file"]).read_text(encoding="utf-8").strip()
    after = str(rec["prompt"]).strip()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_shape(slide, emu(0), emu(0), SLIDE_W, emu(0.08), RED)

    add_runs(
        slide,
        emu(0.36),
        emu(0.18),
        emu(7.3),
        emu(0.38),
        [
            (f"{spec['id']}   {spec['kind']}", False, True),
            (f"    {spec['theme']}", False, False, GRAY),
        ],
        size=18,
    )
    add_runs(
        slide,
        emu(7.7),
        emu(0.18),
        emu(5.2),
        emu(0.38),
        [
            (f"原文 {rec['baseline_S_fp']:.2f} / {rec['baseline_penalty']:.2f}", False, False),
            ("   →   ", False, False),
            (f"入选 {rec['S_fp']:.2f} / {rec['total_penalty']:.2f}", True, True),
        ],
        size=14,
        align=PP_ALIGN.RIGHT,
    )

    col_w = 6.22
    gap = 0.22
    lefts = (0.36, 0.36 + col_w + gap)
    labels = ("改写前", "改写后")
    texts = (before, after)
    reds = (spec["before_red"], spec["after_red"])
    pngs = (
        BEFORE_DIR / spec["before_png"],
        AFTER_DIR / spec["after_png"],
    )
    image_top = 0.62
    image_h = 3.42
    for left, label, text, spans, png in zip(lefts, labels, texts, reds, pngs):
        add_shape(slide, emu(left), emu(image_top), emu(col_w), emu(image_h), PANEL)
        fit_picture(slide, png, emu(left + 0.08), emu(image_top + 0.08), emu(col_w - 0.16), emu(image_h - 0.16))
        add_runs(
            slide,
            emu(left),
            emu(4.08),
            emu(col_w),
            emu(0.28),
            [(label, label == "改写后", True)],
            size=12,
        )
        add_runs(
            slide,
            emu(left),
            emu(4.36),
            emu(col_w),
            emu(1.95),
            paint(text, spans),
            size=10,
        )

    add_shape(slide, emu(0.36), emu(6.42), emu(12.62), emu(0.92), RED_SOFT)
    add_shape(slide, emu(0.36), emu(6.42), emu(0.08), emu(0.92), RED)
    add_runs(slide, emu(0.52), emu(6.48), emu(12.3), emu(0.8), spec["point"], size=13, anchor=MSO_ANCHOR.MIDDLE)


def add_deduction(prs: Presentation, spec: dict, selected: dict[str, dict]) -> None:
    rec = selected[spec["group"]]
    blocks = spec["deduction"]
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_shape(slide, emu(0), emu(0), SLIDE_W, emu(0.08), RED)
    add_runs(
        slide,
        emu(0.36),
        emu(0.18),
        emu(8.2),
        emu(0.38),
        [
            (f"{spec['id']}   入选稿的扣分", False, True),
            (f"    {spec['theme']}", False, False, GRAY),
        ],
        size=18,
    )
    add_runs(
        slide,
        emu(8.4),
        emu(0.18),
        emu(4.5),
        emu(0.38),
        [
            (
                f"覆盖 {rec['scores_coverage']:.2f}   ",
                rec["scores_coverage"] < 0.999,
                False,
            ),
            (
                f"质量 {rec['scores_quality']:.2f}   ",
                rec["scores_quality"] < 0.999,
                False,
            ),
            (f"惩罚 {rec['total_penalty']:.2f}", rec["total_penalty"] > 0, False),
        ],
        size=13,
        align=PP_ALIGN.RIGHT,
    )

    add_shape(slide, emu(0.36), emu(0.7), emu(4.7), emu(6.5), PANEL)
    fit_picture(
        slide,
        AFTER_DIR / spec["after_png"],
        emu(0.48),
        emu(0.82),
        emu(4.46),
        emu(6.26),
    )

    top = 0.7
    gap = 0.14
    height = (6.5 - gap * (len(blocks) - 1)) / len(blocks)
    for block in blocks:
        fill = RED_SOFT if block["tone"] == "bad" else PANEL
        bar = RED if block["tone"] == "bad" else RGBColor(0xC8, 0xC8, 0xC8)
        add_shape(slide, emu(5.24), emu(top), emu(7.72), emu(height), fill)
        add_shape(slide, emu(5.24), emu(top), emu(0.08), emu(height), bar)
        add_runs(
            slide,
            emu(5.44),
            emu(top + 0.1),
            emu(7.35),
            emu(0.36),
            [(block["title"], block["tone"] == "bad", True)],
            size=16,
        )
        add_runs(
            slide,
            emu(5.44),
            emu(top + 0.48),
            emu(7.35),
            emu(height - 0.58),
            block["body"],
            size=14,
        )
        top += height + gap


def main() -> None:
    selected = load_selected()
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    add_title(prs, selected)
    for spec in CASES_SPEC:
        add_case(prs, spec, selected)
        add_deduction(prs, spec, selected)
    OUT.parent.mkdir(parents=True, exist_ok=True)
    try:
        prs.save(OUT)
        print(OUT)
    except PermissionError:
        fallback = OUT.with_name(OUT.stem + "_新.pptx")
        prs.save(fallback)
        print(fallback)


if __name__ == "__main__":
    main()
