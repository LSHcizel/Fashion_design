"""六页 PPT：四面图库 + 两页论文式左右对照（等宽栏、等大图）。"""

from __future__ import annotations

import json
from io import BytesIO
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

ROOT = Path(r"c:\Users\lsh\Desktop\记录10.1\生图对照")
OUT = Path(r"c:\Users\lsh\Desktop\记录10.1\极端案例生图对照.pptx")
INV = ROOT / "逆解析_降分最多10"
GEN = ROOT / "生成图_提分最多10"

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)
BLACK = RGBColor(0x1A, 0x1A, 0x1A)
GRAY = RGBColor(0x5C, 0x5C, 0x5C)
MUTED = RGBColor(0x8A, 0x8A, 0x8A)
RED = RGBColor(0xC0, 0x22, 0x22)
LINE = RGBColor(0xC8, 0xC8, 0xC8)
HEADER_RAW = RGBColor(0xE8, 0xE8, 0xE8)
HEADER_RW = RGBColor(0xD9, 0xE8, 0xF5)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
FONT = "微软雅黑"

# 定性页：左右等宽两栏（描述在上、等大正方形图在下），两页几何完全一致
M = Inches(0.32)
GUTTER = Inches(0.24)
COL_W = (SLIDE_W - M * 2 - GUTTER) / 2
HEAD_H = Inches(0.32)
DESC_H = Inches(1.45)
IMG_SIZE = Inches(3.85)
CAP_H = Inches(0.62)


def _run(p, text, *, size=12, bold=False, color=BLACK):
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = color
    r.font.name = FONT


def clip_text(text: str, limit: int = 420) -> str:
    text = " ".join(text.split())
    if len(text) <= limit:
        return text
    return text[: limit - 1].rstrip(" ,.;:") + "…"


def add_text(slide, l, t, w, h, text, *, size=12, bold=False, color=BLACK, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(l, t, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    _run(p, text, size=size, bold=bold, color=color)
    return box


def fit_picture(slide, path: Path, box_l, box_t, box_w, box_h):
    with Image.open(path) as im:
        iw, ih = im.size
        buf = BytesIO()
        im.convert("RGB").save(buf, format="JPEG", quality=88)
        buf.seek(0)
    scale = min(box_w / iw, box_h / ih)
    pw, ph = iw * scale, ih * scale
    left = box_l + (box_w - pw) / 2
    top = box_t + (box_h - ph) / 2
    slide.shapes.add_picture(buf, int(left), int(top), int(pw), int(ph))


def cases(folder: Path) -> list[Path]:
    return sorted([p for p in folder.iterdir() if p.is_dir()], key=lambda p: p.name)


def meta_of(case: Path) -> dict:
    path = case / "meta.json"
    return json.loads(path.read_text(encoding="utf-8")) if path.is_file() else {}


def add_grid_slide(prs, *, title: str, subtitle: str, folder: Path, image_name: str, score_key: str, emphasize_delta: bool):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = WHITE
    add_text(slide, Inches(0.35), Inches(0.18), Inches(12.5), Inches(0.4), title, size=24, bold=True)
    add_text(slide, Inches(0.35), Inches(0.58), Inches(12.5), Inches(0.35), subtitle, size=12, color=GRAY)

    items = cases(folder)
    cols, rows = 5, 2
    margin_l = Inches(0.28)
    margin_t = Inches(1.05)
    gap_x, gap_y = Inches(0.16), Inches(0.12)
    label_h = Inches(0.38)
    usable_w = SLIDE_W - margin_l * 2 - gap_x * (cols - 1)
    usable_h = SLIDE_H - margin_t - Inches(0.25) - gap_y * (rows - 1)
    cell_w = usable_w / cols
    cell_h = usable_h / rows
    img_h = cell_h - label_h

    for i, case in enumerate(items[:10]):
        r, c = divmod(i, cols)
        left = margin_l + c * (cell_w + gap_x)
        top = margin_t + r * (cell_h + gap_y)
        png = case / image_name
        if png.is_file():
            fit_picture(slide, png, left, top, cell_w, img_h)
        m = meta_of(case)
        score = m.get(score_key)
        delta = m.get("delta")
        score_txt = f"{score:.3f}" if isinstance(score, (int, float)) else "—"
        delta_txt = f"{delta:+.3f}" if isinstance(delta, (int, float)) else ""
        box = slide.shapes.add_textbox(left, top + img_h, cell_w, label_h)
        tf = box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        _run(p, f"{i + 1:02d}  {score_txt}", size=11, bold=True)
        if delta_txt:
            p2 = tf.add_paragraph()
            p2.alignment = PP_ALIGN.CENTER
            _run(p2, f"相对原文 {delta_txt}", size=10, bold=emphasize_delta, color=RED if emphasize_delta else GRAY)


def add_col_block(slide, l, t, w, *, header: str, score: str, body: str, header_fill: RGBColor, img: Path):
    """单栏：顶栏 + 描述 + 等大正方形图（图居中于栏宽）。"""
    # 顶栏
    head = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, l, t, w, HEAD_H)
    head.fill.solid()
    head.fill.fore_color.rgb = header_fill
    head.line.color.rgb = LINE
    hb = slide.shapes.add_textbox(l + Inches(0.1), t + Inches(0.04), w - Inches(2.0), HEAD_H - Inches(0.06))
    _run(hb.text_frame.paragraphs[0], header, size=12, bold=True)
    sb = slide.shapes.add_textbox(l + w - Inches(1.85), t + Inches(0.04), Inches(1.75), HEAD_H - Inches(0.06))
    sp = sb.text_frame.paragraphs[0]
    sp.alignment = PP_ALIGN.RIGHT
    _run(sp, score, size=12, bold=True, color=RED)

    # 描述框
    desc_t = t + HEAD_H
    outer = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, l, desc_t, w, DESC_H)
    outer.fill.solid()
    outer.fill.fore_color.rgb = WHITE
    outer.line.color.rgb = LINE
    tb = slide.shapes.add_textbox(l + Inches(0.1), desc_t + Inches(0.06), w - Inches(0.2), DESC_H - Inches(0.1))
    tf = tb.text_frame
    tf.word_wrap = True
    _run(tf.paragraphs[0], clip_text(body), size=10, color=BLACK)

    # 等大正方形图：在栏内水平居中
    img_t = desc_t + DESC_H + Inches(0.14)
    img_l = l + (w - IMG_SIZE) / 2
    frame = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, img_l, img_t, IMG_SIZE, IMG_SIZE)
    frame.fill.solid()
    frame.fill.fore_color.rgb = RGBColor(0xF2, 0xF2, 0xF2)
    frame.line.color.rgb = LINE
    with Image.open(img) as im:
        buf = BytesIO()
        im.convert("RGB").save(buf, format="JPEG", quality=92)
        buf.seek(0)
    slide.shapes.add_picture(buf, int(img_l), int(img_t), int(IMG_SIZE), int(IMG_SIZE))
    return img_t + IMG_SIZE


def add_rich_caption(slide, l, t, w, h, segments, *, size=10):
    """segments: list[(text, color, bold)]，用于红字标出改好/改坏要点。"""
    box = slide.shapes.add_textbox(l, t, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    for text, color, bold in segments:
        _run(p, text, size=size, bold=bold, color=color)
    return box


def add_qual_slide(
    prs,
    *,
    page_title: str,
    raw_header: str,
    raw_score: str,
    raw_text: str,
    raw_img: Path,
    rw_header: str,
    rw_score: str,
    rw_text: str,
    rw_img: Path,
    caption_segments: list[tuple[str, RGBColor, bool]],
):
    """左右等宽两栏：上描述、下等大正方形生图；两页几何完全一致。"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = WHITE

    add_text(slide, M, Inches(0.1), SLIDE_W - M * 2, Inches(0.34), page_title, size=18, bold=True)

    top = Inches(0.48)
    left0 = M
    left1 = M + COL_W + GUTTER

    bottom0 = add_col_block(
        slide, left0, top, COL_W,
        header=raw_header, score=raw_score, body=raw_text, header_fill=HEADER_RAW, img=raw_img,
    )
    bottom1 = add_col_block(
        slide, left1, top, COL_W,
        header=rw_header, score=rw_score, body=rw_text, header_fill=HEADER_RW, img=rw_img,
    )

    cap_t = max(bottom0, bottom1) + Inches(0.12)
    add_rich_caption(slide, M, cap_t, SLIDE_W - M * 2, CAP_H, caption_segments, size=10)


def main() -> None:
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    prs.core_properties.title = "极端案例生图对照"

    add_grid_slide(
        prs,
        title="逆解析 · 原文生图",
        subtitle="降分最多的 10 组：用原文描述生成。分数为原文总分 S_fp。",
        folder=INV,
        image_name="原文.png",
        score_key="orig_S",
        emphasize_delta=False,
    )
    add_grid_slide(
        prs,
        title="逆解析 · 改写生图",
        subtitle="同一 10 组：用组内最高改写生成。红字为相对原文的总分变化。",
        folder=INV,
        image_name="改写.png",
        score_key="best_S",
        emphasize_delta=True,
    )
    add_grid_slide(
        prs,
        title="生成图 · 原文生图",
        subtitle="提分最多的 10 组：用原文描述生成。分数为原文总分 S_fp。",
        folder=GEN,
        image_name="原文.png",
        score_key="orig_S",
        emphasize_delta=False,
    )
    add_grid_slide(
        prs,
        title="生成图 · 改写生图",
        subtitle="同一 10 组：用组内最高改写生成。红字为相对原文的总分提升。",
        folder=GEN,
        image_name="改写.png",
        score_key="best_S",
        emphasize_delta=True,
    )

    inv = INV / "01_m0.143"
    gen = GEN / "01_p0.686"
    add_qual_slide(
        prs,
        page_title="定性对照 · 逆解析（得分降低）",
        raw_header="原文 description",
        raw_score="S_fp = 0.993",
        raw_text=(inv / "原文.txt").read_text(encoding="utf-8").strip(),
        raw_img=inv / "原文.png",
        rw_header="改写 description",
        rw_score="S_fp = 0.850",
        rw_text=(inv / "改写.txt").read_text(encoding="utf-8").strip(),
        rw_img=inv / "改写.png",
        caption_segments=[
            ("实际变坏：", RED, True),
            ("原文一次写清领口、肩带金扣、红黑竖条、多色飘带、手帕摆、金项链、绿色手包与细带凉鞋；改写", BLACK, False),
            ("压成更短的柱形裙概括", RED, True),
            ("，", BLACK, False),
            ("fluid midi 与多色飘带分层被削弱", RED, True),
            ("。生图上改写版", BLACK, False),
            ("飘带更碎", RED, True),
            ("、", BLACK, False),
            ("场景偏建筑大厅", RED, True),
            ("，", BLACK, False),
            ("少了秀场观众与色彩面板呼应", RED, True),
            ("，辨识变钝。", BLACK, False),
        ],
    )
    add_qual_slide(
        prs,
        page_title="定性对照 · 生成图（得分提升）",
        raw_header="原文 description",
        raw_score="S_fp = 0.307",
        raw_text=(gen / "原文.txt").read_text(encoding="utf-8").strip(),
        raw_img=gen / "原文.png",
        rw_header="改写 description",
        rw_score="S_fp = 0.993",
        rw_text=(gen / "改写.txt").read_text(encoding="utf-8").strip(),
        rw_img=gen / "改写.png",
        caption_segments=[
            ("实际变好：", RED, True),
            ("原文 Theme/Concept 提纲末句写 plain ivory silk with no weave，与编织外层矛盾；改写", BLACK, False),
            ("删掉提纲壳和矛盾句", RED, True),
            ("，只保留焦糖编织皮大衣、象牙白丝质裙与腰间打结皮带。生图上改写版", BLACK, False),
            ("编织纹理清楚、内外层对比稳定", RED, True),
            ("；原文更容易漂成光滑浅色外套。", BLACK, False),
        ],
    )

    # 若原文件被 PPT 占用，另存为明确新文件，避免继续看旧版
    try:
        prs.save(OUT)
        print(OUT)
    except PermissionError:
        alt = OUT.with_name("极端案例生图对照_排版修正.pptx")
        prs.save(alt)
        print(alt)
    print("slides", len(prs.slides))
    print("qual geometry: COL_W=", round(float(COL_W) / 914400, 3), "IMG=", round(float(IMG_SIZE) / 914400, 3))


if __name__ == "__main__":
    main()
