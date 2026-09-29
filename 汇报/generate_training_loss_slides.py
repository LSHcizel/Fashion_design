"""把冷启动 SFT / GRPO 的损失定义和实测变化写进《改进方案及训练》汇报。

GRPO 数字来自 2026-09-13 远程
``training/runs/grpo_chanel_inverse_v2/train_coldstart.log``。
该日志从「加载 SFT 权重、进入第 1 轮 GRPO」开始，没有留下 SFT 逐步交叉熵。
"""

from __future__ import annotations

from io import BytesIO
from pathlib import Path

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

DESKTOP = Path(r"C:\Users\lsh\Desktop\改进方案及训练_v2.pptx")
LOCAL = Path(r"E:\fashion_agents_project\汇报\改进方案及训练_v2.pptx")
MARKERS = ("06 损失函数", "07 冷启动损失")

BLACK = RGBColor(0x1A, 0x1A, 0x1A)
GRAY = RGBColor(0x5A, 0x5A, 0x5A)
MUTED = RGBColor(0x6E, 0x6E, 0x6E)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
ACCENT = RGBColor(0x2F, 0x4B, 0x8A)
CARD = RGBColor(0xF7, 0xF8, 0xFA)
LINE = RGBColor(0xE2, 0xE2, 0xE2)
GREEN = RGBColor(0x1F, 0x7A, 0x4D)
FONT = "微软雅黑"
REL_NS = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"

# 四轮短 GRPO。每行：起点、三次 logging loss、轮均 train_loss。
ROUNDS = [
    ("第 1 轮", "SFT", 0.03251, 0.2727, 0.5743, 0.2809),
    ("第 2 轮", "round_01", 0.03204, 0.2723, 0.5738, 0.2804),
    ("第 3 轮", "round_02", 0.03169, 0.2718, 0.5726, 0.2797),
    ("第 4 轮", "round_03", 0.03086, 0.2716, 0.5719, 0.2792),
]


def _set_run(run, text, *, size=12, bold=False, color=BLACK, font=FONT):
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font


def add_text(
    slide,
    l,
    t,
    w,
    h,
    text,
    *,
    size=12,
    bold=False,
    color=BLACK,
    align=PP_ALIGN.LEFT,
    anchor=MSO_ANCHOR.TOP,
    font=FONT,
):
    box = slide.shapes.add_textbox(l, t, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    tf.paragraphs[0].alignment = align
    box.text_frame._txBody.bodyPr.set("anchor", {MSO_ANCHOR.TOP: "t", MSO_ANCHOR.MIDDLE: "ctr", MSO_ANCHOR.BOTTOM: "b"}[anchor])
    run = tf.paragraphs[0].add_run()
    _set_run(run, text, size=size, bold=bold, color=color, font=font)
    return box


def add_lines(slide, l, t, w, h, lines, *, size=13, color=BLACK, space=4):
    box = slide.shapes.add_textbox(l, t, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    for i, (text, kwargs) in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        p.space_after = Pt(space)
        run = p.add_run()
        _set_run(run, text, size=kwargs.get("size", size), bold=kwargs.get("bold", False), color=kwargs.get("color", color))
    return box


def slide_blob(slide) -> str:
    parts = []
    for shape in slide.shapes:
        if shape.has_text_frame:
            parts.append(shape.text_frame.text)
    return "\n".join(parts)


def delete_slide(prs: Presentation, index: int) -> None:
    sld_id = prs.slides._sldIdLst[index]
    r_id = sld_id.get(REL_NS + "id")
    prs.part.drop_rel(r_id)
    prs.slides._sldIdLst.remove(sld_id)


def move_slide(prs: Presentation, old_index: int, new_index: int) -> None:
    lst = prs.slides._sldIdLst
    el = lst[old_index]
    lst.remove(el)
    lst.insert(new_index, el)


def sidebar_image(prs: Presentation) -> bytes:
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE and shape.width < Emu(4000000):
                return shape.image.blob
    raise RuntimeError("找不到训练页左侧配图")


def paint_header(slide, title: str, image: bytes) -> None:
    add_text(slide, Emu(138430), Emu(103505), Emu(3200000), Emu(520000), "Training", size=32, bold=True, font="Calibri")
    add_text(
        slide,
        Emu(634067),
        Emu(620000),
        Emu(7000000),
        Emu(360000),
        "如何构建我们的训练流程？",
        size=16,
        bold=True,
        color=GRAY,
    )
    add_text(slide, Emu(3963035), Emu(430000), Emu(7600000), Emu(420000), title, size=22, bold=True)
    slide.shapes.add_picture(BytesIO(image), Emu(354330), Emu(1212215), Emu(3049270), Emu(4422775))


def shade(cell, rgb: RGBColor) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    solid = tc_pr.makeelement(qn("a:solidFill"), {})
    srgb = solid.makeelement(qn("a:srgbClr"), {"val": f"{rgb[0]:02X}{rgb[1]:02X}{rgb[2]:02X}"})
    solid.append(srgb)
    tc_pr.append(solid)


def set_cell(cell, text, *, size=11, bold=False, color=BLACK, fill=None, align=PP_ALIGN.CENTER):
    cell.text = ""
    tf = cell.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    _set_run(run, text, size=size, bold=bold, color=color)
    if fill is not None:
        shade(cell, fill)


def build_definition(prs: Presentation, image: bytes):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    paint_header(slide, "06 损失函数：SFT 与 GRPO", image)

    left = Emu(3680000)
    width = Emu(8100000)

    card1 = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, Emu(1000000), width, Emu(2100000))
    card1.fill.solid()
    card1.fill.fore_color.rgb = CARD
    card1.line.color.rgb = LINE
    add_lines(
        slide,
        Emu(3880000),
        Emu(1080000),
        Emu(7700000),
        Emu(1900000),
        [
            ("SFT　交叉熵", {"bold": True, "size": 16, "color": ACCENT}),
            ("一组 K 条里只留 r_Q 最高的改写。用户侧 token 掩掉，只让模型逐词预测这条冠军改写。", {"size": 13}),
            ("学习率 2e-5，最多 1 个 epoch。loss 的指数滑动平均在 100 step 之后，连续 6 次日志不再明显下降就早停。", {"size": 13}),
            ("loss 下降表示：看到同一篇原文时，冠军改写每个 token 的概率在升高。", {"size": 13}),
        ],
        space=3,
    )

    card2 = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, Emu(3240000), width, Emu(2500000))
    card2.fill.solid()
    card2.fill.fore_color.rgb = CARD
    card2.line.color.rgb = LINE
    add_lines(
        slide,
        Emu(3880000),
        Emu(3320000),
        Emu(7700000),
        Emu(2320000),
        [
            ("GRPO　L = −mean(A · log π) + β · mean((log π − log π_ref)²)", {"bold": True, "size": 15, "color": ACCENT}),
            ("策略项：组内优势 A>0 的改写被抬高概率，A<0 的被压低。优势来自冻结的 r_Q，训练中不再重打分。", {"size": 13}),
            ("KL 项：β = 0.04，把当前改写器拴在训练前快照上。日志里的 loss 是两项之和；这次没有把两项拆开记录。", {"size": 13}),
            ("四轮短更新，每轮 0.25 epoch，学习率 1e-6。后一轮从上一轮的 LoRA 接着训，参考模型全程冻结。", {"size": 13}),
        ],
        space=3,
    )

    add_text(
        slide,
        left,
        Emu(5900000),
        width,
        Emu(700000),
        "冷启动日志从「加载 hf_checkpoints/sft、进入第 1 轮 GRPO」开始。SFT 的逐步交叉熵没有留在这份日志里；SFT 权重是第 1 轮的起点，说明监督微调已经写完。下一页是四轮 GRPO 的实测。",
        size=12,
        color=MUTED,
    )
    return slide


def build_curve(prs: Presentation, image: bytes):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    paint_header(slide, "07 冷启动实测：GRPO 损失下降", image)

    headers = ["轮次", "起点", "epoch 0.08", "epoch 0.17", "epoch 0.25", "轮均 train_loss", "相对第 1 轮"]
    rows = [headers]
    base = ROUNDS[0][5]
    for name, start, a, b, c, mean in ROUNDS:
        delta = (mean - base) / base * 100
        rows.append(
            [
                name,
                start,
                f"{a:.5f}".rstrip("0").rstrip("."),
                f"{b:.4f}",
                f"{c:.4f}",
                f"{mean:.4f}",
                "—" if name == "第 1 轮" else f"{delta:.2f}%",
            ]
        )

    table_w = Emu(8100000)
    col_w = [1100000, 1100000, 1200000, 1200000, 1200000, 1400000, 900000]
    shape = slide.shapes.add_table(5, 7, Emu(3680000), Emu(980000), table_w, Emu(1450000))
    table = shape.table
    for i, w in enumerate(col_w):
        table.columns[i].width = Emu(w)
    for r, row in enumerate(rows):
        for c, val in enumerate(row):
            fill = ACCENT if r == 0 else (RGBColor(0xE8, 0xF5, 0xEE) if c == 6 and r == 4 else WHITE)
            color = WHITE if r == 0 else (GREEN if c == 6 and r > 1 else BLACK)
            set_cell(table.cell(r, c), val, size=11, bold=(r == 0 or c == 0), color=color, fill=fill)

    chart_data = CategoryChartData()
    chart_data.categories = [row[0] for row in ROUNDS]
    chart_data.add_series("轮均 train_loss", [row[5] for row in ROUNDS])
    chart = slide.shapes.add_chart(
        XL_CHART_TYPE.LINE_MARKERS,
        Emu(3680000),
        Emu(2520000),
        Emu(4300000),
        Emu(2500000),
        chart_data,
    ).chart
    chart.has_legend = True
    chart.legend.position = XL_LEGEND_POSITION.BOTTOM
    chart.legend.include_in_layout = False
    axis = chart.value_axis
    axis.minimum_scale = 0.2785
    axis.maximum_scale = 0.2815
    axis.has_major_gridlines = True
    axis.tick_labels.number_format = "0.0000"
    chart.has_title = True
    chart.chart_title.text_frame.paragraphs[0].text = "轮均 train_loss（纵轴放大）"

    add_lines(
        slide,
        Emu(8080000),
        Emu(2580000),
        Emu(3700000),
        Emu(2500000),
        [
            ("怎么读这组数", {"bold": True, "size": 14, "color": ACCENT}),
            ("同一 epoch 的三个检查点，四轮全部逐轮下降。", {"size": 12}),
            ("轮均从 0.2809 降到 0.2792，合计 −0.61%。", {"size": 12}),
            ("单轮里 0.03 → 0.57 是不同数据窗口的滑动均值，轮均仍约 0.28。", {"size": 12}),
            ("梯度范数停在 0.31–0.55，训练没有发散。", {"size": 12}),
        ],
        space=6,
    )

    add_text(
        slide,
        Emu(3680000),
        Emu(5200000),
        Emu(8100000),
        Emu(1400000),
        "来源：grpo_chanel_inverse_v2 / train_coldstart.log（2026-09-13）。每轮约 3.7 分钟、0.25 epoch，β=0.04。\n"
        "下降幅度小，和 1e-6 的学习率以及 KL 约束一致：策略在按 GRPO 目标更新，同时留在训练前改写器附近。\n"
        "这组曲线说明优化是稳定生效的。改写文本是否更好，要用 round_04 再做一轮 K 路抽检。",
        size=12,
        color=GRAY,
    )
    return slide


def insert_before_summary(prs: Presentation) -> None:
    image = sidebar_image(prs)
    # 去掉上一次插入的同名页，避免重复。
    kill = [i for i, s in enumerate(prs.slides) if any(m in slide_blob(s) for m in MARKERS)]
    for i in reversed(kill):
        delete_slide(prs, i)

    build_definition(prs, image)
    build_curve(prs, image)
    # 追加后顺序是「总结, 定义, 曲线」。把总结挪到最后。
    n = len(prs.slides)
    move_slide(prs, n - 3, n - 1)


def save_all(prs: Presentation) -> list[Path]:
    written = []
    for path in (DESKTOP, LOCAL):
        try:
            prs.save(str(path))
            written.append(path)
        except PermissionError:
            alt = path.with_name(path.stem + "_含损失曲线.pptx")
            prs.save(str(alt))
            written.append(alt)
    return written


def main() -> None:
    src = DESKTOP if DESKTOP.is_file() else LOCAL
    prs = Presentation(str(src))
    insert_before_summary(prs)
    paths = save_all(prs)
    for p in paths:
        print(p)


if __name__ == "__main__":
    main()
