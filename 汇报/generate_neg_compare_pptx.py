"""16-group original vs rewrite comparison PPT, one look per slide."""

from __future__ import annotations

import json
from io import BytesIO
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Emu, Pt

ROOT = Path(r"c:\Users\lsh\Desktop\grpo_chanel_inverse_v3") / "负样本对照"
ORIG = ROOT / "01_原文原图"
NEW = ROOT / "02_高分改写新图"
OUT = ROOT / "16组原文vs高分改写对照_含设计价值.pptx"
OUT_FALLBACK = ROOT / "16组原文vs高分改写对照_含设计价值_新.pptx"
EXISTING = ROOT / "16组原文vs高分改写对照_完整描述.pptx"
EXISTING_WITH_PRINCIPLES = ROOT / "16组原文vs高分改写对照_完整描述_含改写原则.pptx"
DM_ACCENT = RGBColor(0x2F, 0x4B, 0x8A)

SLIDE_W = Emu(12192000)
SLIDE_H = Emu(6858000)
BLACK = RGBColor(0x1A, 0x1A, 0x1A)
GRAY = RGBColor(0x5A, 0x5A, 0x5A)
MUTED = RGBColor(0x8A, 0x8A, 0x8A)
LINE = RGBColor(0xD8, 0xD8, 0xD8)
HEADER = RGBColor(0xF4, 0xF4, 0xF4)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
ACCENT = RGBColor(0x2F, 0x4B, 0x8A)
FONT = "微软雅黑"


def _set_run(run, text, *, size=12, bold=False, color=BLACK, font=FONT):
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font


def add_text(slide, l, t, w, h, text, *, size=12, bold=False, color=BLACK, align=PP_ALIGN.LEFT, anchor="t"):
    box = slide.shapes.add_textbox(l, t, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    tf._txBody.bodyPr.set("anchor", anchor)
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    _set_run(run, text, size=size, bold=bold, color=color)
    return box


def add_paragraphs(slide, l, t, w, h, text: str, *, size=10, color=BLACK):
    box = slide.shapes.add_textbox(l, t, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    lines = (text or "").split("\n")
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        p.space_after = Pt(3)
        run = p.add_run()
        _set_run(run, line if line else " ", size=size, color=color)
    return box


def fit_picture(slide, path: Path, box_l, box_t, box_w, box_h):
    with Image.open(path) as im:
        iw, ih = im.size
        rgb = im.convert("RGB")
        buf = BytesIO()
        rgb.save(buf, format="JPEG", quality=90)
        buf.seek(0)
    scale = min(box_w / iw, box_h / ih)
    pw = int(iw * scale)
    ph = int(ih * scale)
    left = int(box_l + (box_w - pw) / 2)
    top = int(box_t + (box_h - ph) / 2)
    slide.shapes.add_picture(buf, left, top, pw, ph)


def short_label(gid: str) -> str:
    parts = gid.split("_")
    num = parts[1] if len(parts) > 1 else gid
    rest = "_".join(parts[2:]).replace("looks_original_", "orig_")
    return f"{num}  {rest}"


def load_groups() -> list[dict]:
    data = json.loads((ROOT / "对照一览.json").read_text(encoding="utf-8"))
    return data["groups"]


def score_of(folder: Path) -> dict:
    p = folder / "得分.json"
    if not p.is_file():
        return {}
    return json.loads(p.read_text(encoding="utf-8"))


def read_full(path: Path) -> str:
    if not path.is_file():
        return ""
    return path.read_text(encoding="utf-8").replace("\r\n", "\n").strip()


def delta_zh(delta: float, *, label: str = "分数") -> str:
    if delta > 1e-6:
        return f"{label}提高 {delta:.4f}"
    if delta < -1e-6:
        return f"{label}下降 {abs(delta):.4f}"
    return f"{label}持平"


def fmt_score(value, digits: int = 4) -> str:
    if value is None:
        return "—"
    return f"{float(value):.{digits}f}"


def dm_of(sc: dict, group: dict | None = None, *, side: str = "original"):
    if sc.get("DesignMerit") is not None:
        return sc.get("DesignMerit")
    detail = sc.get("DesignMerit_detail")
    if isinstance(detail, dict) and detail.get("score") is not None:
        return detail.get("score")
    if group:
        key = "original_DesignMerit" if side == "original" else "best_DesignMerit"
        return group.get(key)
    return None


def chunk_text(text: str, max_chars: int = 2200) -> list[str]:
    text = (text or "").strip()
    if not text:
        return [""]
    paras = text.split("\n")
    chunks: list[str] = []
    buf: list[str] = []
    n = 0
    for p in paras:
        extra = len(p) + (1 if buf else 0)
        if buf and n + extra > max_chars:
            chunks.append("\n".join(buf))
            buf = [p]
            n = len(p)
        else:
            buf.append(p)
            n += extra
    if buf:
        chunks.append("\n".join(buf))
    return chunks or [""]


def blank_slide(prs: Presentation):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, SLIDE_H)
    bg.fill.solid()
    bg.fill.fore_color.rgb = WHITE
    bg.line.fill.background()
    return slide


def header_bar(slide, title: str, subtitle: str = "") -> None:
    bar_h = Emu(720000) if subtitle else Emu(640000)
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, bar_h)
    bar.fill.solid()
    bar.fill.fore_color.rgb = HEADER
    bar.line.fill.background()
    add_text(
        slide,
        Emu(360000),
        Emu(140000) if subtitle else Emu(160000),
        Emu(11400000),
        Emu(280000) if subtitle else Emu(360000),
        title,
        size=20,
        bold=True,
        color=BLACK,
        anchor="ctr",
    )
    if subtitle:
        add_text(
            slide,
            Emu(360000),
            Emu(420000),
            Emu(11400000),
            Emu(240000),
            subtitle,
            size=12,
            color=GRAY,
            anchor="t",
        )


def add_card(slide, l, t, w, h, title: str, body: str) -> None:
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, l, t, w, h)
    card.fill.solid()
    card.fill.fore_color.rgb = RGBColor(0xFA, 0xFA, 0xFA)
    card.line.color.rgb = LINE
    add_text(slide, l + Emu(120000), t + Emu(80000), w - Emu(240000), Emu(280000), title, size=14, bold=True)
    add_paragraphs(
        slide,
        l + Emu(120000),
        t + Emu(380000),
        w - Emu(240000),
        h - Emu(460000),
        body,
        size=12,
    )


def rewrite_principle_slides(prs: Presentation) -> int:
    """Explain rewrite rules and why images barely change. Returns slide count added."""
    s1 = blank_slide(prs)
    header_bar(
        s1,
        "改写 prompt 原则",
        "改写器只看见文本、看不见图。目标是把同一套 look 写成更可成像的英文 prompt，不是另做一套设计。",
    )
    cards = [
        (
            "锁住风格概念",
            "识别性想法（若原文已有）、主干衣类、轮廓语言、主色、主面料家族保持不变。\n"
            "禁止换成另一套 look：不能工装变晚装，不能另起一个竞争身份，不能换主干品类。",
        ),
        (
            "允许局部调整",
            "为了更好成像，可以改或换同风格内的局部：领/袖/摆/袋处理、已有沿边饰边的介质或密度、五金、腰带、首饰、已有包、同族鞋、已有内层的表面。\n"
            "优先替换/调整，不新加一件衣服。缺的品类保持缺失，不补全。",
        ),
        (
            "写法，不是新设计",
            "一段英文、信息密集、按 主体/想法 → 轮廓 → 材料颜色 → 配饰 重排。\n"
            "原文若已有识别性想法，放到段首并占主体；明线、暗襟、漫步/沙龙/姿态句压成从句或删掉。",
        ),
        (
            "不能事后发明",
            "原文没有的表面、饰边、第二身份，改写不得编造。\n"
            "局部替换必须服务已有想法，不能用新想法把它换掉。",
        ),
    ]
    gap = Emu(160000)
    card_w = Emu(5600000)
    card_h = Emu(2200000)
    left0 = Emu(280000)
    left1 = left0 + card_w + gap
    top0 = Emu(980000)
    top1 = top0 + card_h + Emu(140000)
    positions = [(left0, top0), (left1, top0), (left0, top1), (left1, top1)]
    for (x, y), (title, body) in zip(positions, cards):
        add_card(s1, x, y, card_w, card_h, title, body)

    s2 = blank_slide(prs)
    header_bar(
        s2,
        "为什么改写几乎没有改图",
        "分数涨的是文本怎么写；生图读的是还在不在的服装事实。",
    )
    add_paragraphs(
        s2,
        Emu(360000),
        Emu(980000),
        Emu(11400000),
        Emu(5400000),
        "1.  生图模型不读 R_content。同一套 GRPO 裁判给改写加分，是因为压缩散文、重排结构、去掉姿态句、把已有事实写清楚；这些大多不改变像素里的衣类、颜色、轮廓。\n\n"
        "2.  主干被锁死。外套 / 内搭 / 下装 / 主色 / 主面料家族必须保持。图模型看见的仍是同一套 SKU，构图和光线会变，设计身份不会变。\n\n"
        "3.  负样本原文往往没有可提前的识别性想法。改写不能编造亮片、沿边饰边或第二身份。原文若是成衣叠穿，改写后仍是成衣叠穿，只是更短、更像 prompt。\n\n"
        "4.  对照里涨分最大的一组，衣服事实几乎没动：仍是白短夹克、条纹内搭、海军短裤。变的是版式——分节长文收成一段——不是另一件设计。\n\n"
        "所以：改写优化的是「同一套衣服如何被写成图 prompt」，不是「做成另一套衣服」。要让图有直观设计差异，识别性想法必须在生成阶段就写出来。",
        size=15,
    )
    return 2


def fill_redesign_cards(slide) -> None:
    cards = [
        (
            "1  拆成两种改写",
            "保真改写（现状）：锁主干衣类、轮廓、主色、主面料，只把同一套 look 写成更好的图 prompt。\n"
            "再设计改写（新）：锁主题与品牌，换识别性想法；为服务新想法，允许换轮廓和品类。\n"
            "两条指令不要写进同一个 prompt，否则模型会继续保真。",
        ),
        (
            "2  按原文质量切换",
            "原文已有清楚的识别性想法 → 走保真，把想法提前写清即可。\n"
            "原文是成衣公式、设计价值低 → 走再设计。新想法只从本章概念/元素里的想法池抽取，不凭空编造亮片或第二身份。",
        ),
        (
            "3  让 K 路真正分叉",
            "现在 K=10 是同一把锁下的温度阶梯，得到的是近义改写，图不会散开。\n"
            "再设计时：每条候选绑定本章一个不同想法（表面场、沿边、开合内层、量感对撞，或其他）。同主题、不同身份，生图才会直观不同。",
        ),
        (
            "4  换一把选优尺子",
            "现在只采纳 R_content 高于原文的改写。换了设计之后，原文衣组覆盖往往下降，好方案会被丢掉。\n"
            "再设计应按设计价值、与原文的视觉差异、是否仍在主题内来选，而不是只比总分。",
        ),
    ]
    gap = Emu(160000)
    card_w = Emu(5600000)
    card_h = Emu(2200000)
    left0 = Emu(280000)
    left1 = left0 + card_w + gap
    top0 = Emu(980000)
    top1 = top0 + card_h + Emu(140000)
    for (x, y), (title, body) in zip(
        [(left0, top0), (left1, top0), (left0, top1), (left1, top1)],
        cards,
    ):
        add_card(slide, x, y, card_w, card_h, title, body)


def add_landing_order_slide(prs: Presentation) -> None:
    s2 = blank_slide(prs)
    header_bar(
        s2,
        "建议落地顺序",
        "改写要「完全不同」，必须同时改指令、改采样、改选优；只放宽一句「可以更大胆」不够。",
    )
    add_paragraphs(
        s2,
        Emu(360000),
        Emu(980000),
        Emu(11400000),
        Emu(5400000),
        "第一步　上游继续把本章 2–4 条互不相同的识别性想法写进概念和元素，作为再设计的合法想法池。没有池子，改写仍只能编造或保真。\n\n"
        "第二步　改写器增加 redesign 模式。extra_context 只注入「本章尚未用过的想法 + 主题约束」，明确允许为该想法更换轮廓与主干品类。\n\n"
        "第三步　K 路按想法分配，而不是只调温度。10 条候选应对 4 个想法做覆盖，而不是 10 次同义改写。\n\n"
        "第四步　负样本 / 低设计价值样本用设计价值与视觉差异选优；正样本或高设计价值样本继续保真，用 R_content 选更可成像的写法。\n\n"
        "不要做的　在保真 prompt 里要求「做成完全不同的设计」——与锁主干直接冲突，模型仍会只改句子。也不要无池子地允许乱编，那会离开主题。",
        size=15,
    )


def redesign_solution_slides(prs: Presentation) -> int:
    s1 = blank_slide(prs)
    header_bar(
        s1,
        "如何通过改写来带来完全不同的设计？",
        "从多个角度来探索其可能性。前提：仍锁主题 / 子主题 / 品牌，换的是识别性想法，不是换一场秀。",
    )
    fill_redesign_cards(s1)
    add_landing_order_slide(prs)
    return 2


def move_slide(prs: Presentation, old_index: int, new_index: int) -> None:
    sld_id_lst = prs.slides._sldIdLst  # noqa: SLF001
    slides = list(sld_id_lst)
    el = slides[old_index]
    sld_id_lst.remove(el)
    sld_id_lst.insert(new_index, el)


def overview_slide(prs: Presentation, groups: list[dict]) -> None:
    slide = blank_slide(prs)
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, Emu(720000))
    bar.fill.solid()
    bar.fill.fore_color.rgb = HEADER
    bar.line.fill.background()
    add_text(
        slide,
        Emu(360000),
        Emu(160000),
        Emu(11400000),
        Emu(420000),
        "负样本 16 组对照  ·  原文原图  vs  最高分改写新图",
        size=22,
        bold=True,
        color=BLACK,
        anchor="ctr",
    )
    add_text(
        slide,
        Emu(360000),
        Emu(760000),
        Emu(11400000),
        Emu(280000),
        "同一套 GRPO R_content。先看改写原则与「为何图几乎不变」，再逐组对照原图 / 新图与完整描述。",
        size=12,
        color=GRAY,
    )

    rows = [["组", "来源", "原文 R", "改写 R", "分数变化"]]
    for g in groups:
        gid = g["group_id"]
        src = (g.get("source_path") or "").replace("fashion_research_dir/workflow_0/", "")
        rows.append(
            [
                gid.split("_")[1],
                src.replace("/look_", " / look_"),
                f"{g['original_R_content']:.4f}",
                f"{g['best_R_content']:.4f}",
                delta_zh(float(g.get("delta_R") or 0)),
            ]
        )

    table = slide.shapes.add_table(
        len(rows), 5, Emu(320000), Emu(1080000), Emu(11560000), Emu(5400000)
    ).table
    widths = [900000, 5400000, 1500000, 1500000, 2260000]
    for i, w in enumerate(widths):
        table.columns[i].width = w
    for r, row in enumerate(rows):
        for c, val in enumerate(row):
            cell = table.cell(r, c)
            cell.text = val
            for p in cell.text_frame.paragraphs:
                p.alignment = PP_ALIGN.CENTER if c != 1 else PP_ALIGN.LEFT
                for run in p.runs:
                    run.font.name = FONT
                    run.font.size = Pt(10 if r else 11)
                    run.font.bold = r == 0
                    run.font.color.rgb = BLACK
            cell.fill.solid()
            cell.fill.fore_color.rgb = HEADER if r == 0 else WHITE


def compare_slide(prs: Presentation, g: dict, idx: int, total: int) -> None:
    gid = g["group_id"]
    slide = blank_slide(prs)
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, Emu(640000))
    bar.fill.solid()
    bar.fill.fore_color.rgb = HEADER
    bar.line.fill.background()

    add_text(
        slide,
        Emu(280000),
        Emu(90000),
        Emu(8000000),
        Emu(240000),
        f"{idx:02d} / {total:02d}    {short_label(gid)}",
        size=16,
        bold=True,
        color=BLACK,
        anchor="ctr",
    )
    add_text(
        slide,
        Emu(8200000),
        Emu(90000),
        Emu(3700000),
        Emu(240000),
        delta_zh(float(g.get("delta_R") or 0)),
        size=16,
        bold=True,
        color=ACCENT,
        align=PP_ALIGN.RIGHT,
        anchor="ctr",
    )
    add_text(
        slide,
        Emu(280000),
        Emu(340000),
        Emu(11600000),
        Emu(240000),
        Path(g.get("source_path") or "").as_posix(),
        size=10,
        color=MUTED,
        anchor="t",
    )

    col_w = Emu(5600000)
    gap = Emu(200000)
    left0 = Emu(280000)
    left1 = left0 + col_w + gap
    label_t = Emu(680000)
    img_t = Emu(1120000)
    img_h = Emu(5400000)

    o_dir = ORIG / gid
    n_dir = NEW / gid
    o_sc = score_of(o_dir)
    n_sc = score_of(n_dir)

    def col_header(x, title, sc):
        r = sc.get("R_content")
        add_text(
            slide,
            x,
            label_t,
            col_w,
            Emu(380000),
            f"{title}    R_content {float(r):.4f}" if r is not None else title,
            size=13,
            bold=True,
            color=BLACK,
            anchor="ctr",
        )

    col_header(left0, "原文 + 原图", o_sc)
    col_header(left1, "最高分改写 + 新图", n_sc)

    divider = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, left1 - Emu(90000), img_t, Emu(18000), img_h
    )
    divider.fill.solid()
    divider.fill.fore_color.rgb = LINE
    divider.line.fill.background()

    o_img = o_dir / "原图.png"
    n_img = n_dir / "新图.png"
    if o_img.is_file():
        fit_picture(slide, o_img, left0, img_t, col_w, img_h)
    else:
        add_text(
            slide,
            left0,
            img_t,
            col_w,
            img_h,
            "原图缺失",
            size=16,
            color=MUTED,
            align=PP_ALIGN.CENTER,
            anchor="ctr",
        )
    if n_img.is_file():
        fit_picture(slide, n_img, left1, img_t, col_w, img_h)
    else:
        add_text(
            slide,
            left1,
            img_t,
            col_w,
            img_h,
            "新图缺失",
            size=16,
            color=MUTED,
            align=PP_ALIGN.CENTER,
            anchor="ctr",
        )


def description_slides(prs: Presentation, g: dict, idx: int, total: int) -> None:
    gid = g["group_id"]
    o_chunks = chunk_text(read_full(ORIG / gid / "原描述.txt"), 2200)
    n_chunks = chunk_text(read_full(NEW / gid / "高分改写.txt"), 2200)
    pages = max(len(o_chunks), len(n_chunks))
    col_w = Emu(5600000)
    gap = Emu(200000)
    left0 = Emu(280000)
    left1 = left0 + col_w + gap
    body_t = Emu(1080000)
    body_h = Emu(5450000)

    for pi in range(pages):
        slide = blank_slide(prs)
        bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, Emu(640000))
        bar.fill.solid()
        bar.fill.fore_color.rgb = HEADER
        bar.line.fill.background()
        cont = f"  ·  描述 {pi + 1}/{pages}" if pages > 1 else ""
        add_text(
            slide,
            Emu(280000),
            Emu(90000),
            Emu(8000000),
            Emu(240000),
            f"{idx:02d} / {total:02d}    {short_label(gid)}    完整描述{cont}",
            size=16,
            bold=True,
            color=BLACK,
            anchor="ctr",
        )
        add_text(
            slide,
            Emu(8200000),
            Emu(90000),
            Emu(3700000),
            Emu(240000),
            delta_zh(float(g.get("delta_R") or 0)),
            size=16,
            bold=True,
            color=ACCENT,
            align=PP_ALIGN.RIGHT,
            anchor="ctr",
        )
        add_text(slide, left0, Emu(680000), col_w, Emu(320000), "原文描述", size=13, bold=True)
        add_text(slide, left1, Emu(680000), col_w, Emu(320000), "最高分改写描述", size=13, bold=True)
        add_paragraphs(
            slide,
            left0,
            body_t,
            col_w,
            body_h,
            o_chunks[pi] if pi < len(o_chunks) else "",
            size=10,
        )
        add_paragraphs(
            slide,
            left1,
            body_t,
            col_w,
            body_h,
            n_chunks[pi] if pi < len(n_chunks) else "",
            size=10,
        )


def _slide_has_text(slide, needle: str) -> bool:
    for shape in slide.shapes:
        if shape.has_text_frame and needle in (shape.text_frame.text or ""):
            return True
    return False


def _find_slide_index(prs: Presentation, needle: str) -> int:
    for i, slide in enumerate(prs.slides):
        if _slide_has_text(slide, needle):
            return i
    return -1


def insert_slides_after(prs: Presentation, needle: str, added: int, orig_n: int) -> None:
    idx = _find_slide_index(prs, needle)
    insert_at = idx + 1 if idx >= 0 else orig_n
    for i in range(added):
        move_slide(prs, orig_n + i, insert_at + i)


def insert_principles_into_existing(path: Path) -> Path:
    prs = Presentation(str(path))
    if any(_slide_has_text(s, "改写 prompt 原则") for s in prs.slides):
        print(f"already present → {path}")
        return path
    orig_n = len(prs.slides)
    added = rewrite_principle_slides(prs)
    for i in range(added):
        move_slide(prs, orig_n + i, 1 + i)
    first = prs.slides[0]
    for shape in first.shapes:
        if not shape.has_text_frame:
            continue
        for p in shape.text_frame.paragraphs:
            for run in p.runs:
                if "每组先看图" in (run.text or "") or "完整原文描述" in (run.text or ""):
                    run.text = (
                        "同一套 GRPO R_content。先看改写原则与「为何图几乎不变」，"
                        "再逐组对照原图 / 新图与完整描述。"
                    )
    return _save_prs(prs, path)


def insert_solutions_into_existing(path: Path) -> Path:
    prs = Presentation(str(path))
    if any(_slide_has_text(s, "拆成两种改写") for s in prs.slides):
        print(f"solutions already present → {path}")
        return path
    title_idx = _find_slide_index(prs, "如何通过改写来带来完全不同的设计")
    if title_idx >= 0:
        fill_redesign_cards(prs.slides[title_idx])
        orig_n = len(prs.slides)
        add_landing_order_slide(prs)
        insert_slides_after(prs, "如何通过改写来带来完全不同的设计", 1, orig_n)
        return _save_prs(prs, path)
    orig_n = len(prs.slides)
    added = redesign_solution_slides(prs)
    insert_slides_after(prs, "为什么改写几乎没有改图", added, orig_n)
    return _save_prs(prs, path)


def _save_prs(prs: Presentation, path: Path) -> Path:
    out = path
    try:
        prs.save(str(out))
    except PermissionError:
        out = path.with_name(path.stem + "_新.pptx")
        prs.save(str(out))
    print(f"saved {len(prs.slides)} slides → {out}")
    return out


def main() -> None:
    groups = load_groups()
    target = EXISTING_WITH_PRINCIPLES if EXISTING_WITH_PRINCIPLES.is_file() else EXISTING
    if target.is_file():
        target = insert_principles_into_existing(target)
        insert_solutions_into_existing(target)
        return
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    overview_slide(prs, groups)
    rewrite_principle_slides(prs)
    redesign_solution_slides(prs)
    for i, g in enumerate(groups, 1):
        compare_slide(prs, g, i, len(groups))
        description_slides(prs, g, i, len(groups))
    prs.save(OUT)
    print(f"slides={len(prs.slides)} → {OUT}")


if __name__ == "__main__":
    main()
