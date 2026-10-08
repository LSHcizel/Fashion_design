"""Paper-quality GRPO slides (PromptEnhancer Fig.2 visual language).

Notes pages 3 & 5: English titles, Chinese body.
"""

from __future__ import annotations

from pathlib import Path

import matplotlib.pyplot as plt
import numpy as np
from matplotlib.patches import (
    Circle,
    Ellipse,
    FancyArrowPatch,
    FancyBboxPatch,
    PathPatch,
    Polygon,
    Rectangle,
)
from matplotlib.path import Path as MPath
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parent
FIG_DIR = ROOT / "_grpo_paper_figs"
OUT = ROOT / "GRPO训练指标说明.pptx"

# Paper-like palette from PromptEnhancer figures
BLUE = "#4A7AB5"
BLUE_SOFT = "#EAF1F8"
BLUE_EDGE = "#6E96C0"
BLUE_TITLE = "#3A6EA5"
GRAY = "#5A5A5A"
GRAY_BOX = "#F5F6F8"
GRAY_EDGE = "#B7BFC9"
DARK = "#1F1F1F"
MUTED = "#6B6B6B"
RED = "#C0392B"
RED_SOFT = "#FDF0EE"
TEAL = "#3D8B84"
TEAL_SOFT = "#E7F4F2"
WHITE = "#FFFFFF"
GOLD = "#C9A227"

FONT = "Calibri"
plt.rcParams.update(
    {
        "font.family": FONT,
        "font.size": 9.5,
        "axes.unicode_minus": False,
        "pdf.fonttype": 42,
        "ps.fonttype": 42,
    }
)

N_GROUPS = 400
K = 8

STEPS = list(range(10, 401, 10))
GRPO_LOSS = [
    -0.010274, 0.006312, -0.000759, -0.021112, -0.00586, -0.004037, -0.005829,
    -0.019473, -0.018871, -0.007303, -0.052356, -0.009312, -0.030264, -0.048467,
    -0.007083, -0.014794, -0.007695, -0.015513, -0.022196, -0.021924, -0.03397,
    -0.02008, -0.041561, -0.018383, -0.017754, -0.030136, -0.022034, -0.051655,
    -0.036806, -0.071586, -0.056007, -0.068061, -0.042892, -0.046843, -0.043867,
    -0.108142, -0.071354, -0.050435, -0.113277, -0.113385,
]
KL_TERM = [
    0.108852, 0.102043, 0.079596, 0.084752, 0.080303, 0.095568, 0.068525,
    0.08181, 0.087994, 0.074789, 0.088487, 0.082941, 0.089219, 0.056166,
    0.072793, 0.055407, 0.047475, 0.081833, 0.051306, 0.050768, 0.052196,
    0.05733, 0.056876, 0.04822, 0.038314, 0.028801, 0.041646, 0.032249,
    0.030067, 0.031558, 0.039144, 0.04254, 0.030854, 0.052956, 0.100262,
    0.101886, 0.1245, 0.115473, 0.223455, 0.141849,
]


# ── primitives ───────────────────────────────────────────────────────────────

def _canvas(ax):
    ax.set_xlim(0, 100)
    ax.set_ylim(0, 100)
    ax.axis("off")
    ax.set_facecolor(WHITE)


def _round(ax, x, y, w, h, *, fc=GRAY_BOX, ec=GRAY_EDGE, lw=1.2, r=0.9, z=2, ls="-"):
    p = FancyBboxPatch(
        (x, y), w, h,
        boxstyle=f"round,pad=0.01,rounding_size={r}",
        linewidth=lw, facecolor=fc, edgecolor=ec, linestyle=ls, zorder=z,
    )
    ax.add_patch(p)
    return p


def _stage(ax, x, y, w, h, title: str):
    """Paper-style stage container: soft blue fill + blue border + title."""
    _round(ax, x, y, w, h, fc=BLUE_SOFT, ec=BLUE_EDGE, lw=1.8, r=1.2, z=1)
    ax.text(x + 1.4, y + h - 3.2, title, fontsize=11, fontweight="bold",
            color=BLUE_TITLE, va="top", ha="left", zorder=4, fontstyle="italic")


def _arrow(ax, x1, y1, x2, y2, *, color=BLUE, lw=1.35, rad=0.0, ls="-", style="-|>"):
    ax.add_patch(
        FancyArrowPatch(
            (x1, y1), (x2, y2),
            arrowstyle=style, mutation_scale=10, linewidth=lw,
            color=color, connectionstyle=f"arc3,rad={rad}",
            linestyle=ls, zorder=5,
        )
    )


def _db(ax, x, y, w, h, label: str):
    """Cylinder database (paper icon)."""
    cx = x + w / 2
    ry = min(1.45, h * 0.16)
    ax.add_patch(Ellipse((cx, y + h - ry), w, 2 * ry, fc="#DDE7F2", ec=BLUE, lw=1.15, zorder=3))
    ax.add_patch(Rectangle((x, y + ry), w, h - 2 * ry, fc="#DDE7F2", ec="none", zorder=2))
    ax.plot([x, x], [y + ry, y + h - ry], color=BLUE, lw=1.15, zorder=3)
    ax.plot([x + w, x + w], [y + ry, y + h - ry], color=BLUE, lw=1.15, zorder=3)
    ax.add_patch(Ellipse((cx, y + ry), w, 2 * ry, fc="#C9D7E8", ec=BLUE, lw=1.15, zorder=3))
    ax.text(cx, y + h / 2, label, fontsize=7.6, fontweight="bold", color=DARK,
            ha="center", va="center", zorder=4)


def _fire(ax, x, y, s=2.2):
    """Trainable marker (paper fire icon)."""
    # outer flame
    verts = [
        (x, y),
        (x - 0.55 * s, y + 0.55 * s),
        (x - 0.25 * s, y + 1.05 * s),
        (x - 0.45 * s, y + 1.45 * s),
        (x, y + 1.85 * s),
        (x + 0.45 * s, y + 1.45 * s),
        (x + 0.25 * s, y + 1.05 * s),
        (x + 0.55 * s, y + 0.55 * s),
        (x, y),
    ]
    ax.add_patch(Polygon(verts, closed=True, fc="#E85D04", ec="#C44500", lw=0.6, zorder=8))
    # inner
    inner = [
        (x, y + 0.25 * s),
        (x - 0.22 * s, y + 0.7 * s),
        (x, y + 1.25 * s),
        (x + 0.22 * s, y + 0.7 * s),
        (x, y + 0.25 * s),
    ]
    ax.add_patch(Polygon(inner, closed=True, fc="#FFBA08", ec="none", zorder=9))


def _snow(ax, x, y, s=1.6):
    """Frozen marker — 6-arm snowflake."""
    for ang in range(0, 360, 60):
        rad = np.deg2rad(ang)
        dx, dy = s * np.cos(rad), s * np.sin(rad)
        ax.plot([x, x + dx], [y, y + dy], color="#4FA3D9", lw=1.55, solid_capstyle="round", zorder=8)
        bx, by = x + 0.55 * dx, y + 0.55 * dy
        px, py = -0.32 * dy, 0.32 * dx
        ax.plot([bx, bx + px], [by, by + py], color="#4FA3D9", lw=1.05, zorder=8)
        ax.plot([bx, bx - px], [by, by - py], color="#4FA3D9", lw=1.05, zorder=8)
    ax.add_patch(Circle((x, y), 0.28 * s, fc="#4FA3D9", ec="none", zorder=9))


def _model(ax, x, y, w, h, name: str, *, trainable: bool, sub: str = ""):
    """Model card with fire/snow status."""
    _round(ax, x, y, w, h, fc=WHITE, ec=BLUE, lw=1.5, r=0.8, z=3)
    # top tint strip
    ax.add_patch(Rectangle((x + 0.15, y + h - 3.2), w - 0.3, 3.05, fc="#EEF4FA", ec="none", zorder=3))
    ax.text(x + w / 2, y + h - 1.55, name, fontsize=9, fontweight="bold", color=DARK,
            ha="center", va="center", zorder=4)
    if sub:
        ax.text(x + w / 2, y + 2.2, sub, fontsize=7.2, color=MUTED, ha="center", va="center", zorder=4)
    if trainable:
        _fire(ax, x + w - 2.0, y + h - 2.6, s=1.7)
    else:
        _snow(ax, x + w - 2.0, y + h - 1.7, s=1.35)


def _dash_box(ax, x, y, w, h, lines, *, title=""):
    """Dashed content sample box (paper style)."""
    _round(ax, x, y, w, h, fc=WHITE, ec=BLUE_EDGE, lw=1.0, r=0.6, z=3, ls="--")
    yy = y + h - 1.6
    if title:
        ax.text(x + 0.7, yy, title, fontsize=7.5, fontweight="bold", color=BLUE_TITLE, ha="left", va="top", zorder=4)
        yy -= 2.0
    for line in lines:
        ax.text(x + 0.7, yy, line, fontsize=6.8, color=GRAY, ha="left", va="top", zorder=4)
        yy -= 1.7


def _pill(ax, x, y, w, h, text, *, fc=BLUE, tc=WHITE):
    _round(ax, x, y, w, h, fc=fc, ec=fc, lw=0, r=0.9, z=4)
    ax.text(x + w / 2, y + h / 2, text, fontsize=7.2, fontweight="bold", color=tc,
            ha="center", va="center", zorder=5)


# ── Figure 1: metrics ──────────────────────────────────────────────────────────

def fig_metrics(path: Path) -> None:
    fig = plt.figure(figsize=(13.2, 6.5), dpi=220, facecolor=WHITE)
    axd = fig.add_axes([0.04, 0.80, 0.92, 0.16])
    axd.set_xlim(0, 100)
    axd.set_ylim(0, 10)
    axd.axis("off")
    _round(axd, 0, 0.3, 48.5, 9.4, fc=RED_SOFT, ec=RED, lw=1.3, r=0.7)
    _round(axd, 51.5, 0.3, 48.5, 9.4, fc=TEAL_SOFT, ec=TEAL, lw=1.3, r=0.7)
    axd.text(2, 7.5, "mean grpo_loss", fontsize=12, fontweight="bold", color=RED)
    axd.text(2, 4.6, "Window mean of  grpo_loss = policy_term + beta * KL", fontsize=9, color=DARK)
    axd.text(2, 2.0, "More negative => policy prefers high-advantage rewrites", fontsize=8.5, color=MUTED)
    axd.text(53.5, 7.5, "mean KL", fontsize=12, fontweight="bold", color=TEAL)
    axd.text(53.5, 4.6, "Mean divergence of pi from frozen reference pi_ref", fontsize=9, color=DARK)
    axd.text(53.5, 2.0, "Moderate = stable; large spikes indicate policy drift", fontsize=8.5, color=MUTED)

    ax1 = fig.add_axes([0.07, 0.10, 0.40, 0.60])
    ax2 = fig.add_axes([0.55, 0.10, 0.40, 0.60])
    for ax, ys, title, color, ylabel in [
        (ax1, GRPO_LOSS, "(a) grpo_loss vs. training step", RED, "grpo_loss"),
        (ax2, KL_TERM, "(b) KL vs. training step", TEAL, "KL term"),
    ]:
        ax.plot(STEPS, ys, color=color, lw=2.0, marker="o", ms=2.8,
                markerfacecolor=WHITE, markeredgewidth=1.15, markeredgecolor=color)
        ax.set_title(title, fontsize=11, fontweight="bold", color=DARK, pad=8, loc="left")
        ax.set_xlabel("global step", fontsize=9, color=MUTED)
        ax.set_ylabel(ylabel, fontsize=9, color=MUTED)
        for s in ("top", "right"):
            ax.spines[s].set_visible(False)
        ax.spines["left"].set_color("#D5D5D5")
        ax.spines["bottom"].set_color("#D5D5D5")
        ax.tick_params(colors=MUTED, labelsize=8)
        ax.grid(True, axis="y", ls="--", lw=0.55, color="#E6E6E6")
        ax.set_facecolor(WHITE)
    ax1.axhline(-0.0344, color="#AAAAAA", ls=":", lw=1.0)
    ax1.text(402, -0.0344, "mean", fontsize=7.5, color=MUTED, va="center")
    fig.savefig(path, dpi=220, bbox_inches="tight", facecolor=WHITE)
    plt.close(fig)


# ── Figure 2: cold start (paper Fig.2 layout) ──────────────────────────────────

def fig_cold_start(path: Path) -> None:
    fig, ax = plt.subplots(figsize=(13.4, 7.0), dpi=220)
    _canvas(ax)
    fig.patch.set_facecolor(WHITE)

    # Stage 1 — SFT (top-left), like paper
    _stage(ax, 1.2, 55, 58, 43, "Stage 1: SFT for Rewriter Initialization")
    _db(ax, 3.2, 78, 11, 12, f"Corpus\n{N_GROUPS}")
    _round(ax, 16.5, 79.5, 12, 9.5, fc=TEAL_SOFT, ec=TEAL, lw=1.2, r=0.6)
    ax.text(22.5, 85.5, "K-way rewrite", fontsize=8, fontweight="bold", color=TEAL, ha="center", va="center")
    ax.text(22.5, 82.2, f"K = {K}", fontsize=7.5, color=MUTED, ha="center", va="center")
    _db(ax, 31, 78, 11, 12, "samples")
    _arrow(ax, 14.2, 84, 16.5, 84, color=BLUE)
    _arrow(ax, 28.5, 84, 31, 84, color=BLUE)

    _db(ax, 3.2, 60, 11, 11, "phase_a\ngroup-best")
    _arrow(ax, 14.2, 65.5, 17.5, 65.5, color=BLUE)
    _model(ax, 17.5, 59, 18, 13, "Rewriter", trainable=True, sub="Qwen2.5-7B")
    _dash_box(ax, 38, 59, 19.5, 13,
              ["group-best rewrite", "next-token prediction"], title="SFT target")
    _arrow(ax, 35.5, 65.5, 38, 65.5, color=BLUE)
    # loss label
    ax.text(26.5, 57.2, "Next Token Prediction Loss", fontsize=7.5, color=RED,
            ha="center", style="italic", zorder=6)
    _arrow(ax, 26.5, 58.5, 26.5, 59.2, color=RED, lw=1.1)

    # AlignEvaluator-like panel — Dual-head RM (top-right)
    _stage(ax, 61, 55, 37.5, 43, "Dual-head RM (train once)")
    _dash_box(ax, 63.5, 80, 15, 11,
              ["pref. pairs from", "within-group ranks"], title="Pairs")
    _model(ax, 81, 78, 15, 13.5, "Dual-head RM", trainable=True, sub="backbone frozen")
    _arrow(ax, 78.5, 85.5, 81, 85.5, color=BLUE)
    _round(ax, 63.5, 60, 32.5, 14, fc=WHITE, ec=GRAY_EDGE, lw=1.0, r=0.6)
    ax.text(79.75, 71, "Keypoint-style heads", fontsize=8, fontweight="bold", color=DARK, ha="center")
    ax.text(79.75, 67.5, "r_Q  (quality)     r_L  (length)", fontsize=8, color=GRAY, ha="center")
    ax.text(79.75, 63.5, "Export: phase_a / phase_b via r_Q", fontsize=7.5, color=MUTED, ha="center")
    _arrow(ax, 88.5, 78, 88.5, 74, color=BLUE)

    # Stage 2 — GRPO (bottom full width), like paper
    _stage(ax, 1.2, 3, 97.5, 49, "Stage 2: Policy Alignment with GRPO")
    _db(ax, 3.2, 22, 11, 14, "phase_b\nadvantages")
    _model(ax, 17, 20, 15, 16, "Rewriter / Policy", trainable=True, sub="from SFT ckpt")
    _arrow(ax, 14.2, 29, 17, 29, color=BLUE)

    # K fan-out
    ax.text(38.5, 45, f"rewrite_1 ... rewrite_K   (K={K})", fontsize=8, fontweight="bold",
            color=TEAL, ha="center")
    for i, lab in enumerate(["rewrite_1", "rewrite_2", "...", f"rewrite_{K}"]):
        yy = 40 - i * 6.2
        _round(ax, 34, yy, 12, 5.2, fc=TEAL_SOFT, ec=TEAL, lw=1.0, r=0.5)
        ax.text(40, yy + 2.6, lab, fontsize=7.5, color=TEAL, ha="center", va="center", fontweight="bold")
        _arrow(ax, 32, 28, 34, yy + 2.6, color=TEAL, lw=0.9, ls="--", style="-")

    _model(ax, 50, 20, 15, 16, "Dual-head RM", trainable=False, sub="r_Q -> advantage")
    for i in range(4):
        yy = 40 - i * 6.2
        _arrow(ax, 46, yy + 2.6, 50, 28, color=MUTED, lw=0.85, ls="--", style="-")

    _round(ax, 68, 22, 16, 12, fc=WHITE, ec=GRAY_EDGE, lw=1.1, r=0.6)
    ax.text(76, 31, "Group-wise", fontsize=8, fontweight="bold", color=DARK, ha="center")
    ax.text(76, 27.5, "advantages", fontsize=8, fontweight="bold", color=DARK, ha="center")
    ax.text(76, 24.2, "A_i within group", fontsize=7, color=MUTED, ha="center")
    _arrow(ax, 65, 28, 68, 28, color=BLUE)

    _model(ax, 86.5, 36, 10.5, 10, "pi_ref", trainable=False, sub="frozen")
    # policy optimization dashed red
    _arrow(ax, 76, 22, 24.5, 20, color=RED, lw=1.5, ls="--", rad=-0.28)
    ax.text(52, 12.5, "Policy Optimization", fontsize=9, color=RED, fontweight="bold",
            style="italic", ha="center")
    _pill(ax, 68, 8, 28, 4.2, f"1 epoch | lr=1e-6 | beta_KL=0.001 | K={K}", fc=RED)

    fig.tight_layout(pad=0.2)
    fig.savefig(path, dpi=220, bbox_inches="tight", facecolor=WHITE)
    plt.close(fig)


# ── Figure 3: online GRPO ──────────────────────────────────────────────────────

def fig_online_grpo(path: Path) -> None:
    fig, ax = plt.subplots(figsize=(13.4, 7.0), dpi=220)
    _canvas(ax)
    fig.patch.set_facecolor(WHITE)

    # Top: numbered pipeline like paper Stage 2 strip
    _stage(ax, 1.2, 72, 97.5, 26, "Online Loop: Sample -> Score -> Short GRPO")
    steps = [
        ("1", "Load policy", "merge LoRA"),
        ("2", "K-way sample", f"{N_GROUPS} x K={K}"),
        ("3", "Frozen RM", "r_Q -> phase_b"),
        ("4", "Short GRPO", "1 epoch"),
        ("5", "Export", "weights + accept"),
    ]
    xs = [4, 23.5, 43, 62.5, 82]
    for i, (x, (n, t, s)) in enumerate(zip(xs, steps)):
        _round(ax, x, 76, 14.5, 14, fc=WHITE, ec=BLUE, lw=1.35, r=0.7)
        ax.add_patch(Circle((x + 1.7, 88.2), 1.35, fc=BLUE, ec="none", zorder=6))
        ax.text(x + 1.7, 88.2, n, fontsize=8, fontweight="bold", color=WHITE, ha="center", va="center", zorder=7)
        ax.text(x + 7.25, 84.5, t, fontsize=8.5, fontweight="bold", color=DARK, ha="center", va="center")
        ax.text(x + 7.25, 80.2, s, fontsize=7.2, color=MUTED, ha="center", va="center")
        if i < 4:
            _arrow(ax, x + 14.5, 83, xs[i + 1], 83, color=BLUE, lw=1.5)

    # Bottom: detailed group update (paper Stage 2 body)
    _stage(ax, 1.2, 3, 97.5, 66, "Group Relative Update (per prompt)")
    _db(ax, 3.5, 32, 11, 16, "prompt\ngroup")
    _model(ax, 17.5, 30, 15, 18, "Rewriter", trainable=True, sub="current policy")
    _arrow(ax, 14.5, 40, 17.5, 40, color=BLUE)

    # CoT-like cloud label
    _round(ax, 19, 52, 12, 7, fc="#F3F7FB", ec=BLUE_EDGE, lw=1.0, r=1.5)
    ax.text(25, 55.5, "K-way", fontsize=8, fontweight="bold", color=BLUE_TITLE, ha="center")

    ys = [52, 40, 28, 16]
    labs = ["rewrite_1", "rewrite_2", "...", f"rewrite_{K}"]
    for y, lab in zip(ys, labs):
        _round(ax, 37, y, 13, 6.5, fc=TEAL_SOFT, ec=TEAL, lw=1.05, r=0.5)
        ax.text(43.5, y + 3.25, lab, fontsize=8, fontweight="bold", color=TEAL, ha="center", va="center")
        _arrow(ax, 32.5, 39, 37, y + 3.25, color=TEAL, lw=0.95, ls="--", style="-")

    _model(ax, 54, 30, 15, 18, "Dual-head RM", trainable=False, sub="r_Q / advantage")
    for y in ys:
        _arrow(ax, 50, y + 3.25, 54, 39, color=MUTED, lw=0.9, ls="--", style="-")

    _round(ax, 72, 32, 12, 14, fc=WHITE, ec=GRAY_EDGE, lw=1.15, r=0.6)
    ax.text(78, 42, "Group-wise", fontsize=8, fontweight="bold", color=DARK, ha="center")
    ax.text(78, 38.5, "Final Score", fontsize=8, fontweight="bold", color=DARK, ha="center")
    ax.text(78, 34.5, "advantages", fontsize=7.2, color=MUTED, ha="center")
    _arrow(ax, 69, 39, 72, 39, color=BLUE)

    _model(ax, 87, 48, 10, 11, "pi_ref", trainable=False, sub="frozen")

    # red policy loop
    _arrow(ax, 78, 32, 25, 28, color=RED, lw=1.6, ls="--", rad=-0.30)
    ax.text(52, 18, "Policy Optimization", fontsize=10, color=RED, fontweight="bold",
            style="italic", ha="center")
    _pill(ax, 58, 8, 36, 5, "pi <- pi + d    |    KL(pi || pi_ref) anchored at cold start", fc=RED)

    # user intent style annotation
    ax.annotate(
        "", xy=(54, 48), xytext=(10, 55),
        arrowprops=dict(arrowstyle="-|>", color=DARK, lw=1.0, connectionstyle="arc3,rad=-0.15"),
    )
    ax.text(18, 58, "prompt / user intent", fontsize=7.5, color=MUTED, style="italic")

    fig.tight_layout(pad=0.2)
    fig.savefig(path, dpi=220, bbox_inches="tight", facecolor=WHITE)
    plt.close(fig)


# ── PPT assembly ─────────────────────────────────────────────────────────────

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)
C_DARK = RGBColor(0x1F, 0x1F, 0x1F)
C_MUTED = RGBColor(0x6B, 0x6B, 0x6B)
C_BLUE = RGBColor(0x3A, 0x6E, 0xA5)
C_SOFT = RGBColor(0xEA, 0xF1, 0xF8)
C_LINE = RGBColor(0xC5, 0xCF, 0xD9)
C_WHITE = RGBColor(0xFF, 0xFF, 0xFF)
FONT_EN = "Calibri"
FONT_ZH = "Microsoft YaHei"


def _blank(prs: Presentation):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.background.fill
    bg.solid()
    bg.fore_color.rgb = C_WHITE
    return slide


def _set_run(run, text, *, size, bold=False, color=C_DARK, font=FONT_EN):
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = FONT_EN
    try:
        from pptx.oxml.ns import qn

        rPr = run._r.get_or_add_rPr()
        rFonts = rPr.get_or_add_rFonts()
        ea = FONT_ZH if font == FONT_ZH else FONT_EN
        rFonts.set(qn("a:ea"), ea)
        rFonts.set(qn("a:latin"), FONT_EN)
        rFonts.set(qn("a:cs"), FONT_EN)
    except Exception:
        pass


def _text(slide, l, t, w, h, text, *, size=12, bold=False, color=C_DARK, align=PP_ALIGN.LEFT, font=FONT_EN):
    box = slide.shapes.add_textbox(l, t, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    _set_run(run, text, size=size, bold=bold, color=color, font=font)
    return box


def _header(slide, section: str, title: str) -> None:
    rule = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.45), Inches(0.28), Inches(12.4), Inches(0.018))
    rule.fill.solid()
    rule.fill.fore_color.rgb = C_BLUE
    rule.line.fill.background()
    _text(slide, Inches(0.45), Inches(0.38), Inches(12.4), Inches(0.28), section, size=11, bold=True, color=C_BLUE)
    _text(slide, Inches(0.45), Inches(0.64), Inches(12.4), Inches(0.40), title, size=22, bold=True, color=C_DARK)


def add_figure_slide(prs, *, section, title, fig_path, caption):
    slide = _blank(prs)
    _header(slide, section, title)
    slide.shapes.add_picture(str(fig_path), Inches(0.30), Inches(1.10), width=Inches(12.7), height=Inches(5.20))
    _text(slide, Inches(0.45), Inches(6.40), Inches(12.4), Inches(0.90), caption, size=11, color=C_MUTED)


def add_notes_zh(prs, *, section, title, lead, bullets, foot=""):
    """English title + Chinese body (pages 3 & 5)."""
    slide = _blank(prs)
    _header(slide, section, title)

    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.45), Inches(1.20), Inches(12.4), Inches(5.65))
    card.adjustments[0] = 0.05
    card.fill.solid()
    card.fill.fore_color.rgb = C_SOFT
    card.line.color.rgb = C_LINE
    card.line.width = Pt(0.75)

    _text(slide, Inches(0.75), Inches(1.40), Inches(11.8), Inches(0.55), lead,
          size=14, bold=True, color=C_DARK, font=FONT_ZH)

    box = slide.shapes.add_textbox(Inches(0.75), Inches(2.15), Inches(11.8), Inches(3.9))
    tf = box.text_frame
    tf.word_wrap = True
    for i, (head, desc) in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(12)
        r1 = p.add_run()
        _set_run(r1, head + "  ", size=13, bold=True, color=C_BLUE, font=FONT_ZH)
        r2 = p.add_run()
        _set_run(r2, desc, size=13, bold=False, color=C_DARK, font=FONT_ZH)

    if foot:
        _text(slide, Inches(0.75), Inches(6.25), Inches(11.8), Inches(0.4), foot,
              size=11, color=C_MUTED, font=FONT_ZH)


def build() -> Path:
    FIG_DIR.mkdir(parents=True, exist_ok=True)
    p1 = FIG_DIR / "fig_metrics.png"
    p2 = FIG_DIR / "fig_cold_start.png"
    p3 = FIG_DIR / "fig_online_grpo.png"
    fig_metrics(p1)
    fig_cold_start(p2)
    fig_online_grpo(p3)

    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H

    add_figure_slide(
        prs,
        section="Figure 1  ·  Optimization Metrics",
        title="Training Dynamics of Short GRPO",
        fig_path=p1,
        caption=(
            "Figure 1. Mean grpo_loss and mean KL over one online GRPO epoch (steps 10-400). "
            "grpo_loss decreases from -0.010 to -0.113; KL reaches a mid-run minimum of 0.029 "
            "and peaks at 0.223 near the end of the epoch."
        ),
    )

    add_figure_slide(
        prs,
        section="Figure 2  ·  Cold Start",
        title="Cold-Start Training Framework",
        fig_path=p2,
        caption=(
            f"Figure 2. Cold-start pipeline (once). {N_GROUPS} prompts with K={K} rewrites; "
            "a dual-head RM is trained once; the rewriter is initialized by SFT then aligned "
            "with one short GRPO epoch (KL to the cold-start reference)."
        ),
    )

    add_notes_zh(
        prs,
        section="Notes to Figure 2",
        title="Cold Start: What Each Stage Trains",
        lead=f"数据规模：{N_GROUPS} 条原文，每组并行 K={K} 路改写（另保留原文作对照）。",
        bullets=[
            ("Stage 1 — 采数与 SFT。",
             "改写器对每条原文生成 K 路候选，导出 phase_a（组内最优）。对 Rewriter 做一次监督微调，目标为最优改写的 next-token prediction。"),
            ("Dual-head RM — 只训一次。",
             "用组内排序构造偏好对，冻住 7B backbone，只训 r_Q / r_L 两个头。之后用 r_Q 覆盖 R_content，并导出 phase_b（组内 advantage）。"),
            ("Stage 2 — 首轮短 GRPO。",
             "从 SFT 权重出发，1 个 epoch；lr=1e-6，beta_KL=0.001。参考策略 pi_ref 锚定冷启动改写器快照并保持冻结。"),
            ("可训 / 冻结。",
             "火焰标记表示本阶段更新参数；雪花标记表示冻结（RM backbone、后续的 pi_ref）。"),
        ],
        foot="冷启动只执行一次；在线轮复用冻结的双头 RM 与 pi_ref。",
    )

    add_figure_slide(
        prs,
        section="Figure 3  ·  Online GRPO",
        title="Online Policy Alignment Loop",
        fig_path=p3,
        caption=(
            f"Figure 3. Each online round loads the latest policy, re-samples {N_GROUPS} x K={K}, "
            "scores with the frozen cold-start RM, runs one short GRPO epoch, then exports "
            "weights and an acceptance summary."
        ),
    )

    add_notes_zh(
        prs,
        section="Notes to Figure 3",
        title="Online GRPO: Loop Mechanics",
        lead="每一轮外循环都在当前策略下重新采样，再做一次短 GRPO 更新。",
        bullets=[
            ("① 加载策略。",
             "将上一轮 GRPO 的 LoRA merge 进改写服务，保证采样使用最新权重。"),
            ("② K 路采样。",
             f"对 {N_GROUPS} 组原文各生成 K={K} 条改写；策略变了，候选分布随之更新。"),
            ("③ 冻结 RM 打分。",
             "复用冷启动双头 RM，用 r_Q 打分并导出新的 phase_b（组内 advantage）。"),
            ("④ 短 GRPO。",
             "1 个 epoch；policy 接上一轮 checkpoint；KL 仍锚定冷启动 pi_ref。"),
            ("⑤ 导出。",
             "写出新权重与 accept.json（质量提升 / 门限）；metrics.jsonl 记录 grpo_loss 与 KL。"),
        ],
        foot="在线阶段只更新改写器策略；奖励模型与 KL 参考策略保持固定。",
    )

    prs.save(str(OUT))
    return OUT


if __name__ == "__main__":
    print(build())
