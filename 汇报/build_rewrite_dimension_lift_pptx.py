"""论文展示：400 组改写前后宏观对比 + 得分轴子维度综合进步。

口径：每组取改写总分最高一条，减同组原文，再对组取平均。
分层：全部 / 逆解析 / 生成图。
"""

from __future__ import annotations

import json
from collections import defaultdict
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

import matplotlib.pyplot as plt
import numpy as np
from matplotlib import font_manager
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

REPO = Path(__file__).resolve().parents[1]
SAMPLES = Path(r"c:\Users\lsh\Desktop\记录10.1\10.2\samples.jsonl")
CORPUS = (
    REPO
    / "fashion_research_dir"
    / "k_rewrite_instructions_2026-09-29"
    / "source_corpus.jsonl"
)
FIG_DIR = Path(r"c:\Users\lsh\Desktop\记录10.1\_rewrite_paper_figs")
OUT = Path(r"c:\Users\lsh\Desktop\记录10.1\改写前后维度对比.pptx")

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)
BLACK = RGBColor(0x1A, 0x1A, 0x1A)
GRAY = RGBColor(0x5C, 0x5C, 0x5C)
MUTED = RGBColor(0x8A, 0x8A, 0x8A)
RED = RGBColor(0xC0, 0x22, 0x22)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
FONT = "微软雅黑"

# 论文图配色（对齐参考横条图）
POS_BAR = "#5B8DB8"
NEG_BAR = "#A8A8A8"
ORIG_BAR = "#B0B0B0"
BEST_BAR = "#C0392B"
ZERO_LINE = "#888888"
MEAN_LINE = "#222222"
GRID = "#D8D8D8"
FACE = "#FFFFFF"

COVERAGE_MODULES = (
    ("GarmentCore", "服装核心"),
    ("MaterialColor", "材质色彩"),
    ("ConstructionDetail", "结构细节"),
    ("StylingSet", "造型配饰"),
    ("Composition", "构图比例"),
)
QUALITY_MODULES = (
    ("DesignMerit", "设计价值"),
    ("InformationDensity", "信息密度"),
    ("ConcisenessAndDensity", "可见性优先级"),
    ("GenerationReadiness", "生成适配"),
    ("BindingAccuracy", "属性绑定"),
    ("LanguageClarity", "语言清晰"),
    ("StructuralClarity", "结构清晰"),
)
ALL_MODULES = COVERAGE_MODULES + QUALITY_MODULES


def _setup_font() -> str:
    for name in ("微软雅黑", "Microsoft YaHei", "SimHei", "Arial Unicode MS"):
        try:
            path = font_manager.findfont(name, fallback_to_default=False)
            if path and "DejaVu" not in path:
                plt.rcParams["font.sans-serif"] = [name]
                plt.rcParams["axes.unicode_minus"] = False
                return name
        except Exception:
            continue
    plt.rcParams["axes.unicode_minus"] = False
    return "sans-serif"


def _f(v: Any) -> Optional[float]:
    if v is None or isinstance(v, bool):
        return None
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def _is_original(rec: Dict[str, Any]) -> bool:
    try:
        if int(rec.get("candidate_index", 0)) < 0:
            return True
    except (TypeError, ValueError):
        pass
    return bool((rec.get("parallel_sampling") or {}).get("injected_original"))


def _module_val(raw: Any) -> Optional[float]:
    if isinstance(raw, dict):
        return _f(raw.get("score"))
    return _f(raw)


def _axis(rec: Dict[str, Any]) -> Dict[str, Optional[float]]:
    sc = rec.get("scores_compact") or {}
    out: Dict[str, Optional[float]] = {
        "S_fp": _f(rec.get("S_fp")) if _f(rec.get("S_fp")) is not None else _f(sc.get("S_fp")),
        "quality": _f(sc.get("quality_base_score")) or _f(sc.get("quality_penalized_score")),
        "coverage": _f(sc.get("coverage_axis_score")),
        "penalty": _f(sc.get("total_penalty")),
    }
    mods = sc.get("module_scores") or {}
    for name, _ in ALL_MODULES:
        out[name] = _module_val(mods.get(name))
    gates = rec.get("gates_compact") or {}
    out["score_gate"] = 1.0 if gates.get("score_gate_passed") else 0.0
    out["penalty_gate"] = 1.0 if gates.get("penalty_gate_passed") else 0.0
    out["both_gate"] = 1.0 if gates.get("both_passed") else 0.0
    # 若 gates 缺失，用阈值回推
    if "score_gate_passed" not in gates and out["S_fp"] is not None:
        out["score_gate"] = 1.0 if out["S_fp"] >= 0.8 else 0.0
    if "penalty_gate_passed" not in gates and out["penalty"] is not None:
        out["penalty_gate"] = 1.0 if out["penalty"] <= 0.25 else 0.0
    if "both_passed" not in gates:
        out["both_gate"] = 1.0 if (out["score_gate"] == 1.0 and out["penalty_gate"] == 1.0) else 0.0
    return out


def _kind(sample_kind: str) -> str:
    return "inverse" if sample_kind == "inverse" else "generated"


def load_pairs() -> List[Dict[str, Any]]:
    corpus_rows = [
        json.loads(x)
        for x in CORPUS.read_text(encoding="utf-8").splitlines()
        if x.strip()
    ]
    corpus = {r["source_id"]: _kind(str(r.get("sample_kind") or "")) for r in corpus_rows}

    by_gid: Dict[str, List[Dict[str, Any]]] = defaultdict(list)
    with SAMPLES.open(encoding="utf-8") as f:
        for line in f:
            line = line.strip()
            if line:
                rec = json.loads(line)
                by_gid[rec["group_id"]].append(rec)

    pairs: List[Dict[str, Any]] = []
    for gid, rows in by_gid.items():
        orig = next((r for r in rows if _is_original(r)), None)
        cands = [r for r in rows if not _is_original(r) and _f(r.get("S_fp")) is not None]
        if orig is None or not cands:
            continue
        best = max(cands, key=lambda r: float(r["S_fp"]))

        # group_id 通常以 source_id 为前缀；取最长匹配
        matches = [k for k in corpus if gid == k or gid.startswith(k)]
        if matches:
            sid = max(matches, key=len)
            kind = corpus[sid]
        else:
            kind = "inverse" if gid.startswith(("pos_", "inv_")) else "generated"

        pairs.append(
            {
                "group_id": gid,
                "kind": kind,
                "orig": _axis(orig),
                "best": _axis(best),
            }
        )
    return pairs


def summarize(pairs: List[Dict[str, Any]]) -> Dict[str, Any]:
    n = len(pairs)
    keys_macro = ["S_fp", "quality", "coverage", "penalty", "score_gate", "penalty_gate", "both_gate"]

    def mean_key(side: str, key: str) -> Optional[float]:
        vals = [p[side][key] for p in pairs if p[side].get(key) is not None]
        return float(np.mean(vals)) if vals else None

    lifts = []
    higher = lower = equal = 0
    for p in pairs:
        o, b = p["orig"]["S_fp"], p["best"]["S_fp"]
        if o is None or b is None:
            continue
        d = b - o
        lifts.append(d)
        if d > 1e-9:
            higher += 1
        elif d < -1e-9:
            lower += 1
        else:
            equal += 1

    module_lifts: Dict[str, float] = {}
    for name, _zh in ALL_MODULES:
        ds = []
        for p in pairs:
            o, b = p["orig"].get(name), p["best"].get(name)
            if o is not None and b is not None:
                ds.append(b - o)
        if ds:
            module_lifts[name] = float(np.mean(ds))

    axis_lifts = {
        "S_fp": float(np.mean(lifts)) if lifts else 0.0,
        "quality": None,
        "coverage": None,
    }
    for key in ("quality", "coverage"):
        ds = []
        for p in pairs:
            o, b = p["orig"].get(key), p["best"].get(key)
            if o is not None and b is not None:
                ds.append(b - o)
        axis_lifts[key] = float(np.mean(ds)) if ds else 0.0

    return {
        "n": n,
        "higher": higher,
        "lower": lower,
        "equal": equal,
        "orig": {k: mean_key("orig", k) for k in keys_macro},
        "best": {k: mean_key("best", k) for k in keys_macro},
        "axis_lifts": axis_lifts,
        "module_lifts": module_lifts,
    }


def fig_macro(stats: Dict[str, Any], title: str, path: Path) -> None:
    _setup_font()
    fig, axes = plt.subplots(1, 2, figsize=(11.2, 4.6), dpi=220)
    fig.patch.set_facecolor(FACE)

    cats_score = ["总分 S_fp", "质量轴 Q", "覆盖轴 C"]
    keys_score = ["S_fp", "quality", "coverage"]
    orig = [stats["orig"][k] or 0 for k in keys_score]
    best = [stats["best"][k] or 0 for k in keys_score]
    x = np.arange(len(cats_score))
    w = 0.34
    ax = axes[0]
    ax.bar(x - w / 2, orig, w, label="原文", color=ORIG_BAR, edgecolor="none")
    ax.bar(x + w / 2, best, w, label="组内最高改写", color=BEST_BAR, edgecolor="none")
    for i, (o, b) in enumerate(zip(orig, best)):
        ax.text(i - w / 2, o + 0.02, f"{o:.3f}", ha="center", va="bottom", fontsize=9, color="#444")
        ax.text(i + w / 2, b + 0.02, f"{b:.3f}", ha="center", va="bottom", fontsize=9, color="#444", fontweight="bold")
    ax.set_xticks(x)
    ax.set_xticklabels(cats_score, fontsize=11)
    ax.set_ylim(0, 1.12)
    ax.set_ylabel("分数", fontsize=11)
    ax.set_title("分数：原文 vs 组内最高改写", fontsize=12, fontweight="bold", pad=10)
    ax.legend(frameon=False, fontsize=10, loc="lower right")
    ax.spines["top"].set_visible(False)
    ax.spines["right"].set_visible(False)
    ax.yaxis.grid(True, color=GRID, linewidth=0.8)
    ax.set_axisbelow(True)

    cats_gate = ["得分门", "惩罚门", "双门"]
    keys_gate = ["score_gate", "penalty_gate", "both_gate"]
    orig_g = [(stats["orig"][k] or 0) * 100 for k in keys_gate]
    best_g = [(stats["best"][k] or 0) * 100 for k in keys_gate]
    ax = axes[1]
    ax.bar(x - w / 2, orig_g, w, label="原文", color=ORIG_BAR, edgecolor="none")
    ax.bar(x + w / 2, best_g, w, label="组内最高改写", color=BEST_BAR, edgecolor="none")
    for i, (o, b) in enumerate(zip(orig_g, best_g)):
        ax.text(i - w / 2, o + 1.2, f"{o:.1f}", ha="center", va="bottom", fontsize=9, color="#444")
        ax.text(i + w / 2, b + 1.2, f"{b:.1f}", ha="center", va="bottom", fontsize=9, color="#444", fontweight="bold")
    ax.set_xticks(x)
    ax.set_xticklabels(cats_gate, fontsize=11)
    ax.set_ylim(0, 112)
    ax.set_ylabel("通过率 (%)", fontsize=11)
    ax.set_title("门通过率（%）", fontsize=12, fontweight="bold", pad=10)
    ax.legend(frameon=False, fontsize=10, loc="lower right")
    ax.spines["top"].set_visible(False)
    ax.spines["right"].set_visible(False)
    ax.yaxis.grid(True, color=GRID, linewidth=0.8)
    ax.set_axisbelow(True)

    fig.suptitle(title, fontsize=14, fontweight="bold", y=1.02)
    fig.tight_layout()
    path.parent.mkdir(parents=True, exist_ok=True)
    fig.savefig(path, dpi=220, bbox_inches="tight", facecolor=FACE)
    plt.close(fig)


def fig_dim_lift(stats: Dict[str, Any], title: str, path: Path) -> None:
    """横条图：各得分轴子维度改写后综合进步（百分点）。"""
    _setup_font()
    rows: List[Tuple[str, float, str]] = []
    for name, zh in ALL_MODULES:
        lift = stats["module_lifts"].get(name)
        if lift is None:
            continue
        axis = "覆盖" if name in {n for n, _ in COVERAGE_MODULES} else "质量"
        rows.append((f"{zh}", lift * 100.0, axis))

    rows.sort(key=lambda x: x[1])  # 升序，最大在上时用 barh 从下往上
    labels = [r[0] for r in rows]
    vals = [r[1] for r in rows]
    colors = [POS_BAR if v >= 0 else NEG_BAR for v in vals]
    overall = float(np.mean(vals)) if vals else 0.0
    sfp_pp = (stats["axis_lifts"]["S_fp"] or 0) * 100

    fig_h = max(5.2, 0.38 * len(rows) + 1.4)
    fig, ax = plt.subplots(figsize=(9.2, fig_h), dpi=220)
    fig.patch.set_facecolor(FACE)
    y = np.arange(len(rows))
    ax.barh(y, vals, color=colors, edgecolor="none", height=0.72)
    ax.axvline(0, color=ZERO_LINE, linewidth=1.0)
    ax.axvline(overall, color=MEAN_LINE, linewidth=1.2, linestyle="--")
    for yi, v in zip(y, vals):
        offset = 0.25 if v >= 0 else -0.25
        ax.text(
            v + offset,
            yi,
            f"{v:+.1f}%",
            va="center",
            ha="left" if v >= 0 else "right",
            fontsize=9,
            color="#333",
            fontweight="bold",
        )
    ax.set_yticks(y)
    ax.set_yticklabels(labels, fontsize=10)
    ax.set_xlabel("改写后综合进步（百分点）", fontsize=11)
    ax.set_title(title, fontsize=13, fontweight="bold", pad=12)
    ax.xaxis.grid(True, color=GRID, linewidth=0.8)
    ax.set_axisbelow(True)
    ax.spines["top"].set_visible(False)
    ax.spines["right"].set_visible(False)
    # 图例
    from matplotlib.lines import Line2D
    from matplotlib.patches import Patch

    legend = [
        Patch(facecolor=POS_BAR, edgecolor="none", label="正向提升"),
        Patch(facecolor=NEG_BAR, edgecolor="none", label="负向变化"),
        Line2D([0], [0], color=MEAN_LINE, linestyle="--", linewidth=1.2, label=f"维度均值 {overall:+.1f}%"),
    ]
    ax.legend(handles=legend, loc="lower right", frameon=True, fontsize=9, framealpha=0.95)
    # 页脚注释
    fig.text(
        0.01,
        0.01,
        f"口径：组内最高改写 − 原文，再对组平均。总分 S_fp 提升 {sfp_pp:+.1f} pp。覆盖 5 维 + 质量 7 维。",
        fontsize=8,
        color="#666",
    )
    fig.tight_layout(rect=(0, 0.03, 1, 1))
    path.parent.mkdir(parents=True, exist_ok=True)
    fig.savefig(path, dpi=220, bbox_inches="tight", facecolor=FACE)
    plt.close(fig)


def _run(p, text, *, size=12, bold=False, color=BLACK):
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = color
    r.font.name = FONT


def add_text(slide, l, t, w, h, text, *, size=12, bold=False, color=BLACK, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(l, t, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    _run(p, text, size=size, bold=bold, color=color)


def add_rich(slide, l, t, w, h, segments, *, size=14):
    box = slide.shapes.add_textbox(l, t, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    for text, color, bold in segments:
        _run(p, text, size=size, bold=bold, color=color)


def add_fig_slide(prs, *, kicker: str, title: str, red_line: list, fig_path: Path, foot: str):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = WHITE
    add_text(slide, Inches(0.4), Inches(0.12), Inches(12.5), Inches(0.28), kicker, size=11, color=GRAY)
    add_text(slide, Inches(0.4), Inches(0.38), Inches(12.5), Inches(0.4), title, size=22, bold=True)
    add_rich(slide, Inches(0.4), Inches(0.82), Inches(12.5), Inches(0.42), red_line, size=14)

    # 图片适配可用高度，避免裁切
    from PIL import Image as PILImage

    box_l, box_t = Inches(0.45), Inches(1.28)
    box_w, box_h = Inches(12.4), Inches(5.55)
    with PILImage.open(fig_path) as im:
        iw, ih = im.size
    scale = min(float(box_w) / iw, float(box_h) / ih)
    pw, ph = iw * scale, ih * scale
    left = box_l + (box_w - pw) / 2
    top = box_t + (box_h - ph) / 2
    slide.shapes.add_picture(str(fig_path), int(left), int(top), int(pw), int(ph))
    add_text(slide, Inches(0.4), Inches(7.05), Inches(12.5), Inches(0.32), foot, size=10, color=MUTED)


def red_macro(stats: Dict[str, Any]) -> list:
    o, b = stats["orig"], stats["best"]
    d = (b["S_fp"] or 0) - (o["S_fp"] or 0)
    sg_o, sg_b = (o["score_gate"] or 0) * 100, (b["score_gate"] or 0) * 100
    pg_o, pg_b = (o["penalty_gate"] or 0) * 100, (b["penalty_gate"] or 0) * 100
    return [
        (f"总分 {(o['S_fp'] or 0):.3f} → {(b['S_fp'] or 0):.3f}（{d:+.3f}）。", RED, True),
        (
            f"  更高 {stats['higher']} 组，更低 {stats['lower']} 组，持平 {stats['equal']} 组。"
            f"  得分门 {sg_o:.1f}% → {sg_b:.1f}%，惩罚门 {pg_o:.1f}% → {pg_b:.1f}%。",
            BLACK,
            False,
        ),
    ]


def red_dim(stats: Dict[str, Any]) -> list:
    lifts = sorted(stats["module_lifts"].items(), key=lambda x: x[1], reverse=True)
    zh = {n: z for n, z in ALL_MODULES}
    top = lifts[0] if lifts else ("", 0.0)
    bot = lifts[-1] if lifts else ("", 0.0)
    mean_pp = float(np.mean([v * 100 for v in stats["module_lifts"].values()])) if stats["module_lifts"] else 0.0
    sfp = (stats["axis_lifts"]["S_fp"] or 0) * 100
    q = (stats["axis_lifts"]["quality"] or 0) * 100
    c = (stats["axis_lifts"]["coverage"] or 0) * 100
    return [
        (f"子维度均值 {mean_pp:+.1f} pp；S_fp {sfp:+.1f} / Q {q:+.1f} / C {c:+.1f}。", RED, True),
        (
            f"  最大提升：{zh.get(top[0], top[0])} {top[1]*100:+.1f}%；"
            f"最小：{zh.get(bot[0], bot[0])} {bot[1]*100:+.1f}%。",
            BLACK,
            False,
        ),
    ]


def main() -> None:
    _setup_font()
    pairs = load_pairs()
    buckets = {
        "all": [p for p in pairs],
        "inverse": [p for p in pairs if p["kind"] == "inverse"],
        "generated": [p for p in pairs if p["kind"] == "generated"],
    }
    stats = {k: summarize(v) for k, v in buckets.items()}
    for k, s in stats.items():
        print(k, "n=", s["n"], "S_fp lift=", round(s["axis_lifts"]["S_fp"], 4),
              "higher/lower/eq", s["higher"], s["lower"], s["equal"])

    FIG_DIR.mkdir(parents=True, exist_ok=True)
    figs = {}
    specs = [
        ("all", "全部 400 组", "全部400组 · 改写前后宏观对比"),
        ("inverse", "逆解析 200 组", "逆解析200组 · 改写前后宏观对比"),
        ("generated", "生成图 200 组", "生成图200组 · 改写前后宏观对比"),
    ]
    for key, label, macro_title in specs:
        p_macro = FIG_DIR / f"macro_{key}.png"
        p_dim = FIG_DIR / f"dim_{key}.png"
        fig_macro(stats[key], macro_title, p_macro)
        fig_dim_lift(stats[key], f"{label} · 得分轴子维度综合进步", p_dim)
        figs[key] = (p_macro, p_dim)

    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    prs.core_properties.title = "改写前后维度对比"

    pages = [
        ("01  /  全部 400 组 · 宏观", "改写器把组内最高分整体抬过原文", "all", "macro",
         "每组取改写总分最高的一条，对比同组原文。得分门阈值 0.8，惩罚门 0.25。"),
        ("02  /  全部 400 组 · 子维度", "得分轴各子维度在改写后的综合进步", "all", "dim",
         "覆盖轴 5 维 + 质量轴 7 维；横轴为组均提升（百分点）。虚线为各维度均值。"),
        ("03  /  逆解析 200 组 · 宏观", "原文已较高，改写仍小幅抬分并收紧门限", "inverse", "macro",
         "逆解析原文基线更高，提升空间相对生成图更小。"),
        ("04  /  逆解析 200 组 · 子维度", "逆解析：各得分子维度综合进步", "inverse", "dim",
         "逆解析子集上按同一口径：组内最高改写 − 原文。"),
        ("05  /  生成图 200 组 · 宏观", "改写器主要拉高的是生成图", "generated", "macro",
         "生成图原文基线低，总分与双门通过率提升最明显。"),
        ("06  /  生成图 200 组 · 子维度", "生成图：各得分子维度综合进步", "generated", "dim",
         "生成图子集上按同一口径：组内最高改写 − 原文。"),
    ]

    for kicker, title, key, kind, foot in pages:
        s = stats[key]
        fig_path = figs[key][0 if kind == "macro" else 1]
        red = red_macro(s) if kind == "macro" else red_dim(s)
        add_fig_slide(prs, kicker=kicker, title=title, red_line=red, fig_path=fig_path, foot=foot)

    # 导出数值摘要，便于论文表格
    summary_path = FIG_DIR / "lift_summary.json"
    summary_path.write_text(
        json.dumps(
            {
                k: {
                    "n": v["n"],
                    "higher": v["higher"],
                    "lower": v["lower"],
                    "equal": v["equal"],
                    "orig": v["orig"],
                    "best": v["best"],
                    "axis_lifts": v["axis_lifts"],
                    "module_lifts": v["module_lifts"],
                }
                for k, v in stats.items()
            },
            ensure_ascii=False,
            indent=2,
        ),
        encoding="utf-8",
    )

    try:
        prs.save(OUT)
        print(OUT)
    except PermissionError:
        alt = OUT.with_name(OUT.stem + "_新.pptx")
        prs.save(alt)
        print(alt)
    print("figs", FIG_DIR)
    print("summary", summary_path)


if __name__ == "__main__":
    main()
